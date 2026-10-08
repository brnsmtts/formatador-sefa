#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
LINKS DE LEGISLAÇÃO - SEFA/PA                                       versão 0.8.2

Lê um .docx já formatado, encontra as citações de leis, decretos, INs,
portarias, resoluções, convênios etc., busca cada ato no site de legislação
da SEFA/PA e grava uma CÓPIA do documento com o hyperlink sobre o nome de
cada ato (em todas as ocorrências), mais um relatório em Excel.

O arquivo original nunca é alterado.

Integrado ao Formatador SEFA (3.7+): o formatador importa este módulo e chama
criar_links() logo depois de gerar o Word, de modo que o documento já sai com
os links. Este arquivo continua funcionando sozinho para Words já prontos.

COMO ABRIR: dois cliques neste arquivo. Na primeira vez, o próprio programa instala
os componentes de que precisa (leva alguns minutos e precisa de internet).

Uso pelo Prompt de Comando (opcional):
    python links_legislacao.py                 -> abre a janela
    python links_legislacao.py arquivo.docx    -> processa direto
    python links_legislacao.py --diagnostico   -> grava como o site responde (p/ ajustes)

Dependências:  pip install python-docx openpyxl truststore playwright
"""

import copy
import csv
import json
import queue
import re
import sys
import threading
import time
import traceback
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime
from pathlib import Path

# -----------------------------------------------------------------------------
# Primeira execução: instala sozinho os componentes que faltarem
# -----------------------------------------------------------------------------
_COMPONENTES = (("docx", "python-docx"), ("openpyxl", "openpyxl"), ("playwright", "playwright"),
                ("truststore", "truststore"))


def _aviso(titulo, texto, erro=False):
    try:
        import tkinter as tk
        from tkinter import messagebox
        r = tk.Tk()
        r.withdraw()
        (messagebox.showerror if erro else messagebox.showinfo)(titulo, texto)
        r.destroy()
    except Exception:
        print(f"{titulo}: {texto}")


def _garantir_componentes():
    import importlib
    import subprocess
    faltando = []
    for modulo, pacote in _COMPONENTES:
        try:
            importlib.import_module(modulo)
        except ImportError:
            faltando.append(pacote)
    if not faltando:
        return
    janela_aviso = None
    try:
        import tkinter as tk
        janela_aviso = tk.Tk()
        janela_aviso.title("Links de Legislação - SEFA/PA")
        tk.Label(janela_aviso, justify="left", padx=30, pady=25, font=("Segoe UI", 11),
                 text="Preparando o programa pela primeira vez...\n\n"
                      "Instalando os componentes necessários.\n"
                      "Isso leva alguns minutos. Não feche esta janela.").pack()
        janela_aviso.update()
    except Exception:
        print("Instalando os componentes necessários (primeira execução)...")
    em_venv = sys.prefix != getattr(sys, "base_prefix", sys.prefix)
    cmd = [sys.executable, "-m", "pip", "install", "--disable-pip-version-check",
           *([] if em_venv else ["--user"]), *faltando]
    try:
        proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
    except Exception as e:
        saida, codigo = str(e), 1
    else:
        linhas = []
        while True:
            linha = proc.stdout.readline()
            if linha:
                linhas.append(linha)
            elif proc.poll() is not None:
                break
            if janela_aviso is not None:
                janela_aviso.update()
        saida, codigo = "".join(linhas[-15:]), proc.returncode
    if janela_aviso is not None:
        janela_aviso.destroy()
    if codigo != 0:
        _aviso("Links de Legislação - erro na instalação",
               "Não foi possível instalar os componentes necessários.\n"
               "Verifique a conexão com a internet e tente de novo.\n\nDetalhes:\n" + saida[-1500:], erro=True)
        sys.exit(1)
    import site
    importlib.invalidate_caches()
    pasta_usuario = site.getusersitepackages()
    if pasta_usuario not in sys.path:
        sys.path.append(pasta_usuario)


if __name__ == "__main__":
    # só quando aberto diretamente; importado pelo Formatador, quem instala é o INSTALAR.bat
    _garantir_componentes()

try:  # usa os certificados do Windows (resolve erro de certificado em sites .gov.br)
    import truststore
    truststore.inject_into_ssl()
except Exception:
    pass

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

# =============================================================================
# CONFIGURAÇÃO
# =============================================================================
CONFIG = {
    # Site de busca de legislação da SEFA/PA (aplicativo que conversa com a API abaixo)
    "pagina_busca": "https://app.sefa.pa.gov.br/busca-legislacao/buscar-legislacoes/buscar",
    "api": "https://app.sefa.pa.gov.br/api-gateway/busca-legislacao-api/extranet",
    # Página que abre um ato. O link gravado segue o mesmo formato que o próprio site usa:
    # ...visualizacao?idStatus=3&numeroAto=5530&ano=1989&ordenacao=MAIS_RECENTES&idAto=105
    "pagina_ato": "https://app.sefa.pa.gov.br/busca-legislacao/buscar-legislacoes/buscar/resultados/visualizacao",
    "timeout": 30,                 # segundos de espera por resposta do site
    "pausa": 0.2,                  # segundos entre uma busca e outra
    "max_falhas_seguidas": 3,      # para de consultar o site após N erros seguidos
    "mostrar_navegador": False,    # True = deixa a janela do Edge visível durante as buscas

    # True -> link em "Lei nº 5.530, de 13 de janeiro de 1989"; False -> só em "Lei nº 5.530"
    "incluir_data_no_link": True,
    # não linka o título do próprio ato ("DECRETO Nº 1.234, DE ...") no topo
    "ignorar_epigrafe": True,
    "linhas_epigrafe": 10,
    # a ementa do ato não recebe links
    "linkar_ementa": False,
    # destaca em amarelo as citações que ficaram sem link
    "destacar_pendentes": True,
}

# Tipos de ato que existem no site da SEFA (os demais - convênios, ajustes, protocolos... -
# não são buscados).
TIPOS_NO_SITE = {"Lei", "Lei Complementar", "Decreto", "Decreto Legislativo", "Instrução Normativa",
                 "Instrução Normativa Conjunta", "Portaria", "Portaria Conjunta", "Resolução", "Ato",
                 "Emenda Constitucional", "Constituição"}

# Esferas do site (categoria.esfera.nome na API), já normalizadas
ESFERAS = {"FEDERAL": "Federal", "ESTADUAL": "Estadual", "CGIBS": "CGIBS"}

PASTA_PROGRAMA = Path(__file__).resolve().parent
ARQ_TABELA = PASTA_PROGRAMA / "tabela_links.tsv"

S_TABELA = "Link criado (tabela local)"
S_SITE = "Link criado (site)"
S_SITE_SEM_ANO = "Link criado (site) - conferir: ano não confirmado"
S_NAO_ACHOU = "Não encontrado no site"
S_FORA_SITE = "Tipo de ato que não consta no site - não buscado"
S_SITE_OFF = "Não buscado (site fora do ar / sem conexão)"
S_INTERROMPIDO = "Não buscado (busca interrompida)"
S_OUTRO_ANO = "Não encontrado: no site há ato com esse número, mas de outro ano"
SEM_DESTAQUE = {S_FORA_SITE}


# =============================================================================
# RECONHECIMENTO DAS CITAÇÕES
# =============================================================================
TIPOS = [
    "Lei Complementar", "Lei Ordinária", "Decreto-Lei", "Decreto Legislativo", "Decreto", "Lei",
    "Instrução Normativa Conjunta", "Instrução Normativa", "Portaria Conjunta", "Portaria",
    "Resolução", "Convênio ICMS", "Convênio", "Ajuste SINIEF", "Protocolo ICMS", "Protocolo",
    "Ato COTEPE/ICMS", "Ato COTEPE", "Ato", "Emenda Constitucional", "Ordem de Serviço",
    "Norma de Execução", "IN", "LC",
]

# Plurais ("Leis Complementares nºs 101, de ..., 105, de ..., e 215, de ...") -> singular
PLURAIS = {
    "Leis Complementares": "Lei Complementar", "Leis": "Lei", "Decretos-Leis": "Decreto-Lei",
    "Decretos-Lei": "Decreto-Lei", "Decretos Legislativos": "Decreto Legislativo", "Decretos": "Decreto",
    "Instruções Normativas": "Instrução Normativa", "Portarias": "Portaria", "Resoluções": "Resolução",
    "Emendas Constitucionais": "Emenda Constitucional", "Convênios ICMS": "Convênio ICMS",
    "Ajustes SINIEF": "Ajuste SINIEF", "Protocolos ICMS": "Protocolo ICMS",
}


_SP, _SPO = r"[ \t\xA0]+", r"[ \t\xA0]*"


def _alternativas(nomes):
    alts = []
    for t in sorted(nomes, key=len, reverse=True):
        alts.append(re.escape(t).replace(r"\ ", _SP))
        if t.upper() != t:
            alts.append(re.escape(t.upper()).replace(r"\ ", _SP))
    return "|".join(alts)


_QUAL = (r"(?:(?:" + _SP + r"|/)(?:(?:do|da)" + _SP + r")?(?:Federal|Federais|Estadual|Estaduais|Municipal|"
         r"FEDERAL|FEDERAIS|ESTADUAL|ESTADUAIS|MUNICIPAL|[A-Z]{2,}[A-Z0-9]*(?:/[A-Z0-9]+)*))")
_MARCA = r"(?:[nN]\.?[º°oO]s?\.?|[nN]\.|n[úu]meros?)"     # nº, n°, n.º, nos, n. (DOE de 02.10.2026)
_DE = r"(?:de|DE)"
# Órgão colado ao número ("Portaria nº 293/2026-SEFA/GS", "IN nº 12-SEFA.GS", "Portaria nº 5/GS"):
# entra no link, mas não muda a chave do ato nem a esfera.
_ORGAO = r"[-\u2013/][A-Z][A-Z0-9]+(?:[-./][A-Z][A-Z0-9]+)*(?![A-Za-z\u00C0-\u00FF])"


def _item(suf=""):
    """número + ano opcional (/89 ou ', de 13 de janeiro de 1989' ou ', de 1995')."""
    return (r"(?P<num" + suf + r">\d{1,3}(?:\.\d{3})+|\d+)(?!\d)(?!\.\d)(?:/[ \t\xA0]?(?P<ano2" + suf + r">\d{4}|\d{2})(?!\d)(?![./-]\d))?"
            r"(?P<orgao" + suf + r">" + _ORGAO + r")?"
            r"(?P<fimnucleo" + suf + r">)"
            r"(?P<data" + suf + r">,?" + _SP + _DE + _SP + r"(?:(?P<dia" + suf + r">\d{1,2})[º°o]?" + _SP + _DE +
            _SP + r"(?P<mes" + suf + r">[A-Za-zçÇ]+)(?:" + _SP + _DE + r")?" + _SP + r"(?P<ano4" + suf + r">\d{4})|(?P<ano5" +
            suf + r">\d{4})(?![\d/])))?")


def _padrao_citacao():
    return re.compile(r"\b(?P<nucleo>(?P<tipo>" + _alternativas(TIPOS) + ")" + _QUAL + "{0,3}(?:" +
                      _SP + _MARCA + _SPO + "|" + _SP + ")" + r"(?P<num>\d{1,3}(?:\.\d{3})+|\d+)(?!\d)(?!\.\d)"
                      r"(?:/[ \t\xA0]?(?P<ano2>\d{4}|\d{2})(?!\d)(?![./-]\d))?"     # 5.810/ 24.01.1994: 24 é dia, não ano
                      r"(?P<orgao>" + _ORGAO + r")?)"
                      r"(?P<data>,?" + _SP + _DE + _SP + r"(?:(?P<dia>\d{1,2})[º°o]?" + _SP + _DE + _SP +
                      r"(?P<mes>[A-Za-zçÇ]+)(?:" + _SP + _DE + r")?" + _SP + r"(?P<ano4>\d{4})|(?P<ano5>\d{4})(?![\d/])))?")


# fim de um item de lista: precisa vir pontuação, "e", parêntese ou fim - evita "e 3 dias"
_FIM_ITEM = r"(?=" + _SPO + r"(?:[,;.:)(]|e" + _SP + r"|$))"
RE_PLURAL = re.compile(r"\b(?P<cab>(?P<tipo>" + _alternativas(PLURAIS) + ")" + _QUAL + "{0,3}(?:" + _SP + _MARCA +
                       _SPO + "|" + _SP + "))" + _item())
RE_PLURAL_SEGUINTE = re.compile(r"(?:" + _SPO + r"\([^()]{1,150}\))?(?:" + _SPO + r"," + _SPO + r"(?:e" + _SP +
                                r")?|" + _SP + r"e" + _SP + r")" + _item())
RE_FIM_ITEM = re.compile(_FIM_ITEM)


def item_valido(m, txt):
    """Item de lista com ano é aceito; sem ano, só se vier pontuação/"e"/fim logo depois."""
    return bool(m.group("ano2") or m.group("ano4") or m.group("ano5")) or bool(RE_FIM_ITEM.match(txt, m.end()))
RE_CITACAO = _padrao_citacao()

# Constituições (citadas sem número): "Constituição Federal", "Constituição da República",
# "CF/88", "Constituição Estadual", "Constituição do Estado do Pará"...
_SPC = r"[ \t\xA0]+"
RE_CONSTITUICAO = re.compile(
    r"\b(?:(?P<fed>Constitui[çc][ãa]o" + _SPC + r"(?:Federal|da" + _SPC + r"Rep[úu]blica(?:" + _SPC +
    r"Federativa" + _SPC + r"do" + _SPC + r"Brasil)?)(?:,?" + _SPC + r"de" + _SPC + r"1988)?"
    r"|CF" + r"[ \t\xA0]*/[ \t\xA0]*(?:19)?88\b)"
    r"|(?P<est>Constitui[çc][ãa]o" + _SPC + r"(?:Estadual|do" + _SPC + r"Estado(?:" + _SPC + r"do" + _SPC +
    r"Par[áa]|(?!" + _SPC + r"(?:de|do|da)" + _SPC + r"[^\W\d_])))(?:,?" + _SPC + r"de" + _SPC + r"1989)?))",
    re.IGNORECASE)

RE_PROJETO = re.compile(r"(?i)\b(?:ante)?projetos?[ \t\xA0]+de[ \t\xA0]*$")

# "Constituição" sozinha (vale a esfera do documento), "Constituição de 1988" e o ADCT
RE_CONST_SOLTA = re.compile(
    r"\b(?:(?P<adct>Ato" + _SP + r"das" + _SP + r"Disposi[çc][õo]es" + _SP + r"Constitucionais" + _SP +
    r"Transit[óo]rias|ADCT)\b|Constitui[çc][ãa]o(?P<c88>" + _SP + r"de" + _SP + r"1988)?(?!" + _SP +
    r"(?:Federal|Estadual|da|do|de)\b)(?![\w]))")

CANON = {
    "LEI": "Lei", "LEI ORDINARIA": "Lei", "LEI COMPLEMENTAR": "Lei Complementar",
    "DECRETO": "Decreto", "DECRETO-LEI": "Decreto-Lei", "DECRETO LEGISLATIVO": "Decreto Legislativo",
    "IN": "Instrução Normativa", "INSTRUCAO NORMATIVA": "Instrução Normativa",
    "INSTRUCAO NORMATIVA CONJUNTA": "Instrução Normativa Conjunta",
    "PORTARIA": "Portaria", "PORTARIA CONJUNTA": "Portaria Conjunta", "RESOLUCAO": "Resolução",
    "CONVENIO ICMS": "Convênio ICMS", "CONVENIO": "Convênio", "AJUSTE SINIEF": "Ajuste SINIEF",
    "PROTOCOLO ICMS": "Protocolo ICMS", "PROTOCOLO": "Protocolo",
    "ATO COTEPE": "Ato COTEPE/ICMS", "ATO COTEPE/ICMS": "Ato COTEPE/ICMS",
    "ATO": "Ato", "EMENDA CONSTITUCIONAL": "Emenda Constitucional", "ORDEM DE SERVICO": "Ordem de Serviço",
    "NORMA DE EXECUCAO": "Norma de Execução", "LC": "Lei Complementar",
}


def sem_acentos(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def limpar_espacos(s):
    return re.sub(r"[\s\xA0]+", " ", s).strip()


def normalizar(s):
    return sem_acentos(limpar_espacos(s).upper())


def canon_tipo(s):
    return CANON.get(normalizar(s), limpar_espacos(s))


def esfera_citada(nucleo):
    """Esfera dita na própria citação: "Lei Complementar Federal nº 87", "Resolução CGIBS nº 1"."""
    n = normalizar(nucleo)
    if re.search(r"\bFEDERA(L|IS)\b", n):
        return "FEDERAL"
    if re.search(r"\bCG[ /-]?IBS\b", n):
        return "CGIBS"
    if re.search(r"\bESTADUA(L|IS)\b", n):
        return "ESTADUAL"
    return ""


def esfera_documento(textos):
    """Esfera do próprio documento, pelos primeiros parágrafos (usada para desempates, para a
    'Constituição' citada sozinha e para decidir quando procurar no Planalto)."""
    t = normalizar(" ".join(textos[:15]))
    marcas = [("FEDERAL", r"PRESIDENTE DA REPUBLICA|CONGRESSO NACIONAL|PUBLICAD[AO] NO DOU\b|"
                          r"DIARIO OFICIAL DA UNIAO"),
              ("ESTADUAL", r"GOVERNADORA? DO ESTADO|ASSEMBLEIA LEGISLATIVA|SECRETARI[OA] DE ESTADO DA FAZENDA|"
                           r"PUBLICAD[AO] NO DOE\b|DIARIO OFICIAL DO ESTADO"),
              ("CGIBS", r"(?:PRESIDENTE|CONSELHO SUPERIOR) DO COMITE GESTOR")]
    achados = []
    for esf, pad in marcas:
        m = re.search(pad, t)
        if m:
            achados.append((m.start(), esf))
    return min(achados)[1] if achados else ""


_VERBOS_EMENTA = (r"^(DISPOE|ALTERA|INSTITUI|REGULAMENTA|ESTABELECE|APROVA|ACRESCENTA|REVOGA|AUTORIZA|"
                  r"CONCEDE|DEFINE|FIXA|CONSOLIDA|PRORROGA|CRIA|INTRODUZ|MODIFICA|REDUZ|TORNA|DA NOVA|"
                  r"HOMOLOGA|DISCIPLINA|DIVULGA|DISPENSA|RATIFICA|INSERE|INCLUI|EXCLUI|SUSPENDE|DECLARA)\b")
_NOTA_PUBLICACAO = r"^(PUBLICAD|REPUBLICAD|RETIFICAD|VIDE\b|ALTERAD|REVOGAD|REGULAMENTAD|CONSOLIDAD)"


def _recuo_esquerdo(par):
    """Recuo à esquerda efetivo (direto ou herdado do estilo), em EMU."""
    if par.paragraph_format.left_indent is not None:
        return par.paragraph_format.left_indent
    est, vistos = par.style, set()
    while est is not None and est.style_id not in vistos:
        vistos.add(est.style_id)
        if est.paragraph_format.left_indent is not None:
            return est.paragraph_format.left_indent
        est = est.base_style
    return 0


def paragrafos_de_ementa(doc, paragrafos, textos):
    """Parágrafos que são a ementa do ato: pelo estilo (modelos da SEFA: "a03_ementa") ou,
    sem esse estilo, o 1º parágrafo depois do título e das notas de publicação, se estiver
    recuado à direita ou começar com "Dispõe", "Altera", "Institui"..."""
    from docx.text.paragraph import Paragraph
    ementa = set()
    for p in paragrafos:
        try:
            est = Paragraph(p, doc._body).style
            if "EMENTA" in normalizar(f"{est.style_id or ''} {est.name or ''}"):
                ementa.add(id(p))
        except Exception:
            pass
    if ementa:
        return ementa
    corpo = [p for p in doc.element.body if p.tag == qn("w:p")]
    txt_de = dict(zip(map(id, paragrafos), textos))
    viu_titulo = False
    for p in [p for p in corpo if txt_de.get(id(p), "").strip(BARREIRA + " \t")][:15]:
        t = normalizar(txt_de[id(p)])
        if not viu_titulo:
            viu_titulo = bool(re.match(r"^(LEI|DECRETO|INSTRUCAO|PORTARIA|RESOLUCAO|EMENDA|ATO|CONVENIO|"
                                       r"AJUSTE|PROTOCOLO|ORDEM|NORMA|CONSTITUICAO)\b", t))
            continue
        if re.match(_NOTA_PUBLICACAO, t) or t.startswith("DOCUMENTO DE TESTE"):
            continue
        par = Paragraph(p, doc._body)
        if _recuo_esquerdo(par) >= 1080000 or re.match(_VERBOS_EMENTA, t):     # >= 3 cm
            ementa.add(id(p))
        break
    return ementa


def norm_numero(s):
    return s.replace(".", "").lstrip("0") or "0"


def ano_completo(a):
    if not a:
        return ""
    if len(a) == 4:
        return a
    yy = int(a)
    return str(2000 + yy if yy <= datetime.now().year % 100 else 1900 + yy)


# =============================================================================
# MANIPULAÇÃO DO .DOCX (nível XML, preservando a formatação)
# =============================================================================
W_R, W_T, W_RPR, W_TAB = qn("w:r"), qn("w:t"), qn("w:rPr"), qn("w:tab")
W_FLDCHAR, W_FLDTYPE = qn("w:fldChar"), qn("w:fldCharType")
SEM_TEXTO = {qn("w:pPr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd"), qn("w:proofErr"),
             qn("w:commentRangeStart"), qn("w:commentRangeEnd"), qn("w:permStart"),
             qn("w:permEnd"), qn("w:moveFromRangeStart"), qn("w:moveFromRangeEnd"),
             qn("w:moveToRangeStart"), qn("w:moveToRangeEnd")}
IGNORAVEIS_NO_RUN = {W_RPR, qn("w:lastRenderedPageBreak")}
BARREIRA = "\x00"   # impede que uma citação "atravesse" links, campos etc.


def _set_texto(t_elem, texto):
    t_elem.text = texto
    t_elem.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def normalizar_runs(p):
    """Divide runs com vários conteúdos (texto+tab+texto...) em runs de um conteúdo só."""
    for r in list(p.findall(W_R)):
        conteudo = [c for c in r if c.tag not in IGNORAVEIS_NO_RUN]
        if len(conteudo) <= 1:
            continue
        rpr = r.find(W_RPR)
        for c in conteudo:
            novo = OxmlElement("w:r")
            if rpr is not None:
                novo.append(copy.deepcopy(rpr))
            novo.append(c)          # move o elemento para o novo run
            r.addprevious(novo)
        p.remove(r)


def segmentos(p):
    """Lista (run|None, texto) na ordem do parágrafo. None = trecho que não pode ser linkado."""
    segs, prof = [], 0
    for ch in p:
        if ch.tag == W_R:
            fld = ch.find(W_FLDCHAR)
            if fld is not None:
                tp = fld.get(W_FLDTYPE)
                if tp == "begin":
                    prof += 1
                elif tp == "end":
                    prof = max(0, prof - 1)
                segs.append((None, BARREIRA))
                continue
            if prof > 0:
                segs.append((None, BARREIRA))
                continue
            conteudo = [c for c in ch if c.tag not in IGNORAVEIS_NO_RUN]
            if not conteudo:
                segs.append((ch, ""))
            elif conteudo[0].tag == W_T:
                segs.append((ch, conteudo[0].text or ""))
            elif conteudo[0].tag == W_TAB:
                segs.append((ch, "\t"))
            else:
                segs.append((None, BARREIRA))
        elif ch.tag in SEM_TEXTO:
            continue
        else:                       # hyperlink existente, campo simples, revisão, controle...
            segs.append((None, BARREIRA))
    return segs


def texto_paragrafo(p):
    return "".join(t for _, t in segmentos(p))


def _dividir_em(p, pos):
    """Garante que exista uma fronteira de run exatamente na posição pos."""
    a = 0
    for run, txt in segmentos(p):
        b = a + len(txt)
        if run is not None and a < pos < b:
            t = run.find(W_T)
            esquerda = copy.deepcopy(run)
            _set_texto(esquerda.find(W_T), txt[:pos - a])
            _set_texto(t, txt[pos - a:])
            run.addprevious(esquerda)
            return
        a = b


def isolar_runs(p, ini, fim):
    """Divide os runs e devolve a lista de elementos (irmãos) que cobrem [ini, fim)."""
    _dividir_em(p, ini)
    _dividir_em(p, fim)
    a, cobertos = 0, []
    for run, txt in segmentos(p):
        b = a + len(txt)
        if run is not None and txt and a >= ini and b <= fim:
            cobertos.append(run)
        a = b
    if not cobertos:
        return []
    filhos = list(p)
    i1, i2 = filhos.index(cobertos[0]), filhos.index(cobertos[-1])
    return filhos[i1:i2 + 1]


def _rpr(run):
    return run.get_or_add_rPr() if hasattr(run, "get_or_add_rPr") else None


_ORDEM_RPR = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
              "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
              "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
              "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
_ORDEM_IDX = {qn("w:" + t): i for i, t in enumerate(_ORDEM_RPR)}
_LIGA_DESLIGA = {qn("w:" + t) for t in ("b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
                                         "outline", "shadow", "emboss", "imprint", "vanish", "webHidden")}
_CARA_DE_LINK = {qn("w:color"), qn("w:u"), qn("w:rStyle")}


def _por_na_ordem(rpr, el):
    """Insere uma propriedade no rPr respeitando a ordem exigida pelo Word."""
    idx = _ORDEM_IDX.get(el.tag, 999)
    for ch in rpr:
        if _ORDEM_IDX.get(ch.tag, 999) > idx:
            ch.addprevious(el)
            return
    rpr.append(el)


class Estilos:
    """Descobre a formatação efetiva de um trecho (direta -> estilo de caractere -> estilo do
    parágrafo -> padrão do documento), para o link não mudar tamanho/fonte do texto."""

    def __init__(self, doc, estilo_link):
        el = doc.styles.element
        self.por_id = {st.get(qn("w:styleId")): st for st in el.findall(qn("w:style"))}
        self.par_padrao = next((sid for sid, st in self.por_id.items()
                                if st.get(qn("w:type")) == "paragraph"
                                and st.get(qn("w:default")) in ("1", "true", "on")), None)
        dd = el.find(qn("w:docDefaults"))
        self.padrao = dd.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}") if dd is not None else None
        # propriedades que o estilo "Hyperlink" do documento impõe além de cor e sublinhado
        # (ex.: nos modelos da SEFA ele tem fonte 8, pensada para as notas de publicação)
        self.extras_do_link = set()
        sid, vistos = estilo_link, set()
        while sid and sid not in vistos:
            vistos.add(sid)
            st = self.por_id.get(sid)
            if st is None:
                break
            rpr = st.find(qn("w:rPr"))
            if rpr is not None:
                self.extras_do_link |= {ch.tag for ch in rpr if ch.tag not in _CARA_DE_LINK}
            b = st.find(qn("w:basedOn"))
            sid = b.get(qn("w:val")) if b is not None else None

    def _na_cadeia(self, sid, tag):
        vistos = set()
        while sid and sid not in vistos:
            vistos.add(sid)
            st = self.por_id.get(sid)
            if st is None:
                return None
            rpr = st.find(qn("w:rPr"))
            if rpr is not None and rpr.find(tag) is not None:
                return rpr.find(tag)
            b = st.find(qn("w:basedOn"))
            sid = b.get(qn("w:val")) if b is not None else None
        return None

    def efetiva(self, p, r, tag):
        rpr = r.find(W_RPR)
        if rpr is not None:
            if rpr.find(tag) is not None:
                return rpr.find(tag)
            rs = rpr.find(qn("w:rStyle"))
            if rs is not None:
                e = self._na_cadeia(rs.get(qn("w:val")), tag)
                if e is not None:
                    return e
        ppr = p.find(qn("w:pPr"))
        ps = ppr.find(qn("w:pStyle")) if ppr is not None else None
        e = self._na_cadeia(ps.get(qn("w:val")) if ps is not None else self.par_padrao, tag)
        if e is not None:
            return e
        return self.padrao.find(tag) if self.padrao is not None else None


def criar_hyperlink(p, nos, rid, estilo_id, estilos=None):
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), rid)
    hl.set(qn("w:history"), "1")
    nos[0].addprevious(hl)
    for n in nos:
        if n.tag == W_R and estilos is not None:
            # guarda o tamanho/fonte que o texto tinha antes de receber o estilo do link
            manter = {}
            atual = n.find(W_RPR)
            for tag in estilos.extras_do_link:
                if atual is None or atual.find(tag) is None:
                    manter[tag] = estilos.efetiva(p, n, tag)
        hl.append(n)
        if n.tag == W_R:
            rpr = _rpr(n)
            if rpr is not None:
                if rpr.highlight_val == WD_COLOR_INDEX.YELLOW:
                    rpr.highlight_val = None
                rpr.style = estilo_id
                if estilos is not None:
                    for tag, el in manter.items():
                        if rpr.find(tag) is not None:
                            continue
                        if el is not None:
                            novo = copy.deepcopy(el)
                        elif tag in (qn("w:sz"), qn("w:szCs")):
                            novo = OxmlElement("w:" + tag.split("}")[1])
                            novo.set(qn("w:val"), "20")          # padrão do Word: 10 pt
                        elif tag in _LIGA_DESLIGA:
                            novo = OxmlElement("w:" + tag.split("}")[1])
                            novo.set(qn("w:val"), "0")
                        else:
                            continue
                        _por_na_ordem(rpr, novo)


def destacar(nos):
    for n in nos:
        if n.tag == W_R:
            rpr = _rpr(n)
            if rpr is not None:
                rpr.highlight_val = WD_COLOR_INDEX.YELLOW


def estilo_hyperlink(doc):
    for s in doc.styles:
        try:
            if s.type == WD_STYLE_TYPE.CHARACTER and s.name == "Hyperlink":
                return s.style_id
        except Exception:
            pass
    st = doc.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    st.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    st.font.underline = True
    st.unhide_when_used = True
    st.priority = 99
    return st.style_id


# =============================================================================
# TABELA LOCAL DE LINKS
# =============================================================================
def carregar_tabela():
    d = {}
    if ARQ_TABELA.exists():
        for linha in ARQ_TABELA.read_text(encoding="utf-8-sig").splitlines():
            linha = linha.strip()
            if linha and not linha.startswith("#") and "\t" in linha:
                k, v = linha.split("\t", 1)
                if k.strip() and v.strip():
                    d[k.strip()] = v.strip()
    return d


def salvar_tabela(d):
    linhas = ["# Tabela de links de legislação (Links de Legislação - SEFA/PA)",
              "# Formato: Tipo|Número|Ano[|Esfera]<TAB>endereço  (número sem pontos; ano com 4 dígitos ou vazio;\n# Esfera = Federal, Estadual ou CGIBS, só quando a citação diz a esfera)",
              "# Pode editar à mão: as linhas daqui têm prioridade sobre a busca no site."]
    linhas += [f"{k}\t{v}" for k, v in d.items()]
    ARQ_TABELA.write_text("\n".join(linhas) + "\n", encoding="utf-8")


# =============================================================================
# CONSULTA AO SITE
# O site da SEFA não responde a programas comuns, só a navegadores. Por isso as
# consultas são feitas por dentro de um Edge (ou Chrome) aberto pelo Playwright,
# exatamente como o próprio site faz quando você pesquisa.
# =============================================================================
_UA_NAVEGADOR = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
                 "Chrome/140.0.0.0 Safari/537.36 Edg/140.0.0.0")

_JS_FETCH = """async ([url, ms]) => {
    const ctl = new AbortController();
    const t = setTimeout(() => ctl.abort(), ms);
    try {
        const r = await fetch(url, {signal: ctl.signal, credentials: 'include',
            headers: {'Accept': 'application/json, text/plain, */*', 'Authorization': 'Bearer'}});
        return {status: r.status, texto: await r.text()};
    } catch (e) {
        return {status: 0, texto: String(e)};
    } finally { clearTimeout(t); }
}"""


class SiteSEFA:
    """Abre o site uma vez e faz todas as buscas por ele. Use com 'with'."""

    def __init__(self, log=print):
        self.log = log
        self.pw = self.navegador = self.pagina = None

    def __enter__(self):
        from playwright.sync_api import sync_playwright
        self.pw = sync_playwright().start()
        visivel = CONFIG["mostrar_navegador"]
        for tentativa in ((visivel,) if visivel else (False, True)):
            if self._abrir(visivel=tentativa):
                return self
            self._fechar_navegador()
        raise ConnectionError("não foi possível abrir o site de legislação da SEFA")

    def __exit__(self, *exc):
        self._fechar_navegador()
        if self.pw:
            self.pw.stop()

    def _abrir(self, visivel):
        args = [] if not visivel else ["--start-minimized"]
        erro = None
        for canal in ("msedge", "chrome", None):
            try:
                self.navegador = self.pw.chromium.launch(channel=canal, headless=not visivel, args=args)
                break
            except Exception as e:
                self.navegador, erro = None, e
        if self.navegador is None:
            self.log(f"  ! não consegui abrir o Edge/Chrome: {str(erro)[:150]}")
            return False
        ctx = self.navegador.new_context(user_agent=_UA_NAVEGADOR, locale="pt-BR")
        self.pagina = ctx.new_page()
        self.log("  Abrindo o site de legislação da SEFA" + (" (janela minimizada)" if visivel else "") + "...")
        try:
            self.pagina.goto(CONFIG["pagina_busca"], wait_until="domcontentloaded",
                             timeout=CONFIG["timeout"] * 3000)
        except Exception as e:
            self.log(f"  ! o site não abriu: {str(e)[:150]}")
            return False
        st, _ = self.get("/situacao")          # teste rápido da API
        if st != 200:
            self.log(f"  ! a API de busca não respondeu (status {st})")
            return False
        return True

    def _fechar_navegador(self):
        try:
            if self.navegador:
                self.navegador.close()
        except Exception:
            pass
        self.navegador = self.pagina = None

    def get(self, caminho, **params):
        """GET na API por dentro do navegador. Devolve (status, texto)."""
        url = CONFIG["api"] + caminho
        if params:
            url += "?" + urllib.parse.urlencode({k: v for k, v in params.items() if v not in (None, "")})
        try:
            r = self.pagina.evaluate(_JS_FETCH, [url, CONFIG["timeout"] * 1000])
            return r["status"], r["texto"]
        except Exception as e:
            return 0, str(e)

    def esfera(self, id_ato):
        """Esfera do ato (FEDERAL/ESTADUAL/CGIBS), lida na ficha do ato no site."""
        if not hasattr(self, "_esferas"):
            self._esferas = {}
        if id_ato not in self._esferas:
            st, txt = self.get(f"/atos-legislativos/{id_ato}")
            if st != 200:
                raise ConnectionError(f"status {st} ao ler o ato {id_ato}")
            try:
                nome = ((json.loads(txt).get("categoria") or {}).get("esfera") or {}).get("nome", "")
            except ValueError:
                nome = ""
            self._esferas[id_ato] = normalizar(nome)
        return self._esferas[id_ato]

    def abrir_externo(self, url):
        """Abre uma página de outro site (ex.: Planalto) no navegador. Devolve (status, html)."""
        try:
            if getattr(self, "_pag_ext", None) is None:
                self._pag_ext = self.pagina.context.new_page()
            r = self._pag_ext.goto(url, wait_until="domcontentloaded", timeout=CONFIG["timeout"] * 1000)
            return (r.status if r else 0), self._pag_ext.content()
        except Exception:
            return 0, ""

    def _json(self, caminho):
        st, txt = self.get(caminho)
        if st != 200:
            raise ConnectionError(f"status {st} em {caminho}")
        try:
            return json.loads(txt)
        except ValueError:
            raise ConnectionError(f"resposta inesperada em {caminho}")

    def id_categoria(self, nome, esfera):
        """Id da categoria (ex.: CONSTITUICAO + FEDERAL -> 1)."""
        for c in self._json("/categorias"):
            if (normalizar(c.get("nome", "")) == nome
                    and normalizar((c.get("esfera") or {}).get("nome", "")) == esfera):
                return c.get("id")
        return None

    def id_subcategoria(self, id_cat, nome):
        """Id da subcategoria dentro da categoria (ex.: TEXTO INTEGRAL). None se não achar."""
        for sc in self._json(f"/categorias/{id_cat}/subcategorias"):
            rotulo = sc.get("nome") or sc.get("descricao") or sc.get("nomeSubCategoria") or ""
            if normalizar(rotulo) == nome:
                return sc.get("id")
        return None

    def pesquisar(self, numero, ano):
        """Lista de atos (dicts da API) com esse número (e ano, se informado)."""
        st, txt = self.get("/atos-legislativos", page=0, size=50, idStatus=3, numeroAto=numero,
                           ano=ano, ordenacao="MAIS_RECENTES")
        if st != 200:
            raise ConnectionError(f"status {st}: {txt[:120]}")
        try:
            return json.loads(txt).get("content", [])
        except ValueError:
            raise ConnectionError("resposta inesperada do site")


NOMES_NO_SITE = {
    "Instrução Normativa": r"INSTRUCAO +NORMATIVA|IN",
    "Lei Complementar": r"LEI +COMPLEMENTAR|LC",
}


def _re_ato(tipo, numero):
    """Reconhece, na epígrafe do site, o tipo (sem confundir Lei com Lei Complementar,
    Decreto com Decreto-Lei etc.) seguido logo adiante do número."""
    nomes = NOMES_NO_SITE.get(tipo, re.escape(normalizar(tipo)).replace(r"\ ", " +"))
    return re.compile(r"(^|[^A-Z-])(" + nomes + r")(?![ -]*(COMPLEMENTAR|LEI\b|LEGISLATIVO|CONJUNTA))"
                      r"[^0-9]{0,30}?0*" + re.escape(numero) + r"(?![0-9])")


def _ano_da_epigrafe(epi):
    m = re.search(r"\bDE (\d{4})\b", epi) or re.search(r"/(\d{4})\b", epi)
    return m.group(1) if m else ""


def candidatos(resultados, tipo, numero, ano):
    """Atos da resposta da API que batem com tipo + número (+ ano). Devolve (lista, ano_confirmado)."""
    r_ato = _re_ato(tipo, numero)
    cand, vistos = [], set()
    for a in resultados:
        epi = re.sub(r"(\d)\.(?=\d)", r"\1", normalizar(a.get("epigrafe") or ""))
        if r_ato.search(epi) and a.get("idAto") not in vistos:
            vistos.add(a.get("idAto"))
            cand.append((a, _ano_da_epigrafe(epi) or (a.get("dataPublicacao") or "")[:4]))
    if not ano:
        return [c[0] for c in cand], True
    com_ano = [c[0] for c in cand if c[1] == ano]
    if com_ano:
        return com_ano, True
    # sem ato do ano citado: só serve o único candidato cujo ano o site não informa; um ato de
    # outro ano é outro ato (ATC nº 6/2026: "Portaria ME nº 284, de 2020" ia para outra Portaria nº 284)
    return ([cand[0][0]] if len(cand) == 1 and not cand[0][1] else []), False


def escolher_ato(site, cands, esfera, preferida="ESTADUAL"):
    """Escolhe o ato considerando a esfera. esfera: "FEDERAL"/"ESTADUAL"/"CGIBS" se a citação
    disser, "" se não disser (aí, havendo empate, prefere a esfera do documento). Devolve (ato|None, situação)."""
    if not cands:
        return None, S_NAO_ACHOU
    if esfera:
        cands = [a for a in cands if site.esfera(a["idAto"]) == esfera]
        if not cands:
            return None, f"Não encontrado no site (esfera {ESFERAS.get(esfera, esfera)})"
    elif len(cands) > 1:
        da_preferida = [a for a in cands if site.esfera(a["idAto"]) == (preferida or "ESTADUAL")]
        if len(da_preferida) == 1:
            cands = da_preferida
    if len(cands) > 1:
        esf = ", ".join(sorted({ESFERAS.get(site.esfera(a["idAto"]), "?") for a in cands}))
        return None, f"Ambíguo: {len(cands)} atos no site ({esf}) - incluir na tabela local"
    return cands[0], S_SITE


def montar_link(id_ato, **filtros):
    """Mesmo formato de endereço que o site gera ao abrir um ato."""
    params = {"idStatus": 3, **{k: v for k, v in filtros.items() if v not in (None, "")},
              "ordenacao": "MAIS_RECENTES", "idAto": id_ato}
    return CONFIG["pagina_ato"] + "?" + urllib.parse.urlencode(params)


def buscar_constituicao(site, esfera):
    """Constituição Federal/Estadual: categoria Constituição + subcategoria Texto Integral."""
    cat = site.id_categoria("CONSTITUICAO", esfera)
    if cat is None:
        return "", f"Categoria Constituição ({ESFERAS[esfera]}) não encontrada no site"
    sub = site.id_subcategoria(cat, "TEXTO INTEGRAL")
    st, txt = site.get("/atos-legislativos", page=0, size=50, idStatus=3, idCategoria=cat,
                       idSubCategoria=sub, ordenacao="MAIS_RECENTES")
    if st != 200:
        raise ConnectionError(f"status {st}")
    ano_const = "1988" if esfera == "FEDERAL" else "1989"
    cands = [a for a in json.loads(txt).get("content", [])
             if normalizar(a.get("epigrafe") or "").startswith("CONSTITUICAO")]
    if len(cands) > 1:
        com_ano = [a for a in cands if ano_const in (a.get("epigrafe") or "")]
        cands = com_ano or cands
    cands = [a for a in cands if site.esfera(a["idAto"]) == esfera]
    if not cands:
        return "", S_NAO_ACHOU
    if len(cands) > 1:
        return "", f"Ambíguo: {len(cands)} resultados - incluir na tabela local"
    return montar_link(cands[0]["idAto"], ano=ano_const), S_SITE


def buscar_no_site(site, o, preferida="ESTADUAL"):
    """Devolve (url, situação). Levanta ConnectionError se o site falhar."""
    if o.tipo == "Constituição":
        return buscar_constituicao(site, o.esfera)
    tentativas = []
    for n in (o.numero, o.numero_orig, o.numero.zfill(4)):
        if n and n not in tentativas:
            tentativas.append(n)
    sit, outro_ano = S_NAO_ACHOU, False
    for n in tentativas:
        for a in ((o.ano, "") if o.ano else ("",)):
            resultados = site.pesquisar(n, a)
            cands, ano_ok = candidatos(resultados, o.tipo, o.numero, o.ano)
            if not cands and o.ano and candidatos(resultados, o.tipo, o.numero, "")[0]:
                outro_ano = True
            ato, sit = escolher_ato(site, cands, o.esfera, preferida)
            if ato:
                if not ano_ok:
                    sit = S_SITE_SEM_ANO
                return montar_link(ato["idAto"], numeroAto=n, ano=a), sit
            if sit.startswith("Ambíguo"):
                return "", sit
            time.sleep(CONFIG["pausa"])
    return "", (S_OUTRO_ANO if outro_ano and sit == S_NAO_ACHOU else sit)


# -----------------------------------------------------------------------------
# PLANALTO (atos federais que não estão no site da SEFA)
# -----------------------------------------------------------------------------
PLANALTO = "https://www.planalto.gov.br/ccivil_03/"
TIPOS_PLANALTO = {"Lei", "Lei Complementar", "Decreto", "Decreto-Lei", "Emenda Constitucional", "Constituição"}
SEMPRE_FEDERAIS = {"Decreto-Lei"}
S_PLANALTO = "Link criado (Planalto)"


def _era(ano):
    for a, b in ((2004, 2006), (2007, 2010), (2011, 2014), (2015, 2018), (2019, 2022), (2023, 2026),
                 (2027, 2030), (2031, 2034)):
        if a <= ano <= b:
            return f"_ato{a}-{b}/{ano}"
    return ""


def planalto_candidatos(tipo, numero, ano):
    """Endereços possíveis no Planalto (o formato muda conforme a época)."""
    if tipo == "Constituição":
        return ["constituicao/constituicao.htm"]
    if not numero.isdigit():
        return []
    n = int(numero)
    ponto = f"{n:,}".replace(",", ".")
    y = int(ano) if ano else 0
    c = []
    if tipo == "Lei Complementar":
        c = [f"leis/lcp/lcp{n}.htm"]
    elif tipo == "Emenda Constitucional":
        c = [f"constituicao/emendas/emc/emc{n}.htm"]
    elif tipo == "Decreto-Lei":
        c = [f"decreto-lei/del{n:04d}.htm", f"decreto-lei/del{n}.htm"]
    elif tipo == "Lei":
        if y >= 2004:
            c = [f"{_era(y)}/lei/l{n}.htm", f"{_era(y)}/lei/l{ponto}.htm"]
        elif y == 2003:
            c = [f"leis/2003/l{ponto}.htm", f"leis/2003/l{n}.htm"]
        elif y == 2002:
            c = [f"leis/2002/l{n}.htm"]
        elif y == 2001:
            c = [f"leis/leis_2001/l{n}.htm"]
        else:
            c = [f"leis/l{n}.htm"] + ([f"leis/l{n:04d}.htm"] if n < 1000 else [])
    elif tipo == "Decreto":
        if y >= 2004:
            c = [f"{_era(y)}/decreto/d{n}.htm"]
        elif y in (2002, 2003):
            c = [f"decreto/{y}/d{n}.htm"]
        else:
            c = [f"decreto/d{n}.htm", f"decreto/1990-1994/d{n}.htm", f"decreto/antigos/d{n}.htm"]
    return [PLANALTO + x for x in c]


def _baixar_direto(url, timeout=20):
    req = urllib.request.Request(url, headers={"User-Agent": _UA_NAVEGADOR})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            bruto = r.read()
            cs = r.headers.get_content_charset()
            if not cs:
                m = re.search(rb"charset=[\"']?([\w-]+)", bruto[:3000], re.I)
                cs = m.group(1).decode() if m else "cp1252"
            try:
                return r.status, bruto.decode(cs, errors="replace")
            except LookupError:
                return r.status, bruto.decode("cp1252", errors="replace")
    except urllib.error.HTTPError as e:
        return e.code, ""
    except Exception:
        return 0, ""


def _titulo_confere(html_txt, tipo, numero, ano):
    """O título da página (início do texto) é o ato procurado?"""
    texto = normalizar(re.sub(r"<[^>]+>", " ", re.sub(r"(?is)<(script|style|head)\b.*?</\1>", " ", html_txt)))
    texto = re.sub(r"(\d)\.(?=\d)", r"\1", texto)[:1500]
    if tipo == "Constituição":
        return "CONSTITUICAO DA REPUBLICA FEDERATIVA DO BRASIL" in texto
    if not _re_ato(tipo, numero).search(texto):
        return False
    return (ano in texto) if ano else True


def buscar_planalto(o, site=None):
    """Procura o ato no Planalto. Devolve (url, situação)."""
    for url in planalto_candidatos(o.tipo, o.numero, o.ano):
        st, txt = _baixar_direto(url)
        if st in (0, 403) and site is not None:      # bloqueado para programas: tenta pelo navegador
            st, txt = site.abrir_externo(url)
        if st == 200 and _titulo_confere(txt, o.tipo, o.numero, o.ano):
            return url, S_PLANALTO
        time.sleep(CONFIG["pausa"])
    return "", "Não encontrado no site da SEFA nem no Planalto"


def http(metodo, url, corpo=None, timeout=60):
    """Acesso direto (sem navegador) - usado só no diagnóstico."""
    req = urllib.request.Request(url, data=corpo.encode() if corpo else None, method=metodo,
                                 headers={"User-Agent": _UA_NAVEGADOR})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, r.read().decode("utf-8", "replace"), r.geturl(), dict(r.headers)
    except urllib.error.HTTPError as e:
        return e.code, "", url, dict(e.headers or {})
    except Exception as e:
        return 0, f"{type(e).__name__}: {e}", url, {}


# =============================================================================
# PROCESSAMENTO DO DOCUMENTO
# =============================================================================
class Ocorrencia:
    __slots__ = ("p", "ini", "fim", "tipo", "numero", "numero_orig", "ano", "texto", "esfera",
                 "chave", "npar")


RE_ENDERECO = re.compile(r"(?<![\w@/])(?:https?://|www\.)[^\s<>\"“”]+", re.I)


def linkar_enderecos(doc, log=print):
    """Endereços de internet escritos no texto ("https://www.gov.br/sped") viram links para eles mesmos;
    o ponto, a vírgula ou o parêntese que fecham a frase ficam de fora. Devolve quantos."""
    estilo_id = estilo_hyperlink(doc)
    estilos = Estilos(doc, estilo_id)
    total, rids = 0, {}
    for p in doc.element.body.iter(qn("w:p")):
        normalizar_runs(p)
        txt = texto_paragrafo(p)
        achados = []
        for m in RE_ENDERECO.finditer(txt):
            fim = m.end()
            while fim > m.start() and txt[fim - 1] in ".,;:)!?»”'\"":
                fim -= 1
            if fim - m.start() > 8:
                achados.append((m.start(), fim, txt[m.start():fim]))
        for ini, fim, url in reversed(achados):
            nos = isolar_runs(p, ini, fim)
            if not nos:
                continue
            alvo = url if url.lower().startswith("http") else "https://" + url
            if alvo not in rids:
                rids[alvo] = doc.part.relate_to(alvo, RT.HYPERLINK, is_external=True)
            criar_hyperlink(p, nos, rids[alvo], estilo_id, estilos)
            total += 1
    if total:
        log(f"Endereços de internet com link: {total}")
    return total


class ResultadoLinks:
    """Resumo de uma passagem sobre o documento (para relatórios)."""

    def __init__(self, atos=None, linkados=0, pendentes=0, destacadas=0, novos=0, enderecos=0):
        self.atos = atos or {}
        self.linkados, self.pendentes, self.destacadas, self.novos = linkados, pendentes, destacadas, novos
        self.enderecos = enderecos

    @property
    def linhas(self):
        """[(ato na 1ª citação, ocorrências, situação, link)] na ordem em que aparecem."""
        return [(a["exemplo"], a["qtd"], a["sit"], a["url"]) for a in self.atos.values()]


def _interrompido(parar):
    return parar is not None and parar.is_set()


def criar_links(doc, log=print, parar=None):
    """Cria os links no documento já aberto (python-docx) e NÃO salva: quem chama decide onde
    gravar. parar: threading.Event opcional; quando acionado, as buscas restantes no site são
    puladas e o documento fica só com o que já foi resolvido. Devolve ResultadoLinks."""
    enderecos = linkar_enderecos(doc, log)
    corpo = doc.element.body
    paragrafos = list(corpo.iter(qn("w:p")))

    for p in paragrafos:
        normalizar_runs(p)
    textos = [texto_paragrafo(p) for p in paragrafos]
    esf_doc = esfera_documento([t for t in textos if t.strip(BARREIRA + " \t")])
    if esf_doc:
        log(f"Documento identificado como {ESFERAS[esf_doc].lower()}.")

    def pode_planalto(o):
        """Fonte principal é o site da SEFA; o Planalto só entra para o que lá não está.
        Vale para citações federais ou sem esfera dita (as que dizem Estadual/CGIBS, não).
        O título da página no Planalto precisa conferir número e ano, o que evita trocar
        uma lei estadual por uma federal de mesmo número."""
        if o.tipo not in TIPOS_PLANALTO:
            return False
        if o.tipo == "Constituição":
            return o.esfera == "FEDERAL"
        return o.esfera in ("", "FEDERAL") or o.tipo in SEMPRE_FEDERAIS

    ementa = set() if CONFIG["linkar_ementa"] else paragrafos_de_ementa(doc, paragrafos, textos)
    if ementa:
        log("Ementa identificada: não recebe links.")

    # 1) localizar citações
    ocorrs, anos_por_ato, nao_vazios = [], {}, 0

    def nova(p, ini, fim, tipo, num_txt, ano, texto, esfera):
        o = Ocorrencia()
        o.p, o.npar, o.ini, o.fim = p, nao_vazios, ini, fim
        o.tipo = tipo
        o.numero = norm_numero(num_txt) if num_txt else ""
        o.numero_orig = num_txt.replace(".", "") if num_txt else ""
        o.ano = ano
        o.texto = limpar_espacos(texto)
        o.esfera = esfera
        ocorrs.append(o)
        if o.ano:
            anos_por_ato.setdefault((o.tipo, o.numero), set()).add(o.ano)
        return o

    def eh_epigrafe(tipo_txt):
        return (CONFIG["ignorar_epigrafe"] and nao_vazios <= CONFIG["linhas_epigrafe"]
                and len(tipo_txt) > 2 and tipo_txt == tipo_txt.upper())

    for p, txt in zip(paragrafos, textos):
        if txt.strip(BARREIRA + " \t"):
            nao_vazios += 1
        if id(p) in ementa:
            continue
        ocupados = []

        def livre(a, b):
            if RE_PROJETO.search(txt[max(0, a - 25):a]):     # "Projeto de Lei nº ..." não é lei
                return False
            return not any(x < b and a < y for x, y in ocupados)

        # listas no plural: cada número vira uma citação
        for m in RE_PLURAL.finditer(txt):
            if eh_epigrafe(m.group("tipo")) or not livre(m.start(), m.end()) or not item_valido(m, txt):
                continue
            tipo = PLURAIS.get(limpar_espacos(m.group("tipo")),
                               next((v for k, v in PLURAIS.items() if k.upper() == limpar_espacos(m.group("tipo"))),
                                    ""))
            esf = esfera_citada(m.group("cab"))
            fim = m.end() if CONFIG["incluir_data_no_link"] else m.end("fimnucleo")
            ano = ano_completo(m.group("ano2") or m.group("ano4") or m.group("ano5") or "")
            nova(p, m.start(), fim, tipo, m.group("num"), ano, txt[m.start():fim], esf)
            ocupados.append((m.start(), m.end()))
            pos = m.end()
            while True:
                s2 = RE_PLURAL_SEGUINTE.match(txt, pos)
                if not s2 or not item_valido(s2, txt):
                    break
                ini = s2.start("num")
                fim = s2.end() if CONFIG["incluir_data_no_link"] else s2.end("fimnucleo")
                ano = ano_completo(s2.group("ano2") or s2.group("ano4") or s2.group("ano5") or "")
                nova(p, ini, fim, tipo, s2.group("num"), ano, txt[ini:fim], esf)
                ocupados.append((ini, s2.end()))
                pos = s2.end()

        for m in RE_CITACAO.finditer(txt):
            if not livre(m.start(), m.end()):
                continue
            tipo_txt = m.group("tipo")
            if eh_epigrafe(tipo_txt):
                continue
            fim = m.end() if CONFIG["incluir_data_no_link"] else m.end("nucleo")
            nucleo = txt[m.start("nucleo"):m.start("orgao")] if m.group("orgao") else m.group("nucleo")
            nova(p, m.start(), fim, canon_tipo(tipo_txt), m.group("num"),
                 ano_completo(m.group("ano2") or m.group("ano4") or m.group("ano5") or ""),
                 m.group(0), esfera_citada(nucleo))
            ocupados.append((m.start(), m.end()))

        for m in RE_CONSTITUICAO.finditer(txt):
            if livre(m.start(), m.end()):
                nova(p, m.start(), m.end(), "Constituição", "", "", m.group(0),
                     "FEDERAL" if m.group("fed") else "ESTADUAL")
                ocupados.append((m.start(), m.end()))
        for m in RE_CONST_SOLTA.finditer(txt):
            if not livre(m.start(), m.end()):
                continue
            if m.group("adct") or m.group("c88"):
                esf = "FEDERAL"          # ADCT sem qualificação = o da Constituição Federal
            else:
                esf = esf_doc if esf_doc in ("FEDERAL", "ESTADUAL") else ""
            if esf:
                nova(p, m.start(), m.end(), "Constituição", "", "", m.group(0), esf)
                ocupados.append((m.start(), m.end()))

    if not ocorrs:
        log("Nenhuma citação de ato sem link foi encontrada.")
        return ResultadoLinks(enderecos=enderecos)

    # 2) completar o ano das citações abreviadas ("Lei nº 5.530") e a esfera das citações
    #    que não a dizem, quando o mesmo ato aparece completo em outro ponto do texto
    for o in ocorrs:
        if not o.ano and len(anos_por_ato.get((o.tipo, o.numero), ())) == 1:
            o.ano = next(iter(anos_por_ato[(o.tipo, o.numero)]))
    esferas_por_ato = {}
    for o in ocorrs:
        if o.esfera:
            esferas_por_ato.setdefault((o.tipo, o.numero, o.ano), set()).add(o.esfera)
    for o in ocorrs:
        conhecidas = esferas_por_ato.get((o.tipo, o.numero, o.ano), set())
        if not o.esfera and len(conhecidas) == 1:
            o.esfera = next(iter(conhecidas))
        o.chave = f"{o.tipo}|{o.numero}|{o.ano}" + (f"|{ESFERAS[o.esfera]}" if o.esfera else "")

    atos = {}
    for o in ocorrs:
        a = atos.setdefault(o.chave, {"exemplo": o.texto, "qtd": 0, "o": o, "url": "", "sit": ""})
        a["qtd"] += 1
    log(f"{len(ocorrs)} citações de {len(atos)} atos diferentes.")

    # 3) resolver links: tabela local -> site
    tabela = carregar_tabela()
    novos = 0
    a_buscar = []
    for chave, a in atos.items():
        o = a["o"]
        if chave in tabela:
            a["url"], a["sit"] = tabela[chave], S_TABELA
        elif o.tipo not in TIPOS_NO_SITE and not pode_planalto(o):
            a["sit"] = S_FORA_SITE
        else:
            a_buscar.append(a)

    if a_buscar and _interrompido(parar):
        for a in a_buscar:
            a["sit"] = S_INTERROMPIDO
    elif a_buscar:
        try:
            with SiteSEFA(log) as site:
                falhas = 0
                for i, a in enumerate(a_buscar, 1):
                    o = a["o"]
                    if _interrompido(parar):
                        if not a["sit"]:
                            a["sit"] = S_INTERROMPIDO
                        continue
                    if falhas >= CONFIG["max_falhas_seguidas"]:
                        a["sit"] = S_SITE_OFF
                        continue
                    log(f"  [{i}/{len(a_buscar)}] {a['exemplo']}")
                    try:
                        if o.tipo in TIPOS_NO_SITE:
                            a["url"], a["sit"] = buscar_no_site(site, o, esf_doc)
                        else:
                            a["sit"] = S_NAO_ACHOU
                        falhas = 0
                    except ConnectionError as e:
                        falhas += 1
                        a["sit"] = f"Erro ao consultar o site ({e})"
                        if falhas >= CONFIG["max_falhas_seguidas"]:
                            log("  ! O site falhou várias vezes seguidas; as demais buscas foram puladas.")
                    if not a["url"] and a["sit"].startswith("Não encontrado") and pode_planalto(o):
                        a["url"], a["sit"] = buscar_planalto(o, site)
                    if a["url"] and a["sit"] in (S_SITE, S_PLANALTO):
                        tabela[a["o"].chave] = a["url"]
                        novos += 1
                    log(f"      -> {a['sit']}")
                    time.sleep(CONFIG["pausa"])
                if _interrompido(parar):
                    log("  Buscas interrompidas: os atos restantes ficaram sem link.")
        except ImportError:
            log("  ! Falta o componente Playwright (no Formatador SEFA, rode de novo o INSTALAR.bat).")
            for a in a_buscar:
                a["sit"] = S_SITE_OFF
        except Exception as e:
            log(f"  ! Não foi possível usar o site: {e}")
            for a in a_buscar:
                if not a["sit"]:
                    a["sit"] = S_SITE_OFF
        # sem o site da SEFA, ainda dá para buscar os federais no Planalto
        for a in a_buscar:
            if a["sit"] == S_SITE_OFF and pode_planalto(a["o"]) and not _interrompido(parar):
                url, sit = buscar_planalto(a["o"])
                if url:      # sem achar, mantém "não buscado": o site da SEFA nem foi consultado
                    a["url"], a["sit"] = url, sit
                    tabela[a["o"].chave] = url
                    novos += 1
    if novos:
        salvar_tabela(tabela)

    # 4) aplicar (de trás para frente dentro de cada parágrafo)
    estilo_id = estilo_hyperlink(doc)
    estilos = Estilos(doc, estilo_id)
    rids, linkados, pendentes, destacadas = {}, 0, 0, 0
    for o in sorted(ocorrs, key=lambda x: (id(x.p), -x.ini)):
        a = atos[o.chave]
        nos = isolar_runs(o.p, o.ini, o.fim)
        if not nos:
            continue
        if a["url"]:
            if a["url"] not in rids:
                rids[a["url"]] = doc.part.relate_to(a["url"], RT.HYPERLINK, is_external=True)
            criar_hyperlink(o.p, nos, rids[a["url"]], estilo_id, estilos)
            linkados += 1
        else:
            if CONFIG["destacar_pendentes"] and a["sit"] not in SEM_DESTAQUE:
                destacar(nos)
                destacadas += 1
            pendentes += 1

    log("")
    log(f"Links criados: {linkados}   |   Sem link: {pendentes}"
        + (f" ({destacadas} destacadas em amarelo)" if destacadas else ""))
    if novos:
        log(f"Atos novos guardados na tabela local: {novos}")
    return ResultadoLinks(atos, linkados, pendentes, destacadas, novos, enderecos)


def processar(caminho_docx, log=print, parar=None):
    """Uso avulso: grava "... - com links.docx" e o relatório em Excel ao lado do original."""
    caminho = Path(caminho_docx)
    if caminho.suffix.lower() != ".docx":
        raise ValueError("O arquivo precisa ser .docx (no Word: Salvar como > Documento do Word).")
    log(f"Abrindo {caminho.name}...")
    doc = Document(str(caminho))
    res = criar_links(doc, log, parar)
    if not res.atos and not res.enderecos:
        return None
    saida = caminho.with_name(caminho.stem + " - com links.docx")
    doc.save(str(saida))
    rel = gerar_relatorio(caminho, res.atos)
    log(f"Documento: {saida}")
    log(f"Relatório: {rel}")
    return saida, rel


def gerar_relatorio(caminho, atos):
    linhas = [(a["exemplo"], a["qtd"], a["sit"], a["url"]) for a in atos.values()]
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Links"
        ws.append(["Ato (1ª citação)", "Ocorrências", "Situação", "Link"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for l in linhas:
            ws.append(list(l))
            if l[3]:
                ws.cell(ws.max_row, 4).hyperlink = l[3]
        for col, larg in zip("ABCD", (55, 12, 55, 70)):
            ws.column_dimensions[col].width = larg
        ws.freeze_panes = "A2"
        destino = caminho.with_name(caminho.stem + " - relatório de links.xlsx")
        wb.save(str(destino))
    except ImportError:
        destino = caminho.with_name(caminho.stem + " - relatório de links.csv")
        with open(destino, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Ato (1ª citação)", "Ocorrências", "Situação", "Link"])
            w.writerows(linhas)
    return destino


# =============================================================================
# DIAGNÓSTICO DO SITE
# =============================================================================
def diagnosticar(log=print, concluir=None):
    """
    Parte A: tenta o site por acesso direto (sem navegador), com espera longa.
    Parte B: abre o site no Edge/Chrome e grava tudo o que o site troca com o navegador
             enquanto você faz uma busca de exemplo. É isso que mostra como a busca funciona.
    concluir: threading.Event acionado pelo botão "Concluir diagnóstico" (na janela).
    """
    pasta = PASTA_PROGRAMA / "diagnostico_sefa"
    pasta.mkdir(exist_ok=True)
    for antigo in pasta.iterdir():
        if antigo.is_file():
            antigo.unlink()
    info = [f"Data: {datetime.now():%d/%m/%Y %H:%M}", f"Python: {sys.version.split()[0]}",
            f"truststore: {'sim' if 'truststore' in sys.modules else 'não'}", ""]
    base = CONFIG["pagina_busca"]

    # ---------------- A) acesso direto, esperando até 60 s ----------------
    log("Parte A: acesso direto ao site (pode levar até 1 minuto)...")
    t0 = time.time()
    st, txt, final, cab = http("GET", base)
    info.append(f"A) GET {base}")
    info.append(f"   -> status {st} em {time.time() - t0:.1f} s | final: {final} | {len(txt)} caracteres")
    if st == 0:
        info.append(f"   ERRO: {txt}")
    info.append("   cabeçalhos: " + " | ".join(f"{k}: {v}" for k, v in cab.items()))
    (pasta / "01_acesso_direto.html").write_text(txt, encoding="utf-8")
    log(f"   -> {'OK' if st == 200 else 'falhou'} (status {st}, {time.time() - t0:.0f} s)")
    info.append("")

    # ---------------- B) navegador de verdade ----------------
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        info.append("B) Playwright não instalado.")
        log("")
        log("! Falta instalar o Playwright: rode de novo o '1 - Instalar dependencias.bat'.")
        _fechar_diagnostico(pasta, info, [], log)
        return

    eventos = []
    fotos = {"n": 0}

    def foto(pg, rotulo):
        try:
            fotos["n"] += 1
            nome = f"{fotos['n']:02d}_{rotulo}"
            (pasta / f"{nome}.html").write_text(pg.content(), encoding="utf-8")
            pg.screenshot(path=str(pasta / f"{nome}.png"), full_page=True)
            info.append(f"   foto {nome}: {pg.url}")
        except Exception as e:
            info.append(f"   (não foi possível fotografar a página: {e})")

    log("")
    log("Parte B: abrindo o site no navegador...")
    with sync_playwright() as pw:
        navegador = None
        for canal in ("msedge", "chrome"):
            try:
                navegador = pw.chromium.launch(channel=canal, headless=False)
                info.append(f"B) navegador: {canal}")
                break
            except Exception as e:
                info.append(f"B) não abriu {canal}: {str(e)[:200]}")
        if navegador is None:
            log("! Não consegui abrir o Edge nem o Chrome.")
            _fechar_diagnostico(pasta, info, eventos, log)
            return

        ctx = navegador.new_context(record_har_path=str(pasta / "rede.har"),
                                    record_har_url_filter=re.compile(r"sefa\.pa\.gov\.br"),
                                    record_har_content="embed")
        t_ini = time.time()

        def marca():
            return f"{time.time() - t_ini:7.1f}s"
        def ao_pedir(r):
            corpo = ""
            if r.method != "GET":
                try:
                    corpo = f"\n{' ' * 18}corpo: {(r.post_data or '')[:3000]}"
                except Exception:
                    corpo = f"\n{' ' * 18}corpo: (binário)"
            eventos.append(f"{marca()} PEDIDO  {r.method:5} {r.resource_type:10} {r.url}{corpo}")
        ctx.on("request", ao_pedir)
        ctx.on("response", lambda r: eventos.append(
            f"{marca()} RESPOSTA {r.status} {r.url}  [{r.headers.get('content-type', '')}]"))
        ctx.on("requestfailed", lambda r: eventos.append(f"{marca()} FALHOU  {r.url}  {r.failure}"))

        pagina = ctx.new_page()
        t0 = time.time()
        try:
            pagina.goto(base, timeout=120000, wait_until="load")
            info.append(f"   página carregou em {time.time() - t0:.1f} s")
            log(f"   -> o site abriu no navegador em {time.time() - t0:.0f} s")
        except Exception as e:
            info.append(f"   página NÃO carregou em {time.time() - t0:.1f} s: {str(e)[:300]}")
            log(f"   -> o site não carregou no navegador ({time.time() - t0:.0f} s)")
        foto(pagina, "navegador_inicio")

        log("")
        log(">>> AGORA, NA JANELA DO NAVEGADOR QUE ABRIU:")
        log("    1. Pesquise a Lei 5.530 (de 1989), como você faria normalmente.")
        log("    2. Clique no resultado para abrir a lei.")
        log("    3. Volte aqui e clique em \"Concluir diagnóstico\" (NÃO feche o navegador).")
        if concluir is None:
            input("\nQuando terminar, pressione Enter aqui... ")

        vistos = set()
        limite = time.time() + 15 * 60
        while time.time() < limite and (concluir is None or not concluir.is_set()):
            if concluir is None:
                break
            try:
                if not navegador.is_connected() or not ctx.pages:
                    info.append("   (o navegador foi fechado antes de concluir)")
                    break
                for pg in list(ctx.pages):
                    chave = (id(pg), pg.url)
                    if chave not in vistos:
                        vistos.add(chave)
                        try:
                            pg.wait_for_load_state("load", timeout=30000)
                        except Exception:
                            pass
                        if len(vistos) > 1:
                            foto(pg, "navegacao")
                ctx.pages[0].wait_for_timeout(1000)
            except Exception:
                time.sleep(1)

        # foto final de todas as abas abertas (resultado da busca, lei aberta etc.)
        try:
            for pg in list(ctx.pages):
                foto(pg, "final")
        except Exception:
            pass
        try:
            ctx.close()          # grava o rede.har
        except Exception as e:
            info.append(f"   (não gravou rede.har: {e})")
        try:
            navegador.close()
        except Exception:
            pass

    _fechar_diagnostico(pasta, info, eventos, log)


def _fechar_diagnostico(pasta, info, eventos, log):
    (pasta / "00_info.txt").write_text("\n".join(info), encoding="utf-8")
    if eventos:
        (pasta / "rede_resumo.txt").write_text("\n".join(eventos), encoding="utf-8")
    zip_path = PASTA_PROGRAMA / "diagnostico_sefa.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for arq in sorted(pasta.iterdir()):
            z.write(arq, arq.name)
    log("")
    log(f"Diagnóstico pronto: {zip_path}")
    log("Envie esse arquivo .zip no chat.")
    return zip_path


# =============================================================================
# JANELA
# =============================================================================
def janela():
    import tkinter as tk
    from tkinter import filedialog, scrolledtext

    raiz = tk.Tk()
    raiz.title("Links de Legislação - SEFA/PA")
    raiz.geometry("760x480")
    fila = queue.Queue()

    topo = tk.Frame(raiz, padx=10, pady=8)
    topo.pack(fill="x")
    caixa = scrolledtext.ScrolledText(raiz, font=("Consolas", 9), state="disabled")
    caixa.pack(fill="both", expand=True, padx=10, pady=(0, 10))

    def log(msg):
        fila.put(msg)

    def bombear():
        while not fila.empty():
            caixa.configure(state="normal")
            caixa.insert("end", fila.get() + "\n")
            caixa.see("end")
            caixa.configure(state="disabled")
        raiz.after(100, bombear)

    botoes = []

    def rodar(func, *args):
        for b in botoes:
            b.configure(state="disabled")

        def trabalho():
            try:
                func(*args, log=log)
            except Exception as e:
                log(f"\nERRO: {e}")
                log(traceback.format_exc())
            finally:
                raiz.after(0, lambda: [b.configure(state="normal") for b in botoes])
        threading.Thread(target=trabalho, daemon=True).start()

    def escolher():
        arqs = filedialog.askopenfilenames(title="Escolha o(s) documento(s)",
                                           filetypes=[("Documento do Word", "*.docx")])
        if arqs:
            rodar(lambda lista, log: [processar(a, log=log) for a in lista], list(arqs))

    def abrir_tabela():
        if not ARQ_TABELA.exists():
            salvar_tabela({})
        import os
        os.startfile(str(ARQ_TABELA)) if hasattr(os, "startfile") else log(str(ARQ_TABELA))

    evento_concluir = threading.Event()

    def diagnostico():
        evento_concluir.clear()
        b_concluir.configure(state="normal")

        def diag(log):
            try:
                diagnosticar(log=log, concluir=evento_concluir)
            finally:
                raiz.after(0, lambda: b_concluir.configure(state="disabled"))
        rodar(diag)

    for texto, cmd in (("Criar links em documento(s)...", escolher),
                       ("Diagnóstico do site", diagnostico),
                       ("Abrir tabela de links", abrir_tabela)):
        b = tk.Button(topo, text=texto, command=cmd, padx=10, pady=4)
        b.pack(side="left", padx=(0, 8))
        botoes.append(b)
    b_concluir = tk.Button(topo, text="Concluir diagnóstico", command=evento_concluir.set,
                           padx=10, pady=4, state="disabled", bg="#ffe08a")
    b_concluir.pack(side="left")

    log("Escolha o .docx já formatado. O original não é alterado: é criada uma cópia")
    log("\"... - com links.docx\" e um relatório em Excel na mesma pasta.\n")
    bombear()
    raiz.mainloop()


if __name__ == "__main__":
    try:
        if len(sys.argv) > 1 and sys.argv[1] == "--diagnostico":
            diagnosticar()
        elif len(sys.argv) > 1:
            for arq in sys.argv[1:]:
                processar(arq)
        else:
            janela()
    except Exception:
        _aviso("Links de Legislação - erro", "O programa encontrou um erro:\n\n" + traceback.format_exc()[-2000:],
               erro=True)
