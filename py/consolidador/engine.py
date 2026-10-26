from __future__ import annotations
from copy import deepcopy
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from io import BytesIO
import json
import os
import re
import tempfile
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree
from .model import Operation, Plan, Target, date_iso, digest, norm
from .index import Legislation, label
from .io import prepare_base, read_amender
from .parser import parse_operations, source_citation, number_value
from .temporal import publication_date, assign_effects, collect_remissions
from .editorial import origin_note, note_for_targets, replace_text, links_signature, update_initial_reference, save_preserving_package
from . import notes as notes_mod
from . import refs as refparser

ROOT = Path(__file__).resolve().parent.parent


def comparable_lines(lines):
    # Acentos e maiúsculas fazem parte da redação. Só uniformiza espaços.
    return [" ".join(s.split()) for s in lines]


def sort_key(target):
    result = []
    for level, part in enumerate(target.parts):
        if not part:
            result.append((-1, ""))
            continue
        if part == "UNICO":
            result.append((1, ""))
            continue
        pieces = part.split("-", 1)
        try:
            n = number_value(pieces[0]) if level == 2 else int(pieces[0]) if level != 3 else ord(part[0])
        except (ValueError, KeyError, TypeError):
            n = 0
        result.append((n, pieces[1] if len(pieces) > 1 else ""))
    return tuple(result)


def tail_anchor(base, node, level=None):
    """Insere antes do histórico de um ancestral, depois do último filho atual."""
    own = node.target.level if level is None else None
    level = node.target.level if level is None else level
    number = node.target.parts[node.target.level]
    children = [n for n in base.nodes if node.start < n.start < node.end and node.target.contains(n.target)]
    begin = max([node.head_end] + [n.head_end for n in children])
    for i in range(begin, node.end):
        b = base.blocks[i]
        info = label(b.text) if b.role == 'old' else None
        # Histórico de um ancestral: rótulo superior, ou do mesmo nível mas de outro irmão.
        if info and (info[0] < level or (own is not None and info[0] == own and info[1] != number)):
            anchor = i
            while anchor > begin and (base.blocks[anchor - 1].role == 'old_note' or not base.blocks[anchor - 1].text):
                anchor -= 1
            return anchor
    return node.end


def insertion_anchor(base, target):
    if not target.article or target.corpus not in base.corpora:
        return None, "O artigo ou o corpo da legislação não foi identificado."
    if target.level > 0:
        parent = base.find(target.parent())
        if not parent:
            return None, "O dispositivo superior não foi encontrado de forma única."
        siblings = [n for n in base.nodes if n.target.level == target.level and n.target.parent() == target.parent()]
        if target.level == 1:
            if any(n.target.paragraph == "UNICO" for n in siblings) and target.paragraph != "UNICO":
                return None, "Existe parágrafo único. Cadastre primeiro sua renumeração se o ato assim determinar."
            if target.paragraph == "UNICO" and siblings:
                return None, "Já existem parágrafos numerados neste artigo."
        before = [n for n in siblings if sort_key(n.target) > sort_key(target)]
        after = [n for n in siblings if sort_key(n.target) < sort_key(target)]
        if after:
            # Logo depois do irmão anterior (e do seu histórico), antes do histórico do superior.
            anchor = tail_anchor(base, max(after, key=lambda n: sort_key(n.target)))
            if before:
                anchor = min(anchor, min(before, key=lambda n: sort_key(n.target)).lead)
            return anchor, ""
        if before:
            return min(before, key=lambda n: sort_key(n.target)).lead, ""
        # Incisos do caput são inseridos antes dos parágrafos do artigo.
        children = [n for n in base.nodes if parent.start < n.start < parent.end and not target.contains(n.target)]
        if children and target.level == 2 and not target.paragraph:
            pars = [n for n in children if n.target.level == 1]
            if pars:
                return min(n.lead for n in pars), ""
        # Sem irmãos, insere após o texto atual e antes do histórico do pai.
        return tail_anchor(base, parent, target.level), ""
    siblings = [n for n in base.nodes if n.target.corpus == target.corpus and n.target.level == 0]
    if not siblings:
        return None, "Nenhum artigo foi localizado neste corpo da legislação."
    before = [n for n in siblings if sort_key(n.target) > sort_key(target)]
    if before:
        return min(before, key=lambda n: sort_key(n.target)).lead, ""
    return max(siblings, key=lambda n: sort_key(n.target)).end, ""


def has_children(base, node):
    return bool(node) and any(node.start < n.start < node.end and node.target.contains(n.target) for n in base.nodes)


def source_target(op):
    """Destino que existe hoje na base (o parágrafo único, numa renumeração)."""
    if op.renumber_from:
        return replace(op.target, paragraph=op.renumber_from)
    return op.target


def _fail(op, reason, status="revisar"):
    op.status = status
    op.reason = reason
    op.selected = False


def validate_operation(op, base, default_effects=""):
    op.applied = False
    op.before = []
    op.base_location = ""
    op.effects = date_iso(op.effects or default_effects)
    op.status = "revisar"
    op.reason = ""
    if op.foreign and not op.manual:
        return _fail(op, op.extraction_warning or "O comando pertence a outra legislação.", "outro ato")
    if op.action not in ("replace", "add", "revoke", "rename"):
        return _fail(op, "Escolha substituir, acrescentar, revogar ou renumerar.")
    if op.action == "add" and op.new_lines and op.target.article:
        # Os irmãos acrescidos na mesma redação ("§ 23 ... § 30") formam uma só nota.
        context, roots = list(op.target.parts), []
        for text in op.new_lines:
            info = label(text)
            if info:
                context[info[0]:] = [info[1]] + [""] * (4 - info[0])
                if info[0] == op.target.level:
                    roots.append(Target(op.target.corpus, *context))
        if roots and roots[0] == op.target and all(r.parent() == op.target.parent() for r in roots):
            op.extra_targets = roots[1:]
    elif op.action not in ("add", "revoke"):
        op.extra_targets = []
    targets = op.targets
    if len(targets) > 1:
        if op.action not in ("add", "revoke"):
            return _fail(op, "Só acréscimos e revogações podem reunir vários dispositivos numa mesma nota.")
        if any(t.level != op.target.level or t.parent() != op.target.parent() for t in targets) or len(set(targets)) != len(targets):
            return _fail(op, "Os dispositivos reunidos precisam ser irmãos do mesmo dispositivo superior.")
    src = source_target(op)
    node = base.find(src)
    op.caput = bool(op.action == "replace" and op.scope == "head" and has_children(base, node))
    if node:
        op.before = base.current_lines(node, op.scope == "subtree")
        op.base_location = base.location(node)
    if op.action == "revoke" and len(targets) > 1:
        nodes = [base.find(t) for t in targets]
        if all(nodes):
            op.before = [s for n in nodes for s in base.current_lines(n, True)]
    for t in targets + ([src] if src != op.target else []):
        if len(base.by_target.get(t, [])) > 1:
            return _fail(op, "Numeração duplicada no texto atual. Resolva a duplicidade no Word antes de aplicar.")
    if not op.target.article:
        return _fail(op, "Informe o artigo e a hierarquia completa do destino.")
    if op.action in ("replace", "revoke", "rename") and not node:
        return _fail(op, "Destino não encontrado no texto atual. O histórico não é usado como destino.")
    if op.renumber_from:
        if op.action != "replace" or op.target.level != 1 or op.renumber_from != "UNICO":
            return _fail(op, "A renumeração automática aceita parágrafo único passando a § 1º.")
        if base.by_target.get(op.target):
            return _fail(op, "O parágrafo de destino da renumeração já existe.")
    if op.action == "revoke" and len(targets) > 1:
        nodes = [base.find(t) for t in targets]
        if not all(nodes):
            missing = [t.label() for t, n in zip(targets, nodes) if not n]
            return _fail(op, "Destino não encontrado no texto atual: " + ", ".join(missing) + ".")
        siblings = sorted([n for n in base.nodes if n.target.level == op.target.level and n.target.parent() == op.target.parent()], key=lambda n: n.start)
        positions = [siblings.index(n) for n in nodes]
        if positions != list(range(positions[0], positions[0] + len(nodes))):
            return _fail(op, "Os dispositivos revogados não estão em sequência na base. Separe a revogação em propostas próprias.")
    if op.action == "add":
        for t in targets:
            existing = base.find(t)
            if existing:
                current = base.current_lines(existing, True)
                if len(targets) == 1 and comparable_lines(current) == comparable_lines(op.new_lines):
                    return _fail(op, "A mesma redação já consta no destino.", "já consta")
                return _fail(op, f"O destino do acréscimo ({t.label()}) já existe. Confira a ação e o rótulo.")
        anchor, reason = insertion_anchor(base, op.target)
        if reason:
            return _fail(op, reason)
        if len(targets) > 1:
            siblings = [n for n in base.nodes if n.target.level == op.target.level and n.target.parent() == op.target.parent()]
            low, high = sort_key(targets[0]), sort_key(targets[-1])
            if any(low < sort_key(n.target) < high for n in siblings):
                return _fail(op, "Há dispositivos existentes entre os acréscimos reunidos. Separe o acréscimo em propostas próprias.")
        op.base_location = f"{op.target.corpus}, inserir antes do bloco {anchor + 1} do Word"
    if op.action in ("replace", "add"):
        if not op.new_lines or any(not s.strip() for s in op.new_lines):
            return _fail(op, "Informe a nova redação completa, com os rótulos dos dispositivos.")
        first = label(op.new_lines[0])
        if not first or first[0] != op.target.level or first[1] != op.target.parts[op.target.level]:
            return _fail(op, "O rótulo da nova redação não corresponde ao destino selecionado.")
        if any(re.search(r"\.{3,}|…", s) for s in op.new_lines):
            return _fail(op, "Há reticências no novo texto. Informe somente a redação completa que será inserida.")
        # Impede que uma citação de outro artigo entre no interior deste.
        context = list(op.target.parts)
        new_targets = []
        roots = set(targets)
        for text in op.new_lines:
            info = label(text)
            if not info:
                continue
            level, num, *_ = info
            context[level:] = [num] + [""] * (4 - level)
            t = Target(op.target.corpus, *context)
            if not any(r.contains(t) for r in roots):
                return _fail(op, "O novo texto contém dispositivo fora da hierarquia do destino.")
            if t in new_targets:
                return _fail(op, "A nova redação repete um destino na mesma hierarquia.")
            new_targets.append(t)
        if [t for t in new_targets if t.level == op.target.level] != targets:
            return _fail(op, "Os dispositivos da nova redação não correspondem aos destinos reunidos.")
        descendants = [t for t in new_targets if t not in roots]
        if op.scope == "head" and descendants:
            # Só aceita filhos novos junto do caput quando o dispositivo ainda não os tinha.
            child_levels = {t.level for t in descendants}
            existing = [n for n in base.nodes if node and node.start < n.start < node.end and n.target.level in child_levels
                        and n.target.parent() == src]
            if op.action != "replace" or existing or any(base.by_target.get(t) for t in descendants):
                return _fail(op, "Há descendentes na nova redação. Escolha Dispositivo completo ou separe as operações.")
        if node and op.action == "replace" and not op.renumber_from and comparable_lines(base.current_lines(node, op.scope == "subtree")) == comparable_lines(op.new_lines):
            return _fail(op, "A mesma redação já consta no destino.", "já consta")
    if op.action == "revoke" and node:
        if all(re.search(r"\bREVOGAD[OA]\b|\(revogad[oa]\)", (base.find(t) or node).text, re.I) for t in targets):
            return _fail(op, "O destino já está marcado como revogado.", "já consta")
        op.scope = "subtree"
    if op.action == "rename":
        if op.target.level != 1 or not re.fullmatch(r"\d+", op.rename_to):
            return _fail(op, "Nesta versão, renumere um parágrafo único para um parágrafo numérico.")
        if op.target.paragraph != "UNICO":
            return _fail(op, "A renumeração automática aceita parágrafo único como origem.")
        destination = replace(op.target, paragraph=op.rename_to)
        if base.by_target.get(destination):
            return _fail(op, "O número de destino da renumeração já existe.")
    for t in targets:
        n = base.find(source_target(op) if t == op.target else t)
        if not n:
            continue
        stop = n.end if op.scope == "subtree" else n.head_end
        if any(b.role in ("table", "other") for b in base.blocks[n.start:stop]):
            return _fail(op, "O intervalo contém tabela ou estrutura não textual. Faça essa alteração no Word.")
    if node and op.scope == "subtree" and op.action == "replace" and not op.manual and not op.full_text:
        old_children = [n for n in base.nodes if node.start < n.start < node.end]
        mapped = set(new_targets)
        if op.renumber_from:
            mapped = {replace(t, paragraph=op.renumber_from) if t.paragraph == op.target.paragraph else t for t in mapped}
        missing = [n for n in old_children if n.target not in mapped]
        if missing:
            return _fail(op, "A citação completa omitiria descendentes existentes. Confira se é substituição integral ou parcial.")
    if op.effects and node:
        for b in base.blocks[node.lead:node.start]:
            if b.role == "new_note" and note_matches(b.text, src):
                m = re.search(r"efeitos\s+a\s+partir\s+de\s+(\d{1,2}[./]\d{1,2}[./]\d{2,4})", b.text, re.I)
                if m:
                    try:
                        prior = date_iso(m[1])
                    except ValueError:
                        continue
                    future_replacement = (op.action == "replace" and op.effects == prior
                                          and op.publication and op.publication < prior)
                    if op.effects <= prior and not future_replacement:
                        return _fail(op, "A data proposta não é posterior à data da redação atual. Confira as versões temporais antes de consolidar.")
    if op.extraction_warning and not op.manual:
        return _fail(op, op.extraction_warning.strip())
    op.status = "pronto"
    op.reason = "Destino e redação conferidos por regras." if not op.manual else "Correção manual validada pela estrutura."
    if op.note and not op.manual:
        op.reason += " " + op.note


DEPENDENT = ("Destino não encontrado", "superior não foi encontrado", "parágrafo único", "já existe", "Existe parágrafo")


def validate_plan(plan, base=None):
    base = base or Legislation(plan.converted_base)
    for op in plan.operations:
        was_selected = op.selected
        validate_operation(op, base, plan.effects)
        if op.status == "pronto":
            op.selected = was_selected
    # Uma mudança pode depender de outra do mesmo plano (renumeração, novo artigo ou
    # novo parágrafo), venha ela antes ou depois na ordem do ato: "acrescido o § 2º
    # ao art. 5º" (art. 2º do decreto) só cabe depois de "renumerado o parágrafo único
    # do art. 5º para § 1º" (art. 4º). Simula as mudanças numa cópia, na mesma ordem
    # da geração: as dependentes que ainda não cabem esperam a rodada seguinte.
    waiting = [i for i, op in enumerate(plan.operations)
               if op.status == "revisar" and any(s in op.reason for s in DEPENDENT)
               and any(x.status == "pronto" and x.selected for j, x in enumerate(plan.operations) if j != i)]
    if waiting:
        stream = BytesIO()
        base.doc.save(stream)
        stream.seek(0)
        simulation = Document(stream)
        settings = note_settings(plan, base)
        queue = [i for i, op in enumerate(plan.operations) if i in waiting or (op.status == "pronto" and op.selected)]
        while queue:
            later = []
            for i in queue:
                op = plan.operations[i]
                try:
                    if i in waiting:
                        validate_operation(op, Legislation(simulation), plan.effects)
                        if op.status != "pronto":
                            if any(s in op.reason for s in DEPENDENT):
                                later.append(i)
                            continue
                        op.selected = True
                        op.base_location = "Aplicada depois das outras mudanças deste plano de que depende."
                    apply_operation(simulation, deepcopy(op), plan.citation or "Ato alterador", base.profile, **settings)
                except ValueError:
                    continue
            if len(later) == len(queue) or not later:
                break
            queue = later
    for i, op in enumerate(plan.operations):
        if op.status != "pronto":
            continue
        for other in plan.operations[i + 1:]:
            if other.status != "pronto":
                continue
            conflict = any(a == b or (op.scope == "subtree" and a.contains(b)) or (other.scope == "subtree" and b.contains(a))
                           for a in op.targets for b in other.targets)
            if conflict:
                op.status = other.status = "revisar"
                op.selected = other.selected = False
                op.reason = other.reason = f"Operações sobrepostas: {op.id} e {other.id}. Una ou separe as mudanças antes de aplicar."
    return plan


def note_settings(plan, base=None):
    style = plan.resolved_date_style if plan.date_style in ("", "auto") else plan.date_style
    return {"style": style or "4", "corpus_suffix": plan.adct_suffix}


def analyze(base_path, amender_path, cache, citation="", publication="", effects="", date_style="auto"):
    converted = prepare_base(base_path, cache)
    base = Legislation(converted)
    if not base.nodes:
        raise ValueError("Não foram encontrados artigos no texto atual da legislação base.")
    lines, warnings = read_amender(amender_path, cache)
    resolved_style = notes_mod.detect_date_style(base.blocks)
    style = resolved_style if date_style in ("", "auto", None) else date_style
    citation = citation.strip() or notes_mod.format_citation(source_citation(lines), style)
    publication = date_iso(publication) or publication_date(lines)
    plan = Plan(str(Path(base_path).resolve()), str(Path(amender_path).resolve()), str(converted), citation, publication, date_iso(effects),
                digest(base_path), digest(amender_path), warnings=base.warnings + warnings, style_profile=base.profile, identity=base.identity)
    plan.date_style = date_style or "auto"
    plan.resolved_date_style = resolved_style
    plan.adct_suffix = notes_mod.detect_adct_suffix(base)
    plan.operations = parse_operations(lines, base)
    plan.warnings.extend(assign_effects(lines, plan.operations, base, plan.effects, publication))
    plan.related_provisions = collect_remissions(lines, plan.operations, base, citation)
    for op in plan.operations:
        op.selected = True
        op.publication = publication
    if not citation:
        plan.warnings.append("Informe a identificação do ato alterador antes de gerar o documento.")
    if any(not op.effects for op in plan.operations):
        plan.warnings.append("Data de efeitos não identificada para todas as mudanças. As notas não inventam datas; informe a publicação ou a data de efeitos, se necessário.")
    if publication:
        plan.warnings.append("A data de publicação só é usada como data de efeitos quando o ato diz que entra em vigor ou produz efeitos na data de sua publicação.")
    style = note_settings(plan)["style"]
    plan.warnings.append(f"Datas e anos nas notas: {'dois dígitos (dd.mm.aa)' if style == '2' else 'quatro dígitos (dd.mm.aaaa)'}, conforme as notas mais recentes da base.")
    if any(re.search(r"consolidad[ao]\s+ate|atualizad[ao]\s+ate|alterad[ao]\s+pel", norm(b.text)) for b in base.blocks[:18]):
        plan.warnings.append("Confira ao final as referências iniciais de atualização da legislação. Essas notas são preservadas e não afirmam a completude das mudanças selecionadas.")
    if str(base_path).lower().endswith(".doc"):
        plan.warnings.append("A base .doc foi convertida para .docx. Confira no Word a fidelidade da conversão.")
    validate_plan(plan, base)
    if str(amender_path).lower().endswith((".pdf", ".html", ".htm")):
        for op in plan.operations:
            if op.status == "pronto":
                op.status = "revisar"
                op.reason = "Confira a extração de PDF/HTML com a origem e clique em Validar correção (ou em Aprovar extração do PDF/HTML para todas)."
                op.selected = False
                op.extraction_warning = op.reason
    for op in plan.operations:
        op.selected = op.status == "pronto"
    return plan, base


def ensure_styles(doc, profile):
    source = etree.parse(str(ROOT / "modelos/estilos_sefa.xml")).getroot()
    by_name = {s.find(qn("w:name")).get(qn("w:val")): s for s in source if s.find(qn("w:name")) is not None}
    by_id = {s.get(qn("w:styleId")): s for s in source}

    def copy_style(s):
        sid = s.get(qn("w:styleId"))
        if any(x.get(qn("w:styleId")) == sid for x in doc.styles.element):
            return
        parent = s.find(qn("w:basedOn"))
        if parent is not None and parent.get(qn("w:val")) in by_id:
            copy_style(by_id[parent.get(qn("w:val"))])
        doc.styles.element.append(deepcopy(s))
    for name in profile.values():
        if name not in doc.styles:
            if name not in by_name:
                raise ValueError("Estilo não disponível: " + name)
            copy_style(by_name[name])


def paragraph(doc, text, style, bold_label=False):
    el = OxmlElement("w:p")
    p = Paragraph(el, doc._body)
    p.style = style
    if p._p.pPr is not None:
        for name in ("ind", "spacing", "jc"):
            for prop in list(p._p.pPr.findall(qn("w:" + name))):
                p._p.pPr.remove(prop)
    info = label(text) if bold_label else None
    if info and info[0] == 0:
        end = info[2]
        p.add_run(text[:end]).bold = True
        p.add_run(text[end:])
    else:
        p.add_run(text)
    return el


def historical(el, doc, style):
    el = deepcopy(el)
    if el.tag != qn("w:p"):
        return el
    p = Paragraph(el, doc._body)
    p.style = style
    if p._p.pPr is not None:
        for name in ("ind", "spacing", "jc"):
            for prop in list(p._p.pPr.findall(qn("w:" + name))):
                p._p.pPr.remove(prop)
    # A cor do histórico vem do estilo, inclusive em hyperlinks arquivados.
    # Mantém negrito, itálico, campos e os relacionamentos da própria base.
    for rpr in el.iter(qn("w:rPr")):
        for name in ("color", "sz", "szCs", "rStyle"):
            for attr in list(rpr.findall(qn("w:" + name))):
                rpr.remove(attr)
    return el


def prior_note(text, effects, citation="", publication="", style="4"):
    return notes_mod.prior_note(text, effects, citation, publication, style)


NOTE_START = ("redacao dada", "nova redacao", "acrescido", "acrescida", "renumerado", "renumerada", "revogado", "revogada")


def note_targets(text, corpus="PRINCIPAL"):
    """Destinos de uma nota de redação, sem combinar números de dispositivos diferentes."""
    found = refparser.parse(text)
    targets = []
    for r in found:
        if refparser.is_self(r.law):
            continue
        ts = r.targets(corpus)
        if r.warning:
            return None
        targets.extend(ts)
    return targets


def note_matches(text, target):
    """Nota individual do próprio destino (uma nota coletiva fica com os demais)."""
    n = norm(text)
    if not n.startswith(NOTE_START):
        return False
    if re.search(r"\b(?:incisos|alineas|paragrafos|arts\.|artigos|itens)\b", n) or "§§" in text:
        return False
    ts = note_targets(text.split(" pel")[0] if " pel" in text else text, target.corpus)
    if not ts:
        return False
    if n.startswith(("renumerado", "renumerada")):
        return any(t.article == target.article and t.level <= target.level and t.contains(target) or t == target for t in ts)
    return len(ts) == 1 and ts[0] == target


def attached_note(text, target):
    """Nota colocada logo acima do dispositivo que remete a ele ou a um ancestral
    ("Redação dada ao art. 78 ..." acima da alínea “b”): pertence a este texto.
    Também vale quando só o nível intermediário mudou de nome (inciso I do
    parágrafo único que passou a ser inciso I do § 1º)."""
    n = norm(text)
    if not n.startswith(NOTE_START):
        return False
    ts = note_targets(text.split(" pel")[0] if " pel" in text else text, target.corpus)
    if not ts or len(ts) != 1:
        return False
    t = ts[0]
    if t.contains(target):
        return True
    level = target.level
    return (t.article == target.article and t.level == level and level >= 2
            and t.parts[level] == target.parts[level] and t.parts[1] in ("UNICO", "1") and target.parts[1] in ("UNICO", "1"))


def is_original(base, node):
    """Sem nota de origem: 'Redação original', salvo se a base já guarda outra versão deste rótulo."""
    article = base.find(Target(node.target.corpus, node.target.article))
    start, end = (article.start, article.end) if article else (node.start, node.end)
    level = node.target.level
    number = node.target.parts[level]
    for b in base.blocks[start:end]:
        if b.role == "old":
            info = label(b.text)
            if info and info[0] == level and info[1] == number:
                return False
    return True


def new_note(op, citation, style="4", corpus_suffix=True):
    return "\n".join(notes_mod.new_notes(op, citation, style, corpus_suffix))


def effects_text(effects, style="4"):
    return notes_mod.effects_text(effects, style)


def highlight_new_elements(doc, elements, enabled):
    """Realça só os parágrafos novos; os estilos e a cor da fonte são mantidos."""
    if enabled:
        for el in elements:
            for run in Paragraph(el, doc._body).runs:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _history(doc, base, node, op, citation, profile, style, scope):
    """Notas e textos que vão para o histórico de um dispositivo; e o que sai da base."""
    end = node.end if scope == "subtree" else node.head_end
    target = node.target
    attached = []
    for b in base.blocks[node.lead:node.start]:
        if b.role == "new_note" and (note_matches(b.text, target) or attached_note(b.text, target)):
            attached.append(b)
    origin, origin_targets = origin_note(base, node)
    collective = bool(origin and len(origin_targets) > 1 and target in origin_targets and origin not in attached)
    inserts = []
    archived = []
    if collective:
        remaining = [t for t in origin_targets if t != target]
        remaining_nodes = [base.find(t) for t in remaining]
        remaining_nodes = [n for n in remaining_nodes if n]
        if remaining_nodes:
            first = min(remaining_nodes, key=lambda n: n.start)
            inserts.append((base.blocks[first.start].element, paragraph(doc, note_for_targets(origin.text, remaining), profile["new_note"])))
            attached.append(origin)
    if attached:
        for b in attached:
            text = note_for_targets(b.text, [target]) if collective and b is origin else b.text
            archived.append(replace_text(historical(b.element, doc, profile["old_note"]), prior_note(text, op.effects, citation, op.publication, style)))
    elif origin:
        text = note_for_targets(origin.text, [target])
        archived.append(paragraph(doc, prior_note(text, op.effects, citation, op.publication, style), profile["old_note"]))
    else:
        archived.append(paragraph(doc, notes_mod.fallback_note(is_original(base, node), op.effects, style), profile["old_note"]))
    for b in base.blocks[node.start:end]:
        if b.role == "current":
            archived.append(historical(b.element, doc, profile["old"]))
        elif b.role == "new_note":
            archived.append(replace_text(historical(b.element, doc, profile["old_note"]), prior_note(b.text, op.effects, citation, op.publication, style)))
        else:
            archived.append(deepcopy(b.element))
    removed = [b.element for b in base.blocks[node.start:end]] + [b.element for b in attached]
    return archived, removed, inserts


def revoked_line(node):
    gender = "REVOGADA" if node.target.level == 3 else "REVOGADO"
    return node.label_text.rstrip() + " " + gender


def apply_operation(doc, op, citation, profile, highlight_changes=False, style="4", corpus_suffix=True):
    base = Legislation(doc)
    original_before = list(op.before)
    original_location = op.base_location
    validate_operation(op, base)
    if op.status != "pronto":
        raise ValueError(f"{op.id}: {op.reason}")
    if original_location:
        op.before = original_before
        op.base_location = original_location
    node = base.find(source_target(op))
    body = doc.element.body
    fresh = [paragraph(doc, text, profile["new_note"]) for text in notes_mod.new_notes(op, citation, style, corpus_suffix)]
    if op.action in ("replace", "add"):
        cite = notes_mod.format_citation(citation, style)
        fresh.extend(paragraph(doc, text.replace(citation, cite), profile["new_note"]) for text in op.remissions)
    if op.action == "add":
        anchor, reason = insertion_anchor(base, op.target)
        if reason:
            raise ValueError(reason)
        fresh.extend(paragraph(doc, s, profile["current"], True) for s in op.new_lines)
        highlight_new_elements(doc, fresh, highlight_changes)
        for offset, el in enumerate(fresh):
            body.insert(anchor + offset, el)
        op.applied = True
        return
    nodes = [node] if op.action != "revoke" else sorted((base.find(t) for t in op.targets), key=lambda n: n.start)
    if op.action == "revoke":
        fresh.extend(paragraph(doc, revoked_line(n), profile["current"], True) for n in nodes)
    elif op.action == "rename":
        info = label(node.text)
        fresh.append(paragraph(doc, f"§ {op.rename_to}" + ('º' if int(op.rename_to) <= 9 else '.') + " " + node.text[info[2]:], profile["current"], True))
    else:
        fresh.extend(paragraph(doc, s, profile["current"], True) for s in op.new_lines)
    highlight_new_elements(doc, fresh, highlight_changes)
    scope = "subtree" if op.action == "revoke" else op.scope
    archived, removed, inserts = [], [], []
    for n in nodes:
        a, r, ins = _history(doc, base, n, op, citation, profile, style, scope)
        archived.extend(a)
        removed.extend(r)
        inserts.extend(ins)
    anchor = nodes[0].element if hasattr(nodes[0], "element") else base.blocks[nodes[0].start].element
    for el in fresh + archived:
        anchor.addprevious(el)
    for where, el in inserts:
        where.addprevious(el)
    for el in removed:
        if el.getparent() is not None:
            body.remove(el)
    op.applied = True


def protect_outputs(plan, paths):
    protected = {Path(p).resolve() for p in (plan.base_path, plan.amender_path, plan.converted_base)}
    if any(Path(p).resolve() in protected for p in paths):
        raise ValueError("Escolha outro nome para a saída. Os arquivos de entrada são preservados.")


def generate(plan, output_path, highlight_changes=False):
    if not plan.citation.strip():
        raise ValueError("Informe a identificação do ato alterador, como Lei 12.345/2026.")
    if digest(plan.base_path) != plan.base_hash or digest(plan.amender_path) != plan.amender_hash:
        raise ValueError("Um dos arquivos mudou depois da análise. Analise novamente antes de gerar.")
    out = Path(output_path).resolve()
    protect_outputs(plan, [out])
    if out.suffix.lower() != ".docx":
        raise ValueError("A saída deve ter a extensão .docx.")
    selected = [op for op in plan.operations if op.selected]
    if not selected:
        raise ValueError("Selecione pelo menos uma mudança validada.")
    if any(op.status != "pronto" for op in selected):
        raise ValueError("Há mudanças selecionadas que ainda precisam de revisão.")
    # Revalida conflitos sem alterar o plano que a interface apresenta.
    check = Plan.from_dict(plan.to_dict())
    validate_plan(check)
    if any(op.selected and op.status != "pronto" for op in check.operations) or any(check.operations[i].status != "pronto" for i, op in enumerate(plan.operations) if op.selected):
        raise ValueError("As mudanças selecionadas têm conflito estrutural. Revise o plano.")
    doc = Document(plan.converted_base)
    original_links = links_signature(doc)
    profile = Legislation(doc).profile
    ensure_styles(doc, profile)
    settings = note_settings(plan)
    originals = {op.id: op.to_dict() for op in selected}
    try:
        # Na ordem do ato; a mudança que depende de outra posterior (o § 2º novo depois
        # da renumeração do parágrafo único) é aplicada na rodada seguinte.
        pending = list(selected)
        while pending:
            later, errors = [], []
            for op in pending:
                saved = (list(op.before), op.base_location)
                try:
                    apply_operation(doc, op, plan.citation, profile, highlight_changes=highlight_changes, **settings)
                    op.selected = True
                except ValueError as error:
                    if not any(s in str(error) for s in DEPENDENT):
                        raise
                    op.before, op.base_location = saved
                    later.append(op)
                    errors.append(error)
            if len(later) == len(pending):
                raise errors[0]
            pending = later
        update_initial_reference(doc, notes_mod.format_citation(plan.citation, settings["style"]))
        if links_signature(doc) != original_links:
            raise ValueError("A geração alteraria os hyperlinks originais. O Word não foi salvo; revise os trechos com links.")
        # Continua sendo obrigatório abrir como DOCX válido antes da publicação.
        out.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(suffix=".docx", dir=out.parent, delete=False) as temp:
            temporary = Path(temp.name)
        try:
            save_preserving_package(doc, plan.converted_base, temporary)
            Document(temporary)
            if links_signature(Document(temporary)) != original_links:
                raise ValueError("A validação final detectou alteração nos hyperlinks. O Word não foi salvo.")
            os.replace(temporary, out)
        finally:
            temporary.unlink(missing_ok=True)
    except Exception:
        for op in selected:
            restored = Operation.from_dict(originals[op.id])
            op.__dict__.update(restored.__dict__)
        raise
    return out


def save_plan(plan, path):
    protect_outputs(plan, [path])
    if Path(path).suffix.lower() != ".json":
        raise ValueError("Salve o plano com a extensão .json.")
    value = {"format": "consolidador-sefa", "version": 1, "created": datetime.now().isoformat(timespec="seconds"), "plan": plan.to_dict()}
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def load_plan(path, cache):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if data.get("format") != "consolidador-sefa" or data.get("version") != 1:
        raise ValueError("Esse arquivo não é um plano desta versão do Consolidador.")
    plan = Plan.from_dict(data["plan"])
    if digest(plan.base_path) != plan.base_hash or digest(plan.amender_path) != plan.amender_hash:
        raise ValueError("Os arquivos deste plano foram alterados. Faça uma nova análise.")
    plan.converted_base = str(prepare_base(plan.base_path, cache))
    base = Legislation(plan.converted_base)
    validate_plan(plan, base)
    return plan, base
