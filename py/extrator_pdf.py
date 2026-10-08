"""Importação conservadora de atos em PDF com tabelas vetoriais editáveis.

A extração nunca calcula valores. Cada número é lido de suas coordenadas e
comparado com o conjunto de números presente nas colunas da tabela do PDF.
PDF digitalizado ou grade não reconhecida requer revisão manual.
"""

from __future__ import annotations

import html
import io
import os
import re
import tempfile
import unicodedata
from collections import Counter
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pdfplumber
import fitz
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt

from formatador_sefa import (
    ANNEX, ARTICLE, CARGO, COMMAND, EMENTA_START, ENDS_WITH_ABBREVIATION, EPIGRAPH, HEADING, NAME, PREAMBLE,
    SIGNER_CARGO, Item, continues_epigraph, join_line, looks_like_signer_name, looks_uppercase, vocabulary,
    add_formatted_paragraph, apply_styles, classify, clear_model_contents,
    ementa_spacer_before, looks_like_person_name, paragraph_style_names,
)


MONEY = re.compile(r"(?<!\d)-?(?:\d{1,3}(?:\.\d{3})+|\d+),\d{2}(?![\d%])")
PREFIXED_CURRENCY = re.compile(r"R\$\s*(-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d{2})?)")
PERCENT = re.compile(r"-?\d+(?:\.\d{3})*,\d+(?:\s*%)")
NUMERIC_CELL = re.compile(r"\s*(?:R\$\s*)?-?\d[\d.,]*(?:\s*%)?\s*")
ROW_START = re.compile(r"^(?:\d+(?:\.\d+)*\s*[-–.]|[IVXLCDM]+\s*[-–]|TOTAIS\b)", re.I)
FOOTNOTE = re.compile(r"^[¹²³⁴⁵⁶⁷⁸⁹]")
PROTOCOL = re.compile(r"^PROTOCOLO\s*:\s*\d+", re.I)
SECTION_HEAD = re.compile(r"^ATOS\s+(?:DO|DA|DOS|DAS)\b", re.I)
ANNEX_HEAD = re.compile(r"^(?:GOVERNO DO ESTADO|DEMONSTRATIVO\b|ANEXO\b|R\$\s*1,00)", re.I)
PDF_DATE = re.compile(r"(\d{4})[._-](\d{2})[._-](\d{2})")
PORTUGUESE_DATE = re.compile(
    r"\b(\d{1,2})\s+de\s+(janeiro|fevereiro|março|abril|maio|junho|julho|agosto|setembro|outubro|novembro|dezembro)\s+de\s+(\d{4})\b",
    re.I,
)
MONTH_NUMBERS = {name: f"{index:02d}" for index, name in enumerate(
    ("janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto",
     "setembro", "outubro", "novembro", "dezembro"), 1)}


@dataclass
class ParagraphBlock:
    text: str
    role: str = "texto"
    page: int = 0
    kind: str = "norma"  # norma, annex_head, unit, footnote, date, signer
    item: Item | None = None
    attention: list[str] = field(default_factory=list)


@dataclass
class ImageBlock:
    data: bytes
    page: int
    kind: str = "image"


@dataclass
class TableBlock:
    matrix: list[list[str]]
    column_widths: list[float]
    page: int
    bbox: tuple[float, float, float, float]
    source_numbers: Counter
    source_signature: Counter = field(default_factory=Counter)
    original_matrix: list[list[str]] = field(default_factory=list)
    strategy: str = "linhas"
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)
    reviewed: bool = False
    kind: str = "table"
    source_parts: list[TablePagePart] = field(default_factory=list)


@dataclass
class TablePagePart:
    page: int
    bbox: tuple[float, float, float, float]
    first_row: int
    last_row: int  # exclusivo
    column: str = ""  # esquerda, direita ou vazio em páginas sem colunas


def table_page_parts(table: TableBlock) -> list[TablePagePart]:
    return table.source_parts or [TablePagePart(table.page, table.bbox, 0, len(table.matrix))]


@dataclass
class Extraction:
    pdf_path: Path
    query: str
    blocks: list[ParagraphBlock | ImageBlock | TableBlock]
    pages: list[int]
    issues: list[str] = field(default_factory=list)
    blocking_issues: list[str] = field(default_factory=list)
    limits_reviewed: bool = False
    protocol: str = ""
    intertable_segments_checked: int = 0

    @property
    def tables(self) -> list[TableBlock]:
        return [b for b in self.blocks if isinstance(b, TableBlock)]

    @property
    def paragraphs(self) -> list[ParagraphBlock]:
        return [b for b in self.blocks if isinstance(b, ParagraphBlock)]


def normalized(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).upper()
    return re.sub(r"\s+", " ", "".join(c for c in text if not unicodedata.combining(c))).strip()


def character_signature(text: str) -> Counter:
    """Compara todo o texto da tabela, independentemente das quebras e espaços."""
    return Counter(c for c in unicodedata.normalize("NFKC", text) if not c.isspace())


def matrix_signature(matrix: list[list[str]]) -> Counter:
    return character_signature("".join(value for row in matrix for value in row))


def monetary_values(matrix: list[list[str]]) -> Counter:
    numbers: Counter = Counter()
    for row in matrix:
        for value in row:
            for line in value.splitlines():
                if PERCENT.fullmatch(line.strip()):
                    continue
                prefixed = list(PREFIXED_CURRENCY.finditer(line))
                numbers.update(match.group(1) for match in prefixed)
                numbers.update(match.group() for match in MONEY.finditer(line)
                               if not any(m.start() <= match.start() < m.end() for m in prefixed))
    return numbers


def pdf_lines(pdf_path: Path, page_number: int, document: Any | None = None) -> list[dict]:
    """Lê linhas e coordenadas, preservando a caixa das letras sem programas externos."""
    if document is None:
        with fitz.open(str(pdf_path)) as pdf:
            return pdf_lines(pdf_path, page_number, pdf)
    page = document[page_number - 1]
    lines = []
    for block in page.get_text("dict")["blocks"]:
        if block["type"] != 0:
            continue
        for line in block["lines"]:
            value = re.sub(r"\s+", " ", "".join(span["text"] for span in line["spans"])).strip()
            if value:
                x0, top, x1, bottom = line["bbox"]
                lines.append({"text": value, "x0": x0, "top": top,
                              "x1": x1, "bottom": bottom})
    return sorted(lines, key=lambda l: (l["top"], l["x0"]))


def grouped_words(words: list[dict], tolerance: float = 1.8) -> list[tuple[float, list[dict]]]:
    groups: list[tuple[float, list[dict]]] = []
    for word in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if groups and abs(groups[-1][0] - word["top"]) <= tolerance:
            groups[-1][1].append(word)
        else:
            groups.append((word["top"], [word]))
    return groups


def words_inside(words: list[dict], box: tuple[float, float, float, float]) -> list[dict]:
    x0, top, x1, bottom = box
    return [w for w in words if top - .15 <= w["top"] < bottom + .15 and x0 - .15 <= (w["x0"] + w["x1"]) / 2 < x1 + .15]


def cell_lines(words: list[dict], box: tuple[float, float, float, float], concatenate: bool = False) -> list[tuple[float, str]]:
    lines = []
    for top, members in grouped_words(words_inside(words, box)):
        tokens = [w["text"] for w in sorted(members, key=lambda w: w["x0"])]
        lines.append((top, ("" if concatenate else " ").join(tokens)))
    return lines


def source_numbers_for_table(words: list[dict], table: Any) -> Counter:
    numbers: Counter = Counter()
    for row in table.rows:
        for cell in row.cells:
            if cell is None:
                continue
            for _, value in cell_lines(words, cell, concatenate=True):
                if not PERCENT.fullmatch(value):
                    numbers.update(monetary_values([[value]]))
    return numbers


def extract_table(page: Any, table: Any, words: list[dict], page_number: int) -> TableBlock:
    rows: list[list[str]] = []
    issues: list[str] = []
    output_map: list[list[int]] = []
    extracted = table.extract()
    for row_index, geometry in enumerate(table.rows):
        boxes = geometry.cells
        first = boxes[0]
        if first is None:
            if any(extracted[row_index]):
                issues.append(f"Linha {row_index + 1} com célula mesclada: confirme cabeçalho e colunas no PDF.")
                output_map.append([len(rows)])
                rows.append([cell or "" for cell in extracted[row_index]])
            else:
                output_map.append([])
            continue
        labels: list[list] = []
        for top, text in cell_lines(words, first):
            if ROW_START.match(text):
                labels.append([top, text])
            elif labels:
                labels[-1][1] += " " + text
            elif row_index > 0 and text and not (extracted[row_index][0] or "").startswith("ITEM"):
                # Pode ser um cabeçalho de coluna, tratado como linha única.
                pass
        if len(labels) <= 1:
            raw = [cell or "" for cell in extracted[row_index]]
            if not any(raw):
                output_map.append([])
                continue
            # pdfplumber.Table.extract pode cortar o primeiro dígito junto à
            # borda de uma célula. Para linhas simples, reconstituímos o valor
            # a partir das palavras da coluna inteira, incluindo fragmentos.
            for ci in range(1, len(boxes)):
                cell = boxes[ci]
                if cell is None:
                    continue
                values = [text for _, text in cell_lines(words, cell, concatenate=True) if MONEY.fullmatch(text)]
                if len(values) == 1:
                    raw[ci] = values[0]
                elif len(values) > 1:
                    issues.append(f"Linha visual {row_index + 1}, coluna {ci + 1}: vários valores sem rótulo individual.")
            output_map.append([len(rows)])
            rows.append(raw)
            continue
        split = [[label] + [""] * (len(boxes) - 1) for _, label in labels]
        for ci in range(1, len(boxes)):
            cell = boxes[ci]
            if cell is None:
                issues.append(f"Linha visual {row_index + 1}, coluna {ci + 1}: célula irregular.")
                continue
            for top, value in cell_lines(words, cell):
                compact = "".join(value.split())
                if MONEY.fullmatch(compact) or PERCENT.fullmatch(compact):
                    value = compact
                candidates = [i for i, (label_top, _) in enumerate(labels) if label_top <= top + 1.5]
                if not candidates:
                    issues.append(f"Conteúdo {value!r} sem rótulo correspondente na coluna {ci + 1}.")
                target = candidates[-1] if candidates else 0
                if split[target][ci]:
                    issues.append(f"Várias linhas na coluna {ci + 1} do mesmo rótulo: confira a associação.")
                    split[target][ci] += "\n" + value
                else:
                    split[target][ci] = value
        output_map.append(list(range(len(rows), len(rows) + len(split))))
        rows.extend(split)

    merges: list[tuple[int, int, int, int]] = []
    for ri, visual in enumerate(table.rows):
        if len(output_map[ri]) != 1:
            continue
        for ci, cell in enumerate(visual.cells):
            if cell is None:
                continue
            last_col = ci
            for next_col in range(ci + 1, len(visual.cells)):
                if visual.cells[next_col] is not None:
                    break
                if any(other.cells[next_col] and other.cells[next_col][0] < cell[2] - .5
                       for other in table.rows):
                    last_col = next_col
            last_row = ri
            for next_row in range(ri + 1, len(table.rows)):
                top = min((b[1] for b in table.rows[next_row].cells if b), default=cell[3])
                if top >= cell[3] - .5 or table.rows[next_row].cells[ci] is not None:
                    break
                last_row = next_row
            if (last_col > ci or last_row > ri) and all(
                    len(output_map[index]) == 1 for index in range(ri, last_row + 1)):
                merges.append((output_map[ri][0], ci, output_map[last_row][0], last_col))

    expected = source_numbers_for_table(words, table)
    actual = monetary_values(rows)
    if expected != actual:
        issues.append(
            f"A conferência dos números falhou: {sum(expected.values())} lidos no PDF, "
            f"{sum(actual.values())} nas células exportáveis. Revise a tabela antes de usar."
        )
    column_geometry = table.rows[0].cells
    widths = [min((row.cells[ci][2] - row.cells[ci][0]
                   for row in table.rows if row.cells[ci]), default=30.0)
              for ci in range(len(column_geometry))]
    source = character_signature("".join(word["text"] for word in words_inside(words, table.bbox)))
    if matrix_signature(rows) != source:
        issues.append("Parte do texto da tabela no PDF não corresponde às células extraídas; confira todas as células.")
    strategy = getattr(table, "strategy", "linhas")
    if strategy == "alinhamento":
        issues.append("Tabela sem grade identificada por alinhamento: confira linhas e colunas no PDF.")
    return TableBlock(rows, widths, page_number, table.bbox, expected, source,
                      deepcopy(rows), strategy, merges, issues)


# Linha que termina no meio da frase: a seguinte é continuação, ainda que comece por
# "Portaria ME nº 284" (parece epígrafe), "DIRETOR-EXECUTIVO" (parece cargo) ou "o DIRETOR ...,
# no uso" (parece preâmbulo). Ato Técnico Conjunto nº 6/2026, DOU de 01.10.2026.
MID_SENTENCE = re.compile(r"(?:\b(?:de|da|do|das|dos|e|em|o|a|os|as|ao|aos|à|às|no|na|nos|nas|pelo|pela|"
                          r"pelos|pelas|com|para|por|que|sobre|entre|um|uma))\s*$", re.I)   # sem vírgula: "...," / "R E S O L V E:"
ITEM_NUMBER = re.compile(r"^\d{1,2}\.\s+\S")      # "1. Anexo VI ..." / "2. Anexo VII ..."
# Mobília do DOU: cabeçalho, expediente e rodapé de autenticação
DOU_HEADER = re.compile(r"^(?:ISSN\s+\d{4}-\d{3}[\dX]\b|REP[ÚU]BLICA FEDERATIVA DO BRASIL\b|"
                        r"Ano\s+[IVXLCDM]+\s+N[º°o]|Bras[íi]lia\s*-\s*DF\s*,|Di[áa]rio Oficial da Uni[ãa]o\b|"
                        r"EDI[ÇC][ÃA]O EXTRA\b|SE[ÇC][ÃA]O\s+\d\b)", re.I)
DOU_FOOTER = re.compile(r"PRESID[ÊE]NCIA DA REP[ÚU]BLICA\s*[•·]|Este documento pode ser verificado no endere[çc]o|"
                        r"Documento assinado digitalmente conforme MP\b", re.I)   # busca em qualquer ponto da linha


def dou_margins(lines: list[dict], page_height: float) -> tuple[float, float]:
    """Faixa útil de uma página do DOU: abaixo do cabeçalho (ISSN, "Ano CLXIV Nº 186-A", "Brasília - DF, ...",
    número da página) e acima do expediente ("PRESIDÊNCIA DA REPÚBLICA • CASA CIVIL ...") e do rodapé de
    autenticação. Sem essas linhas, devolve a página inteira."""
    header = [l["bottom"] for l in lines if l["top"] < page_height * .15 and DOU_HEADER.match(l["text"].strip())]
    header += [l["bottom"] for l in lines if l["top"] < page_height * .085 and re.fullmatch(r"\d{1,4}", l["text"].strip())]
    footer = [l["top"] for l in lines if l["top"] > page_height * .5 and DOU_FOOTER.search(l["text"])]
    footer += [l["top"] for l in lines if l["top"] > page_height * .92 and re.fullmatch(r"\d{1,4}", l["text"].strip())]
    return (max(header) + .5 if header else 0.0), (min(footer) - 1 if footer else page_height)


PAGE_HEADER = re.compile(r"DI[ÁA]RIO OFICIAL N[º°]\s*[\d.]+|^(?:Segunda|Ter[çc]a|Quarta|Quinta|Sexta)-feira,\s*\d"
                         r"|^(?:S[áa]bado|Domingo),\s*\d|DO-e/SEFA N[º°]:", re.I)
PAGE_HEADER_TOP = re.compile(r"^\d{1,2} de [a-zç]+ de \d{4}$|^\d{1,4}$", re.I)   # só bem no alto (DO-e/SEFA)


def page_top_margin(lines: list[dict], page_height: float) -> float | None:
    """Fim do cabeçalho de uma página do DOE ou do DO-e/SEFA ("2 ▪ DIÁRIO OFICIAL Nº 36.789",
    "Sexta-feira, 02 de OUTUBRO de 2026"; "2 de outubro de 2026 / DO-e/SEFA Nº: 646 / 2"). No DOE o texto
    começa a 40 pt; com a faixa fixa de 52 pt, as primeiras linhas de cada continuação se perdiam."""
    # Só a faixa de cima (até 36 pt): o texto dos atos começa a 40 pt no DOE e pode citar "DIÁRIO OFICIAL Nº ...".
    header = [l["bottom"] for l in lines if l["top"] < 36 and PAGE_HEADER.search(l["text"].strip())]
    if not header:
        return None
    header += [l["bottom"] for l in lines if l["top"] < 36 and PAGE_HEADER_TOP.match(l["text"].strip())]
    return max(header) + .5


def three_columns(lines: list[dict], width: float) -> bool:
    """Página com três colunas de texto (cada uma com ao menos 8 linhas cheias dentro de um terço)."""
    third = width / 3
    counts = [0, 0, 0]
    for line in lines:
        size = line["x1"] - line["x0"]
        if .2 * width <= size <= .34 * width:
            column = int(line["x0"] // third)
            if column < 3 and line["x1"] <= (column + 1) * third + 10:
                counts[column] += 1
    return min(counts) >= 8


PLACE_DATE = re.compile(r"^[^,;:]+,\s*\d{1,2}[º°]?\s+de\s+\w+\s+de\s+\d{4}\.?$", re.I)
OPEN_ENDING = re.compile(r"(?:\b(?:de|da|do|das|dos|e|em)|[-,])\s*$", re.I)


def signature_break(paragraphs: list[str], text: str) -> bool:
    """Nome / cargo / nome / cargo em linhas próprias, sem maiúsculas nem pontuação que
    separem os parágrafos: "Ana Carolina Lobo Gluck Paúl" / "Procuradora-Geral do Estado do Pará"."""
    last = paragraphs[-1]
    if looks_like_signer_name(last) and SIGNER_CARGO.match(text):
        return True                                   # cargo logo abaixo do nome
    if (len(paragraphs) >= 2 and looks_like_signer_name(paragraphs[-2]) and SIGNER_CARGO.match(last)
            and len(last) <= 90 and not OPEN_ENDING.search(last) and looks_like_signer_name(text)):
        return True                                   # próximo signatário, depois de nome + cargo
    return bool(PLACE_DATE.match(last) and looks_like_signer_name(text))   # "Belém, 1º de ... de 2026"


_VOCABULARY = None   # palavras do diário em processamento (montado em extract_act)


def reflow_lines(lines: list[dict]) -> list[str]:
    """Une somente quebras físicas curtas dentro do mesmo parágrafo."""
    paragraphs: list[str] = []
    previous_bottom = 0.0
    for line in sorted(lines, key=lambda l: (l["top"], l["x0"])):
        text = line["text"].strip()
        if not text:
            continue
        if (len(paragraphs) == 1 and line["top"] - previous_bottom <= 12
                and continues_epigraph(paragraphs[0], text)):
            paragraphs[0] = join_line(paragraphs[0], text)   # "... Nº 024 -" + "SEFA.GS, DE ... 2026"; ".../" + "PRODEPA"
            previous_bottom = line["bottom"]
            continue
        new_start = bool(EPIGRAPH.match(text) or (EMENTA_START.match(text) and text[:1].isupper())
                         or PREAMBLE.match(text) or COMMAND.match(text) or ARTICLE.match(text))
        if paragraphs and EPIGRAPH.match(paragraphs[-1]) and re.search(r"\b\d{4}\s*$", paragraphs[-1]):
            new_start = True
        if CARGO.match(text) or NAME.fullmatch(text) or looks_uppercase(text):
            new_start = True
        if paragraphs and signature_break(paragraphs, text):
            new_start = True
        if ITEM_NUMBER.match(text):
            new_start = True
        if paragraphs and ANNEX.match(paragraphs[-1]) and not re.search(r"\s[-–]\s", paragraphs[-1]):
            new_start = True      # "ANEXO ÚNICO" sozinho: a linha seguinte é o conteúdo do anexo
        if (new_start and paragraphs and line["top"] - previous_bottom <= 6
                and MID_SENTENCE.search(paragraphs[-1])
                and not any(rx.match(text) for rx in (ARTICLE, COMMAND, ANNEX, HEADING, EMENTA_START))
                and not ITEM_NUMBER.match(text)
                and not (PREAMBLE.match(text) and text[:1].isupper())):   # "..., e" / "Considerando ..." (IN 024)
            new_start = False
        if (new_start and paragraphs and line["top"] - previous_bottom <= 6 and MID_SENTENCE.search(paragraphs[-1])
                and not re.search(r";\s*(?:e|ou)\s*$", paragraphs[-1])
                and re.match(r"(?:Art|ART|§|Par[aá]grafo|PAR[AÁ]GRAFO)", text)):
            new_start = False     # "...conforme previsto no" / "§ 1º O ..." nunca abre dispositivo no meio da frase
        if paragraphs and re.search(r"\b(?:do|da|dos|das)\s*$", paragraphs[-1], re.I) and EPIGRAPH.match(text):
            # "... e do / Decreto Estadual nº ..." continua a frase; o nome
            # de uma lei citada não é a epígrafe de um segundo ato.
            new_start = False
        if (paragraphs and re.search(r"\b(?:desta|deste)\s*$", paragraphs[-1], re.I)
                and re.fullmatch(r"(?:LEI|DECRETO|PORTARIA|RESOLUÇÃO)\.", text, re.I)):
            new_start = False
        too_far = bool(paragraphs and line["top"] - previous_bottom > 6)
        if (paragraphs and not new_start and not too_far
                and paragraphs[-1].endswith("\xad") and text[:1].islower()):
            paragraphs[-1] = paragraphs[-1][:-1] + text
        elif paragraphs and not new_start and not too_far and (
                not paragraphs[-1].rstrip().endswith(('.', ':', ';'))
                or (ENDS_WITH_ABBREVIATION.search(paragraphs[-1].rstrip()) and text[:1].isdigit())):
            paragraphs[-1] = join_line(paragraphs[-1], text, _VOCABULARY)
        else:
            paragraphs.append(text)
        previous_bottom = line["bottom"]
    return paragraphs


def pdf_issue_date(pdf_path: Path, document: Any) -> tuple[str, str, str] | None:
    """Usa apenas a data explicitamente impressa no cabeçalho da primeira página."""
    for line in pdf_lines(pdf_path, 1, document)[:8]:
        match = PORTUGUESE_DATE.search(line["text"])
        if match and line["top"] < 140:
            day, month, year = match.groups()
            return year, MONTH_NUMBERS[month.lower()], f"{int(day):02d}"
    return None


def publication_note(pdf_path: Path, epigraph: str, source_name: str | None,
                     issue_date: tuple[str, str, str] | None = None, extra_edition: bool = False) -> str | None:
    found = PDF_DATE.search(pdf_path.name)
    if not source_name or (not found and not issue_date):
        return None
    year, month, day = found.groups() if found else issue_date
    feminine = normalized(epigraph).startswith(("PORTARIA", "RESOLUCAO", "INSTRUCAO", "LEI"))
    extra = " Ed. Extra." if extra_edition or "EXTRA" in pdf_path.name.upper() else ""
    date = f"{day}.{month}.{year[-2:]}" + ("." if extra else "")   # "de 02.10.26. Ed. Extra." (todos os diários)
    return f"{'Publicada' if feminine else 'Publicado'} no {source_name} de {date}{extra or '.'}"


def is_extra_edition(pdf_path: Path, document: Any) -> bool:
    """Edição extra indicada no próprio cabeçalho: "EDIÇÃO EXTRA" ou, no DOU, número com letra
    ("Ano CLXIV Nº 186-A")."""
    for line in pdf_lines(pdf_path, 1, document)[:12]:
        text = line["text"].strip()
        if re.search(r"\bEDI[ÇC][ÃA]O\s+EXTRA\b", text, re.I) or re.match(r"Ano\s+[IVXLCDM]+\s+N[º°o]\s*\d+-[A-Z]\b", text):
            return True
    return False


def detect_publication(pdf: Any) -> str | None:
    """Diário de origem, lido primeiro no cabeçalho da 1ª página (quarto superior: título, número da
    edição, ISSN). Assim, um ato do DOE que cita "publicado no Diário Oficial da União" não faz o
    diário inteiro ser tratado como DOU. Sem marca no cabeçalho, procura na página inteira,
    começando pelos diários do Pará."""
    first = pdf.pages[0]
    x0, top, x1, bottom = first.bbox            # pode não começar exatamente em zero (DOU)
    header = normalized(first.crop((x0, top, x1, top + (bottom - top) * .25)).extract_text() or "")
    if "ISSN 1677-7042" in header or "IMPRENSA NACIONAL" in header or "DIARIO OFICIAL DA UNIAO" in header:
        return "DOU"
    if "DO-E/SEFA" in header or "DIARIO ELETRONICO DA SEFA" in header or "DIARIO OFICIAL ELETRONICO DA SECRETARIA" in header:
        return "DO-e/SEFA"
    if "DIARIO OFICIAL" in header and ("ESTADO DO PARA" in header or " IOE" in header):
        return "DOE (PA)"
    page = normalized(first.extract_text() or "")
    if "DO-E/SEFA" in page or "DIARIO ELETRONICO DA SEFA" in page:
        return "DO-e/SEFA"
    if "DIARIO OFICIAL" in page and "ESTADO DO PARA" in page:
        return "DOE (PA)"
    if "DIARIO OFICIAL DA UNIAO" in page or "ISSN 1677-7042" in page:
        return "DOU"
    if "CGIBS" in page:
        return "CGIBS"
    return None


def normative_blocks(lines: list[dict], page: int, publication: str | None = None) -> list[ParagraphBlock]:
    texts = reflow_lines(lines)
    if publication and texts:
        texts.insert(1, publication)
    items = classify("\n".join(texts))
    blocks = []
    for index, item in enumerate(items):
        if publication and item.text == publication:
            item.attention.append("Linha criada a partir da identificação e da data da edição do PDF; confirme no diário.")
        if index + 1 < len(items) and looks_like_person_name(item.text) and SIGNER_CARGO.match(items[index + 1].text):
            item.role, item.reason = "autoridade", "Assinatura seguida de cargo"
        if index and SIGNER_CARGO.match(item.text) and items[index - 1].role == "autoridade":
            item.role, item.reason = "cargo", "Cargo após assinatura"
        blocks.append(ParagraphBlock(item.text, item.role, page, "norma", item, list(item.attention)))
    return blocks


def logo_block(page: Any, upper: float, lower: float, page_number: int) -> ImageBlock | None:
    candidates = [im for im in page.images if upper < im["top"] < im["bottom"] < lower and im["width"] < 120]
    if not candidates:
        return None
    im = min(candidates, key=lambda x: abs((x["x0"] + x["x1"]) / 2 - page.width / 2))
    box = (max(0, im["x0"] - 2), im["top"] - 2, min(page.width, im["x1"] + 2), im["bottom"] + 2)
    raster = page.crop(box).to_image(resolution=300).original
    stream = io.BytesIO()
    raster.save(stream, format="PNG")
    return ImageBlock(stream.getvalue(), page_number)


def after_table_blocks(lines: list[dict], top: float, bottom: float, page_number: int) -> list[ParagraphBlock]:
    lines = [line for line in lines if top < line["top"] < bottom]
    first_article = next((i for i, line in enumerate(lines) if ARTICLE.match(line["text"].strip())), None)
    if first_article is not None:
        before = after_table_blocks(lines[:first_article], top, bottom, page_number)
        signer = next((i for i in range(first_article + 1, len(lines) - 1)
                       if looks_like_person_name(lines[i]["text"].strip())
                       and CARGO.match(lines[i + 1]["text"].strip())), len(lines))
        return (before + normative_blocks(lines[first_article:signer], page_number) +
                after_table_blocks(lines[signer:], top, bottom, page_number))
    result: list[ParagraphBlock] = []
    previous_bottom = top
    for line in lines:
        text = line["text"].strip()
        if not text or PROTOCOL.match(text):
            continue
        if (result and result[-1].kind == "footnote" and line["top"] - previous_bottom < 8
                and not FOOTNOTE.match(text) and not CARGO.match(text)
                and not re.match(r"^(?:Belém|PAL[AÁ]CIO|GABINETE)\b", text, re.I)):
            result[-1].text += "\n" + text
            previous_bottom = line["bottom"]
            continue
        if FOOTNOTE.match(text):
            result.append(ParagraphBlock(text, "texto", page_number, "footnote"))
        elif re.match(r"^(?:Belém|PAL[AÁ]CIO|GABINETE)\b", text, re.I):
            result.append(ParagraphBlock(text, "texto", page_number, "date"))
        elif CARGO.match(text):
            result.append(ParagraphBlock(text, "cargo", page_number, "signer"))
        else:
            if len(text.split()) >= 2 and not re.search(r"[;:]$", text):
                result.append(ParagraphBlock(text, "autoridade", page_number, "signer"))
            else:
                result.append(ParagraphBlock(text, "texto", page_number, "footnote", attention=["Trecho após a tabela requer conferência."]))
        previous_bottom = line["bottom"]
    # Nas assinaturas em duas colunas, organiza nome e cargo da esquerda antes
    # do nome e cargo da direita, preservando a associação entre cada par.
    for offset in range(len(result) - 3):
        subset = result[offset:offset + 4]
        if [b.role for b in subset] == ["autoridade", "autoridade", "cargo", "cargo"]:
            result[offset + 1], result[offset + 2] = result[offset + 2], result[offset + 1]
    return result


def has_two_columns(page: Any, lines: list[dict], tables: list[Any]) -> bool:
    """Identifica colunas, inclusive antes de uma tabela larga de outro ato."""
    middle = page.width / 2
    wide = [t for t in tables if t.bbox[0] < middle - 15 and t.bbox[2] > middle + 15]
    if wide:
        # Uma tabela de outra publicação, abaixo dos protocolos das duas
        # colunas, não muda a leitura do trecho anterior da página.
        first_wide_top = min(t.bbox[1] for t in wide)
        earlier_tables = [t for t in tables if t.bbox[3] < first_wide_top]
        earlier_lines = [line for line in lines if line["top"] < first_wide_top]
        ended_sides = {(line["x0"] + line["x1"]) / 2 < middle
                       for line in earlier_lines if PROTOCOL.match(line["text"].strip())}
        if not earlier_tables or ended_sides != {True, False}:
            return False
        lines, tables = earlier_lines, earlier_tables
    left = [line for line in lines if line["x1"] < middle + 12]
    right = [line for line in lines if line["x0"] > middle - 12]
    # Uma tabela curta em cada margem pode não produzir as 12 linhas de texto
    # exigidas pelo detector geral, mas ainda indicar uma virada entre colunas.
    left_end = any(t.bbox[2] < middle + 12 and
                   t.bbox[3] > page.height - 75 for t in tables)
    right_start = any(t.bbox[0] > middle - 12 and
                      t.bbox[1] < 150 for t in tables)
    if left_end and right_start and len(left) >= 2 and len(right) >= 2:
        return True
    return len(left) >= 12 and len(right) >= 12 and \
        min(max(l["top"] for l in left), max(l["top"] for l in right)) - \
        max(min(l["top"] for l in left), min(l["top"] for l in right)) > 150


class TextTable:
    """Recorte de uma tabela encontrada por alinhamento, sem inventar células."""

    strategy = "alinhamento"

    def __init__(self, rows: list[Any], extracted: list[list[str | None]]):
        used = [i for i in range(len(extracted[0])) if any(
            line[i] and line[i].strip() for line in extracted)]
        self.rows = [SimpleNamespace(cells=[row.cells[i] for i in used]) for row in rows]
        self._extracted = [[line[i] for i in used] for line in extracted]
        boxes = [cell for row in self.rows for cell in row.cells if cell]
        self.bbox = (min(b[0] for b in boxes), min(b[1] for b in boxes),
                     max(b[2] for b in boxes), max(b[3] for b in boxes))

    def extract(self) -> list[list[str | None]]:
        return self._extracted


def _numeric_cells(row: list[str | None]) -> int:
    return sum(bool(NUMERIC_CELL.fullmatch(line.strip()))
               for cell in row[1:] if cell for line in cell.splitlines() if line.strip())


def find_text_tables(page: Any, start_y: float, stop_y: float, grid_tables: list[Any]) -> list[TextTable]:
    """Tenta uma segunda estratégia apenas em grupos tabulares reconhecíveis."""
    if stop_y - start_y < 35:
        return []
    region = page.crop((0, max(0, start_y), page.width, min(page.height, stop_y)))
    settings = {"vertical_strategy": "text", "horizontal_strategy": "text",
                "min_words_vertical": 2, "min_words_horizontal": 1}
    result = []
    for candidate in region.find_tables(table_settings=settings):
        raw = candidate.extract()
        if not raw or len(raw[0]) < 2:
            continue
        data = [i for i, row in enumerate(raw) if row[0] and _numeric_cells(row)]
        if len(data) < 2:
            continue
        header = next((i for i in range(data[0] - 1, -1, -1)
                       if sum(bool(v and v.strip()) for v in raw[i]) >= 2), data[0])
        indices = [i for i in range(header, data[-1] + 1) if any(v and v.strip() for v in raw[i])]
        table = TextTable([candidate.rows[i] for i in indices], [raw[i] for i in indices])
        if table.bbox[2] - table.bbox[0] < 180:
            continue
        if any(max(0, min(table.bbox[2], other.bbox[2]) - max(table.bbox[0], other.bbox[0])) *
               max(0, min(table.bbox[3], other.bbox[3]) - max(table.bbox[1], other.bbox[1])) >
               (table.bbox[2] - table.bbox[0]) * (table.bbox[3] - table.bbox[1]) * .15
               for other in grid_tables):
            continue
        result.append(table)
    return result


def unrecognized_table_lines(words: list[dict], tables: list[Any], start_y: float, stop_y: float) -> int:
    """Conta linhas de valores alinhados que ficaram fora de qualquer tabela."""
    outside = [w for w in words if start_y <= w["top"] < stop_y and not any(
        t.bbox[0] - 2 <= (w["x0"] + w["x1"]) / 2 <= t.bbox[2] + 2 and
        t.bbox[1] - 2 <= w["top"] <= t.bbox[3] + 2 for t in tables)]
    likely = []
    for top, row in grouped_words(outside):
        fields = [w for w in row if MONEY.fullmatch(w["text"]) or PERCENT.fullmatch(w["text"])]
        if fields and any(w["x0"] + 65 < fields[0]["x0"] for w in row):
            likely.append((top, fields[0]["x0"], len(fields)))
    return sum(1 for i, (top, x, _) in enumerate(likely[1:], 1)
               if top - likely[i - 1][0] <= 38 and abs(x - likely[i - 1][1]) <= 12)


def line_midpoint(line: dict) -> float:
    """PyMuPDF e pdfplumber podem discordar na borda de uma tabela."""
    return (line["top"] + line["bottom"]) / 2


def line_inside_table(line: dict, table: Any) -> bool:
    mid_x = (line["x0"] + line["x1"]) / 2
    x0, top, x1, bottom = table.bbox
    return x0 - 1 <= mid_x <= x1 + 1 and top - 1 <= line_midpoint(line) <= bottom + 1


def append_region(page: Any, pdf_path: Path, extraction: Extraction, text_lines: list[dict],
                  page_tables: list[Any], words: list[dict], start_y: float, stop_y: float,
                  page_number: int, publication: str | None, split_columns: bool) -> None:
    tables = [table for table in page_tables if table.bbox[1] >= start_y
              and table.bbox[3] < stop_y + 1 and table.bbox[2] - table.bbox[0] > 200]
    if not split_columns:
        tables.extend(find_text_tables(page, start_y, stop_y, tables))
    tables.sort(key=lambda t: t.bbox[1])
    if unrecognized_table_lines(words, tables, start_y, stop_y):
        extraction.blocking_issues.append(
            f"Página {page_number}: há valores alinhados fora das tabelas reconhecidas. "
            "Use o PDF original para reconstruir este anexo antes de exportar.")
    if not tables:
        region = [line for line in text_lines if start_y <= line["top"] < stop_y]
        extraction.blocks.extend(normative_blocks(region, page_number, publication))
        return
    outside_lines = [line for line in text_lines if start_y <= line["top"] < stop_y
                     and not any(line_inside_table(line, table) for table in tables)]
    first_top = tables[0].bbox[1]
    images = [im for im in page.images if start_y < im["top"] < im["bottom"] < first_top]
    logo_top = min((im["top"] for im in images), default=first_top)
    all_lines = [line for line in outside_lines if line_midpoint(line) < first_top]
    annex_heading = min((line["top"] for line in all_lines
                         if ANNEX_HEAD.match(line["text"].strip())), default=first_top)
    normative_end = min(logo_top, annex_heading)
    before = [line for line in all_lines if line["top"] + 2 < normative_end]
    extraction.blocks.extend(normative_blocks(before, page_number, publication))
    logo = logo_block(page, normative_end - 2, first_top, page_number)
    if logo:
        extraction.blocks.append(logo)
    # Entre o título do anexo e a primeira tabela, todas as linhas entram (antes só as do tipo "ANEXO ...",
    # "GOVERNO DO ESTADO", "DEMONSTRATIVO", "R$ 1,00"; o nome do anexo, como "SERVIDORES REMOVIDOS –
    # AUDITOR", era descartado). A conferência abaixo bloqueia se algum caractere não chegar ao Word.
    heading_lines = [line for line in all_lines if line["top"] >= normative_end and line["text"].strip()]
    heading_blocks = []
    for line in heading_lines:
        text = line["text"].strip()
        kind = "unit" if text.startswith("R$") else "annex_head"
        heading_blocks.append(ParagraphBlock(text, "anexo" if ANNEX.match(text) else "texto", page_number, kind))
    extraction.blocks.extend(heading_blocks)
    if logo is None and character_signature("".join(l["text"] for l in heading_lines)) != character_signature(
            "".join(b.text for b in heading_blocks)):
        extraction.blocking_issues.append(f"Página {page_number}: há texto antes da tabela que não chegou ao Word. "
                                          "Confira o anexo no PDF.")
    for table_index, table in enumerate(tables):
        if table_index:
            previous_bottom = tables[table_index - 1].bbox[3]
            gap_lines = [line for line in outside_lines
                         if previous_bottom < line_midpoint(line) < table.bbox[1]]
            if any(ARTICLE.match(line["text"].strip()) for line in gap_lines):
                gap_blocks = normative_blocks(gap_lines, page_number)
            else:
                gap_blocks = [ParagraphBlock(line["text"].strip(),
                                            "anexo" if ANNEX.match(line["text"].strip()) else "texto",
                                            page_number, "annex_head")
                              for line in gap_lines if line["text"].strip()]
            if gap_lines:
                if character_signature("".join(line["text"] for line in gap_lines)) != character_signature(
                        "".join(block.text for block in gap_blocks)):
                    extraction.blocking_issues.append(
                        f"Página {page_number}: há texto entre tabelas que não chegou ao Word. "
                        "Confira os artigos e títulos dos anexos antes de exportar.")
                else:
                    extraction.intertable_segments_checked += 1
            extraction.blocks.extend(gap_blocks)
        block = extract_table(page, table, words, page_number)
        if split_columns:
            column = "esquerda" if (table.bbox[0] + table.bbox[2]) / 2 < page.width / 2 else "direita"
            block.source_parts = [TablePagePart(page_number, table.bbox, 0,
                                                len(block.matrix), column)]
        extraction.blocks.append(block)
    after = [line for line in outside_lines if line_midpoint(line) > tables[-1].bbox[3]]
    if after:
        extraction.blocks.extend(after_table_blocks(after, min(line["top"] for line in after) - 1,
                                                  stop_y, page_number))


def _query_key(value: str) -> str:
    return re.sub(r"\s*([/.-])\s*", r"\1", normalized(value))


_ACT_TYPES = (("INSTRUCAO NORMATIVA", ("INSTRUCAO NORMATIVA", "IN")), ("LEI COMPLEMENTAR", ("LEI COMPLEMENTAR", "LC")),
              ("LEI", ("LEI",)), ("DECRETO", ("DECRETO",)), ("PORTARIA", ("PORTARIA",)),
              ("RESOLUCAO", ("RESOLUCAO",)), ("ATO COTEPE", ("ATO COTEPE",)), ("CONVENIO", ("CONVENIO",)),
              ("AJUSTE", ("AJUSTE",)), ("PROTOCOLO", ("PROTOCOLO",)),
              ("ATO TECNICO", ("ATO TECNICO",)), ("ATO DECLARATORIO", ("ATO DECLARATORIO", "ADE")),
              ("MEDIDA PROVISORIA", ("MEDIDA PROVISORIA", "MP")), ("DESPACHO", ("DESPACHO",)))


def _act_parts(text: str, epigraph: bool) -> tuple[str | None, str | None, str | None]:
    """(tipo, número sem pontos nem zeros à esquerda, ano) de uma epígrafe ou do que foi digitado."""
    key = normalized(text)
    kind = None
    for name, spellings in _ACT_TYPES:
        if any(re.search(rf"(?<![A-Z]){re.escape(word)}(?![A-Z])", key[:40] if epigraph else key)
               for word in spellings):
            kind = name
            break
    found = (re.search(r"\bN\s*[O.]?\s*(\d[\d.]*)", key) if epigraph else None) or re.search(r"(\d[\d.]*)", key)
    if not found:
        return kind, None, None
    number = found.group(1).replace(".", "").lstrip("0") or "0"
    rest = key[found.end():]
    year = re.match(r"\s*/\s*(\d{4}|\d{2})(?!\d)", rest)
    if year:
        year_value = year.group(1)
    else:
        years = re.findall(r"(?<!\d)((?:19|20)\d{2})(?!\d)", rest)
        year_value = years[-1] if years else None
    return kind, number, year_value


def _same_act(query: str, epigraph: str) -> bool:
    q_kind, q_number, q_year = _act_parts(query, epigraph=False)
    e_kind, e_number, e_year = _act_parts(epigraph, epigraph=True)
    if not q_number or q_number != e_number:
        return False
    if q_kind and e_kind and not e_kind.startswith(q_kind):
        return False
    return not q_year or not e_year or q_year[-2:] == e_year[-2:]


def _query_hit(query: str, epigraph: str) -> bool:
    """O identificador digitado aparece na epígrafe sem estar colado a outros algarismos:
    "4/2026" casa com "Nº 04/2026" (zeros à esquerda), mas não com "Nº 1.884/2026"."""
    q, e = _query_key(query), _query_key(epigraph)
    start = e.find(q) if q else -1
    while start >= 0:
        before = re.search(r"[\d.]*$", e[:start]).group(0).replace(".", "")
        after = e[start + len(q):start + len(q) + 1]
        if ((not q[0].isdigit() or not before.strip("0"))
                and (not q[-1].isdigit() or not after.isdigit())):
            return True
        start = e.find(q, start + 1)
    return False


CGIBS_SITE = re.compile(r"S[IÍ]TIO ELETR[OÔ]NICO DO (?:COMIT[EÊ] GESTOR DO IMPOSTO SOBRE BENS E SERVI[CÇ]OS|CGIBS)", re.I)


def add_cgibs_availability(extraction: Extraction) -> None:
    """DOU: quando o ato diz que entra em vigor com a publicação no DOU "e no sítio eletrônico do Comitê
    Gestor do IBS", acrescenta "Disponibilizado no site do CGIBS em dd.mm.aa." logo depois da linha de
    publicação, com a data do DOU (ATC nº 6/2026, corrigido à mão). A data deve ser confirmada no site."""
    blocks = extraction.blocks
    publication = next((i for i, b in enumerate(blocks) if isinstance(b, ParagraphBlock) and b.role == "publicacao"), None)
    if publication is None or not any(isinstance(b, ParagraphBlock) and CGIBS_SITE.search(b.text) for b in blocks):
        return
    found = re.search(r"\bde (\d{2}\.\d{2}\.\d{2})\b", blocks[publication].text)
    if not found:
        return
    verb = "Disponibilizada" if blocks[publication].text.startswith("Publicada") else "Disponibilizado"
    text = f"{verb} no site do CGIBS em {found.group(1)}."
    attention = ["Linha criada porque o ato também é publicado no site do CGIBS; a data é a do DOU: confirme no site."]
    item = Item(text, 0, 0, "publicacao", "Disponibilização no site do CGIBS", list(attention))
    blocks.insert(publication + 1, ParagraphBlock(text, "publicacao", blocks[publication].page, "norma", item, attention))


def line_above(lines: list[dict], index: int) -> dict | None:
    """Linha imediatamente acima, na mesma coluna (faixas horizontais que se sobrepõem)."""
    current = lines[index]
    candidates = [l for l in lines[:index] if l["top"] < current["top"] - 1
                  and l["x0"] < current["x1"] and current["x0"] < l["x1"]]
    return max(candidates, key=lambda l: l["top"], default=None)


def _heading_like(text: str) -> bool:
    text = text.strip()
    return bool(text) and len(text) <= 120 and not text.endswith((".", ";", ":", ",")) and text[:1].isupper()


def trim_trailing_headings(extraction: Extraction) -> None:
    """DOU: entre o fim de um ato e a epígrafe do seguinte vêm os títulos do órgão ("Ministério da
    Fazenda", "SECRETARIA EXECUTIVA"). São descartados só quando vêm logo após uma assinatura
    reconhecida ou ao fim de um ANEXO que segue a assinatura; uma assinatura não reconhecida
    (nome e cargo como texto) nunca é apagada."""
    blocks = extraction.blocks
    start = len(blocks)
    while (start > 1 and isinstance(blocks[start - 1], ParagraphBlock) and blocks[start - 1].role == "texto"
           and _heading_like(blocks[start - 1].text)):
        start -= 1
    if start == len(blocks):
        return
    before = blocks[start - 1]
    roles = [b.role if isinstance(b, ParagraphBlock) else "tabela" for b in blocks[:start]]
    after_signature = isinstance(before, ParagraphBlock) and before.role == "cargo"
    annex_after_signature = ("cargo" in roles and "anexo" in roles[len(roles) - 1 - roles[::-1].index("cargo"):]
                             and (not isinstance(before, ParagraphBlock)
                                  or before.text.rstrip().endswith((".", ";", ":"))))
    if after_signature or annex_after_signature:
        tail = blocks[start:]
        del blocks[start:]
        extraction.issues.append("Títulos de órgão após o fim do ato foram deixados de fora ("
                                 + "; ".join(b.text for b in tail)[:120] + "); confira.")


def epigraph_text(lines: list[dict], index: int) -> str:
    """Epígrafe completa a partir da linha index, reunindo a continuação na linha seguinte."""
    text = lines[index]["text"].strip()
    for following in lines[index + 1:index + 3]:
        if not continues_epigraph(text, following["text"].strip()):
            break
        text += " " + following["text"].strip()
    return text


def merge_annex_subtitles(extraction: Extraction) -> None:
    """Depois de "ANEXO I", as linhas em maiúsculas do nome do anexo ("SERVIDORES REMOVIDOS – AUDITOR" /
    "FISCAL DE RECEITAS ESTADUAIS – AFRE") formam um só parágrafo, com o estilo de subtítulo. Antes saíam
    partidas e, entre duas tabelas, sem estilo da SEFA (Portaria nº 522/2026-SEFA.GS)."""
    blocks, merged, index = extraction.blocks, [], 0
    while index < len(blocks):
        block = blocks[index]
        merged.append(block)
        index += 1
        if not (isinstance(block, ParagraphBlock) and block.role == "anexo"):
            continue
        parts = []
        while (index < len(blocks) and len(parts) < 4 and isinstance(blocks[index], ParagraphBlock)
               and blocks[index].role == "texto" and blocks[index].kind in ("norma", "annex_head")
               and looks_uppercase(blocks[index].text) and len(blocks[index].text) <= 120
               and not blocks[index].text.rstrip().endswith((".", ";", ":"))
               and not ANNEX.match(blocks[index].text)
               # nome de anexo é frase sem algarismos ("SERVIDORES REMOVIDOS – AUDITOR"); cabeçalho de coluna de
               # tabela não lida ("ÁREA/UNIDADE", "3º QUADRIMESTRE - 2026") não entra
               and not re.search(r"\d", blocks[index].text)
               and (len(blocks[index].text) >= 8 if parts
                    else len(blocks[index].text) >= 15 and len(blocks[index].text.split()) >= 2)):
            parts.append(blocks[index])
            index += 1
        if parts:
            text = parts[0].text
            for part in parts[1:]:
                text = join_line(text, part.text, _VOCABULARY)
            item = Item(text, 0, 0, "subtitulo", "Nome do anexo")
            merged.append(ParagraphBlock(text, "subtitulo", parts[0].page, "norma", item, []))
    extraction.blocks[:] = merged


def merge_across_regions(extraction: Extraction) -> None:
    """Parágrafo que vira de coluna ou de página: termina sem pontuação e o seguinte começa com minúscula.
    Ex.: "... e a empre-" / "sa AVANÇO FACILITIES ..." (Portaria nº 0628/2026-CRG) e o cargo "Presidente do
    Instituto ... do Es-" / "tado do Pará" (Portaria nº 649/2026-IGEPPS), ambos no DOE de 02.10.2026.
    Dentro de uma mesma coluna isso já é feito ao juntar as linhas."""
    blocks, merged = extraction.blocks, []
    for block in blocks:
        previous = merged[-1] if merged else None
        if (isinstance(previous, ParagraphBlock) and isinstance(block, ParagraphBlock)
                and previous.kind == block.kind and block.kind in ("norma", "signer")
                and previous.role in ("texto", "cargo", "ementa")
                and block.role not in ("epigrafe", "publicacao", "anexo", "ordem")    # "resolvem:" fica separado
                and block.text[:1].islower() and not re.match(r"[a-z]\s*\)", block.text)
                and not previous.text.rstrip("\xad ").endswith((".", ":", ";", "!", "?"))):
            head = previous.text.rstrip()
            text = head[:-1] + block.text if head.endswith("\xad") else join_line(head, block.text, _VOCABULARY)
            previous.text = text
            previous.attention = previous.attention + [a for a in block.attention if a not in previous.attention]
            if previous.item:
                previous.item.text = text
            continue
        merged.append(block)
    extraction.blocks[:] = merged


def link_split_signatures(extraction: Extraction) -> None:
    """Nome no pé de uma coluna (ou página) e cargo no topo da seguinte: cada região é
    classificada em separado, então a regra "nome seguido de cargo" é refeita aqui."""
    paragraphs = [b for b in extraction.blocks]
    for name, cargo in zip(paragraphs, paragraphs[1:]):
        if not (isinstance(name, ParagraphBlock) and isinstance(cargo, ParagraphBlock)
                and name.kind == cargo.kind == "norma" and name.role != "autoridade"
                and looks_like_person_name(name.text) and SIGNER_CARGO.match(cargo.text)):
            continue
        for block, role, reason in ((name, "autoridade", "Assinatura seguida de cargo"),
                                    (cargo, "cargo", "Cargo após assinatura")):
            block.role = role
            block.attention = []   # avisos de quando nome e cargo foram classificados separados
            if block.item:
                block.item.role, block.item.reason = role, reason
                block.item.attention = list(block.attention)


def _row_key(row: list[str]) -> tuple[str, ...]:
    return tuple(" ".join(value.split()).upper() for value in row)


def merge_spanning_tables(extraction: Extraction, pdf: Any) -> None:
    """Une fragmentos contíguos ao virar coluna ou página, preservando sua origem."""
    index = 0
    while index + 1 < len(extraction.blocks):
        first, second = extraction.blocks[index:index + 2]
        if not isinstance(first, TableBlock) or not isinstance(second, TableBlock):
            index += 1
            continue
        previous_part = table_page_parts(first)[-1]
        next_part = table_page_parts(second)[0]
        widths_a, widths_b = first.column_widths, second.column_widths
        same_columns = len(widths_a) == len(widths_b) and len(widths_a) > 1
        same_widths = same_columns and all(a > 0 and b > 0 and
            .75 <= (a / sum(widths_a)) / (b / sum(widths_b)) <= 1.25
            for a, b in zip(widths_a, widths_b))
        closing_row = "".join(first.matrix[-1]).strip().endswith(('”', '"'))
        page_continuation = (next_part.page == previous_part.page + 1
                             and previous_part.bbox[3] > pdf.pages[previous_part.page - 1].height - 50
                             and next_part.bbox[1] < 100)
        page = pdf.pages[previous_part.page - 1]
        middle = page.width / 2 if hasattr(page, "width") else 0
        column_continuation = (next_part.page == previous_part.page
                               and previous_part.column == "esquerda"
                               and next_part.column == "direita"
                               and previous_part.bbox[2] < middle + 12
                               and next_part.bbox[0] > middle - 12
                               and previous_part.bbox[3] > page.height - 75
                               and next_part.bbox[1] < 150)
        same_column_continuation = (next_part.page == previous_part.page
                                    and bool(previous_part.column)
                                    and next_part.column == previous_part.column
                                    and 0 <= next_part.bbox[1] - previous_part.bbox[3] <= 7
                                    and abs(next_part.bbox[0] - previous_part.bbox[0]) < 3
                                    and abs(next_part.bbox[2] - previous_part.bbox[2]) < 3)
        last_label = first.matrix[-1][0].strip() if first.matrix[-1] else ""
        finished_total = bool(re.match(r"^TOTA(?:L|IS)\b", last_label, re.I))
        if not ((page_continuation or column_continuation or same_column_continuation)
                and same_widths and not closing_row and not finished_total
                and first.strategy == second.strategy):
            index += 1
            continue
        if (len(second.matrix) > 1 and first.matrix
                and _row_key(second.matrix[0]) == _row_key(first.matrix[0])
                and not any(r0 == 0 for r0, _, _, _ in second.merges)):
            # O diário repete o cabeçalho ao virar a página; o Word já o repete sozinho
            # (Portaria nº 522/2026-SEFA.GS, Anexo II, p. 3 -> 4).
            repeated = second.matrix.pop(0)
            if second.original_matrix:
                second.original_matrix.pop(0)
            second.source_signature = second.source_signature - matrix_signature([repeated])
            second.merges = [(r0 - 1, c0, r1 - 1, c1) for r0, c0, r1, c1 in second.merges]
            second.source_parts = [TablePagePart(p.page, p.bbox, max(0, p.first_row - 1), p.last_row - 1, p.column)
                                   for p in second.source_parts]
            extraction.issues.append(f"Página {next_part.page}: o cabeçalho repetido da tabela foi retirado "
                                     "(no Word, a linha de cabeçalho se repete sozinha a cada página).")
        offset = len(first.matrix)
        old_parts = table_page_parts(first)
        new_parts = table_page_parts(second)
        first.matrix.extend(second.matrix)
        first.original_matrix.extend(second.original_matrix)
        first.source_numbers.update(second.source_numbers)
        first.source_signature.update(second.source_signature)
        first.issues.extend(second.issues)
        first.merges.extend((r0 + offset, c0, r1 + offset, c1)
                            for r0, c0, r1, c1 in second.merges)
        first.source_parts = old_parts + [TablePagePart(part.page, part.bbox,
                                                       part.first_row + offset, part.last_row + offset,
                                                       part.column)
                                          for part in new_parts]
        first.column_widths = [(a * offset + b * len(second.matrix)) / len(first.matrix)
                               for a, b in zip(widths_a, widths_b)]
        first.reviewed = False
        del extraction.blocks[index + 1]


def extract_act(pdf_path: Path, query: str, max_pages: int = 12) -> Extraction:
    if not query.strip():
        raise ValueError("Informe um trecho único do título do ato, como 494/2026-SEFA.GS.")
    if not 1 <= max_pages <= 100:
        raise ValueError("O limite deve estar entre 1 e 100 páginas.")
    with pdfplumber.open(pdf_path) as pdf, fitz.open(str(pdf_path)) as fitz_pdf:
        source_name = detect_publication(pdf)
        issue_date = pdf_issue_date(pdf_path, fitz_pdf) if source_name else None
        extra_edition = bool(source_name) and is_extra_edition(pdf_path, fitz_pdf)
        hits: list[tuple[int, dict, str]] = []
        by_number: list[tuple[int, dict, str]] = []
        global _VOCABULARY
        all_text: list[str] = []
        for page_index, page in enumerate(pdf.pages):
            page_lines = pdf_lines(pdf_path, page_index + 1, fitz_pdf)
            all_text += [line["text"] for line in page_lines]
            for line_index, line in enumerate(page_lines):
                if not EPIGRAPH.match(line["text"].strip()):
                    continue
                above = line_above(page_lines, line_index)
                if above and MID_SENTENCE.search(above["text"]):
                    continue
                complete = epigraph_text(page_lines, line_index)
                if _query_hit(query, complete):
                    hits.append((page_index, line, complete))
                elif _same_act(query, complete):   # "024/2026" acha "... Nº 024 - / SEFA.GS, DE ... 2026"
                    by_number.append((page_index, line, complete))
        _VOCABULARY = vocabulary(all_text)
        chosen = hits or by_number
        first = chosen[0][:2] if chosen else None
        others = [c for c in chosen[1:] if normalized(c[2]) != normalized(chosen[0][2])] if chosen else []
        if first is None:
            raise ValueError("Não encontrei uma epígrafe com esse identificador no texto do PDF. PDF digitalizado ou título diferente requer seleção manual.")
        start_index, heading = first
        extraction = Extraction(pdf_path, query, [], [])
        if others:
            lista = "; ".join(f"p. {pi + 1}: {txt[:80]}" for pi, _, txt in others[:3])
            extraction.issues.append(
                f"O identificador digitado também corresponde a outro(s) ato(s) ({lista}"
                f"{'; ...' if len(others) > 3 else ''}). Foi usado o da p. {start_index + 1}; se não for o desejado, "
                "digite o identificador completo, com a sigla do órgão (ex.: 04/2026-PGE).")
        if "PREVISTA" in normalized(pdf_path.name):
            extraction.issues.append("O nome do PDF indica edição prevista; confirme a publicação final antes de usar o ato.")
        found_end = signed_off = ended_by_next = False
        for page_index in range(start_index, min(start_index + max_pages, len(pdf.pages))):
            page = pdf.pages[page_index]
            page_number = page_index + 1
            blocks_before_page = len(extraction.blocks)
            all_lines = pdf_lines(pdf_path, page_number, fitz_pdf)
            all_tables = page.find_tables()
            all_words = page.extract_words(x_tolerance=.5, y_tolerance=1.5, keep_blank_chars=False)
            split_columns = has_two_columns(page, all_lines, all_tables)
            if three_columns(all_lines, page.width):
                extraction.blocking_issues.append(
                    f"A página {page_number} tem três colunas de texto, formato que o programa ainda não lê "
                    "(só uma ou duas colunas). Envie este PDF para ajuste.")
            dou = source_name == "DOU"
            top_margin, bottom_margin = dou_margins(all_lines, page.height) if dou else (0.0, page.height)
            page_header = None if dou else page_top_margin(all_lines, page.height)
            middle = page.width / 2
            heading_left = (heading["x0"] + heading["x1"]) / 2 < middle
            columns = ([heading_left] if page_index == start_index and not heading_left else [True, False])
            if not split_columns:
                columns = [None]
            for column in columns:
                first_region = page_index == start_index and (column is None or column == heading_left)
                if column is None:
                    text_lines, page_tables, words = all_lines, all_tables, all_words
                else:
                    text_lines = [l for l in all_lines if
                                  (((l["x0"] + l["x1"]) / 2 < middle) == column)]
                    page_tables = [t for t in all_tables if
                                   (((t.bbox[0] + t.bbox[2]) / 2 < middle) == column)]
                    words = [w for w in all_words if
                             (((w["x0"] + w["x1"]) / 2 < middle) == column)]
                start_y = (heading["top"] - .5 if first_region else top_margin if dou and top_margin
                           else page_header if page_header is not None else 52.0)
                if not first_region and column is not None:
                    # Grades logo abaixo do cabeçalho do diário podem começar
                    # acima de 52 pt (p. ex. 42,6 pt no DOE de 01.10.2026).
                    early = [t.bbox[1] for t in page_tables
                             if 35 <= t.bbox[1] < start_y and
                             (t.bbox[2] < middle + 12 if column else t.bbox[0] > middle - 12)]
                    if early:
                        start_y = min(start_y, max(35.0, min(early) - 1))
                stop_y = min(page.height - 25, bottom_margin)
                if dou and not first_region and signed_off:
                    # DOU: depois da assinatura, o ato só continua se a região seguinte começar por ANEXO
                    first_line = next((l["text"].strip() for l in text_lines
                                       if start_y < l["top"] < stop_y and l["text"].strip()), "")
                    if not ANNEX.match(first_line):
                        found_end = True
                        extraction.issues.append("O fim foi identificado pela assinatura; confirme a delimitação.")
                        break
                previous_text = ""
                for line in text_lines:
                    if line["top"] <= start_y + (8 if first_region else 0):
                        continue
                    text = line["text"].strip()
                    before_text = previous_text
                    after_mid_sentence, previous_text = bool(MID_SENTENCE.search(previous_text)), text
                    if PROTOCOL.match(text):
                        stop_y = min(stop_y, line["top"] - 1)
                        extraction.protocol = text
                        found_end = True
                        break
                    # Nos 45 pt logo abaixo da epígrafe, uma linha que começa por "Decreto n. 2.703, ..." ou
                    # "PORTARIA nº 107/2026, ..." costuma ser citação; o ato seguinte só começa ali se o tipo estiver
                    # em maiúsculas e a linha anterior fechar a frase (Resoluções SENAI nº 09 a 20, DOE de 05.10.2026).
                    next_act_close = (first_region and EPIGRAPH.match(text) and line["top"] > heading["top"] + 2
                                      and text.split()[0].isupper() and len(text.split()[0]) > 2
                                      and before_text.rstrip().endswith((".", ";", ":")))
                    if (SECTION_HEAD.match(text) or next_act_close or
                            (EPIGRAPH.match(text) and line["top"] > start_y + (45 if first_region else 8)
                             and not after_mid_sentence)):
                        stop_y = min(stop_y, line["top"] - 1)
                        found_end = ended_by_next = True
                        extraction.issues.append("O fim foi identificado pelo início do ato ou seção seguinte; confira a delimitação.")
                        break
                start_block = len(extraction.blocks)
                publication = (publication_note(pdf_path, heading["text"], source_name, issue_date, extra_edition)
                               if first_region else None)
                append_region(page, pdf_path, extraction, text_lines, page_tables, words,
                              start_y, stop_y, page_number, publication, split_columns)
                if column is False and start_block and len(extraction.blocks) > start_block:
                    previous = extraction.blocks[start_block - 1]
                    current = extraction.blocks[start_block]
                    if (isinstance(previous, ParagraphBlock) and isinstance(current, ParagraphBlock)
                            and previous.kind == current.kind == "norma"
                            and previous.role == "ementa" and current.text[:1].islower()
                            and not previous.text.rstrip().endswith((".", ";", ":"))):
                        previous.text += " " + current.text
                        if previous.item:
                            previous.item.text = previous.text
                        del extraction.blocks[start_block]
                last_block = extraction.blocks[-1] if extraction.blocks else None
                signed_off = isinstance(last_block, ParagraphBlock) and last_block.role == "cargo"
                if found_end:
                    break
                if column is True and split_columns and extraction.blocks and not dou:
                    last = extraction.blocks[-1]
                    if (isinstance(last, ParagraphBlock) and last.role == "cargo"
                            and not last.text.rstrip().endswith(("-", "\xad"))     # "... do Es-" / "tado do Pará"
                            and not MID_SENTENCE.search(last.text)):
                        found_end = True
                        extraction.issues.append("O fim foi identificado pela assinatura na coluna esquerda; confirme a delimitação.")
                        break
            if page_index == start_index or len(extraction.blocks) > blocks_before_page:
                extraction.pages.append(page_number)
            if found_end:
                break
        if not found_end:
            extraction.issues.append("Não foi encontrado protocolo ou início do próximo ato. Confira onde o ato termina.")
            if start_index + max_pages < len(pdf.pages):
                extraction.blocking_issues.append(
                    f"A busca parou após {max_pages} páginas sem encontrar o fim do ato. "
                    "Aumente o limite antes de exportar.")
        merge_spanning_tables(extraction, pdf)
        merge_across_regions(extraction)
        merge_annex_subtitles(extraction)
        link_split_signatures(extraction)
        if source_name == "DOU" and ended_by_next:
            trim_trailing_headings(extraction)
        if source_name == "DOU":
            add_cgibs_availability(extraction)
        return extraction


def set_cell_border(cell: Any, edge: str, color: str = "595959", size: int = 4) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    name = qn("w:" + edge)
    rule = borders.find(name)
    if rule is None:
        rule = OxmlElement("w:" + edge)
        borders.append(rule)
    rule.set(qn("w:val"), "single")
    rule.set(qn("w:sz"), str(size))
    rule.set(qn("w:color"), color)


def set_cell_margins(cell: Any) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcPr.append(mar)
    for side in ("top", "left", "bottom", "right"):
        el = mar.find(qn("w:" + side))
        if el is None:
            el = OxmlElement("w:" + side)
            mar.append(el)
        el.set(qn("w:w"), str(4 if side in ("top", "bottom") else 22))
        el.set(qn("w:type"), "dxa")


CELL_LIST_START = re.compile(r"^(?:[a-zA-Z]\)|[IVXLCDM]+\s*[-–]|\d+\s*[.)–-]|[-–•·])\s")


def unwrap_cell(value: str) -> str:
    """Desfaz as quebras de linha do diário dentro de uma célula de texto ("CHARLES JOHNSON DA" /
    "SILVA ALCANTARA", "DIRETORIA DE" / "FISCALIZAÇÃO – DFI"): a coluna do Word tem outra largura.
    Linhas com algarismos, itens de lista e linhas depois de ":" ou ";" continuam separadas."""
    lines = value.split("\n")
    joined = [lines[0]]
    for line in lines[1:]:
        previous = joined[-1]
        if (re.search(r"[A-Za-zÀ-ÿ]", previous) and re.search(r"[A-Za-zÀ-ÿ]", line)
                and not re.search(r"\d", previous + line) and not CELL_LIST_START.match(line.strip())
                and not previous.rstrip().endswith((":", ";"))):
            joined[-1] = join_line(previous, line.strip())
        else:
            joined.append(line)
    return "\n".join(joined)


def formatted_table_matrix(block: TableBlock, normalize_municipalities: bool = True
                           ) -> tuple[list[list[str]], list[tuple[int, int, str, str, str]]]:
    """Separa a grafia de apresentação da matriz auditada do PDF.

    Somente colunas explicitamente identificadas como Município recebem iniciais
    maiúsculas. Códigos, produtos, siglas e demais colunas permanecem intactos.
    Tabelas com cabeçalhos mesclados exigem escolha manual do campo apropriado.
    """
    matrix = [[unwrap_cell(value) for value in row] for row in block.matrix]
    changes = []
    if not normalize_municipalities or not matrix or block.merges:
        return matrix, changes
    columns = [ci for ci, value in enumerate(matrix[0])
               if re.fullmatch(r"(?:NOME(?: D[OE]S?)? )?MUNICIPIOS?", normalized(value))]
    if not columns:
        return matrix, changes
    for ri, row in enumerate(matrix):
        for ci, before in enumerate(row):
            if ri == 0:
                # Não muda símbolos ou unidades de outros campos do cabeçalho.
                label = normalized(before)
                if ci in columns or label.startswith(("INDICE", "INDICADOR DE QUALIDADE")):
                    after = " ".join(before.split()).upper()
                    rule = "Cabeçalho em maiúsculas e sem quebra manual de linha"
                else:
                    continue
            elif ci in columns and re.fullmatch(r"[^\W\d_]+(?:[\s'’\-][^\W\d_]+)*", before):
                if normalized(before) in {"TOTAL", "TOTAIS"}:
                    continue
                after = re.sub(r"(^|[\s-])([^\W\d_])",
                               lambda m: m[1] + m[2].upper(), before.lower())
                rule = "Iniciais maiúsculas no nome do município, conforme o exemplo revisado"
            else:
                continue
            if before != after:
                if re.sub(r"\s", "", before).casefold() != re.sub(r"\s", "", after).casefold():
                    raise ValueError("Um ajuste de grafia tentou alterar o conteúdo da célula.")
                row[ci] = after
                changes.append((ri, ci, before, after, rule))
    return matrix, changes


def add_table(document: Document, block: TableBlock, normalize_municipalities: bool = True) -> None:
    matrix, _ = formatted_table_matrix(block, normalize_municipalities)
    table = document.add_table(rows=len(matrix), cols=len(matrix[0]))
    table.autofit = False
    section = document.sections[-1]
    usable_cm = Emu(section.page_width - section.left_margin - section.right_margin).cm - .1
    source_widths = block.column_widths[:]
    dense = len(source_widths) >= 7
    if dense:
        # O diário pode atribuir nove colunas de largura idêntica a nomes,
        # códigos longos e valores de quatro dígitos. Distribua o espaço
        # conforme a palavra mais longa de cada coluna, sem alterar células.
        for ci, original in enumerate(source_widths):
            longest = max((len(part) for row in block.matrix
                           for part in re.findall(r"[^\s]+", row[ci])
                           if re.search(r"\w", part)), default=7)
            source_widths[ci] = original * max(.68, min(1.55, (longest + 2) / 9))
    width_sum = sum(source_widths)
    widths = [Cm(usable_cm * w / width_sum) for w in source_widths]
    for col, width in zip(table.columns, widths):
        col.width = width
    header_rows = max((r1 + 1 for r0, _, r1, _ in block.merges if r0 == 0), default=1)
    dotted_header = bool(dense and all(re.fullmatch(r'[“”".\s…]+', value)
                                       for value in block.matrix[0]))
    if dotted_header:
        header_rows = 0
    if (len(block.matrix) > 1 and any(r0 == 0 and c0 == 0 and c1 == len(widths) - 1
                                      for r0, c0, _, c1 in block.merges)
            and not monetary_values([block.matrix[1]])):
        header_rows = max(header_rows, 2)
    for heading_row in table.rows[:header_rows]:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        heading_row._tr.get_or_add_trPr().append(header)
    available = paragraph_style_names(document)
    font_size = 8.0 if len(widths) <= 3 else (6.2 if len(widths) == 4 else
                                              (5.9 if len(widths) <= 6 else 6.0))
    for row_index, (row, values) in enumerate(zip(table.rows, matrix)):
        is_heading = row_index < header_rows or (not any(MONEY.fullmatch(v) for v in values[1:]) and re.match(r"^\d+\s*[-–]", values[0]))
        is_total = values[0].strip().upper() == "TOTAIS"
        is_major = bool(re.match(r"^\d+\s*[-–]", values[0]))
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for col_index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = widths[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for edge in ("left", "right", "top", "bottom"):
                set_cell_border(cell, edge, color="595959" if row_index == 0 and edge != "bottom" else "000000")
            cell.text = ""
            p = cell.paragraphs[0]
            compact = "".join(value.split())
            alignment = (WD_ALIGN_PARAGRAPH.CENTER if len(widths) <= 3 or is_heading else
                         WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else
                           WD_ALIGN_PARAGRAPH.CENTER if is_heading or
                           (dense and not re.fullmatch(r"(?:R\$)?[\d.,/%-]+", compact)) else
                           WD_ALIGN_PARAGRAPH.RIGHT)
            suffix = {WD_ALIGN_PARAGRAPH.LEFT: "esquerda", WD_ALIGN_PARAGRAPH.CENTER: "centro",
                      WD_ALIGN_PARAGRAPH.RIGHT: "direita"}[alignment]
            style = f"a50_tab_{suffix}_8" if len(widths) <= 3 else f"a60_tab_{suffix}_6"
            p.style = style if style in available else "Normal"
            if style not in available:
                p.alignment = alignment
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
            if len(widths) > 3:
                p.paragraph_format.line_spacing = .95
            if is_heading or (len(matrix) <= 6 and row_index < len(matrix) - 1):
                p.paragraph_format.keep_with_next = True
            for line_index, line in enumerate(value.split("\n")):
                if line_index:
                    p.add_run().add_break()
                run = p.add_run(line)
                if style not in available:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
                elif len(widths) > 3 and font_size != 6.0:
                    run.font.size = Pt(font_size)
                if is_heading or is_total or is_major:
                    run.bold = True
    for first_row, first_col, last_row, last_col in block.merges:
        merged = table.cell(first_row, first_col).merge(table.cell(last_row, last_col))
        if last_row > first_row:
            merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        for paragraph in list(merged.paragraphs)[1:]:
            if not paragraph.text.strip():
                paragraph._element.getparent().remove(paragraph._element)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build_word(extraction: Extraction, model_path: Path, output_path: Path, overwrite: bool = False,
               normalize_municipalities: bool = True) -> None:
    if output_path.resolve() == model_path.resolve():
        raise ValueError("Escolha um nome de saída diferente do documento modelo.")
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"O arquivo já existe: {output_path}")
    if any(table.issues and not table.reviewed for table in extraction.tables):
        raise ValueError("Há tabela com associação incerta. Confira-a com o PDF e marque a revisão antes de exportar.")
    model = Document(str(model_path))
    for block in extraction.paragraphs:
        if block.kind == "annex_head" and ANNEX.match(block.text):
            block.role = "anexo"
        if block.kind in {"norma", "signer", "date"} or block.role == "anexo":
            item = block.item or Item(block.text, 0, 0, block.role, "Importado do PDF")
            item.text, item.role = block.text, block.role
            apply_styles([item], model)
            block.item = item
    items = [b.item for b in extraction.paragraphs if b.item]
    spacer = ementa_spacer_before(items) if "a03_ementa" in paragraph_style_names(model) else None
    clear_model_contents(model)
    for block_index, block in enumerate(extraction.blocks):
        if isinstance(block, TableBlock):
            add_table(model, block, normalize_municipalities)
        elif isinstance(block, ImageBlock):
            p = model.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.add_run().add_picture(io.BytesIO(block.data), width=Inches(.46))
        elif block.kind in {"norma", "signer", "date"} or block.role == "anexo":
            if block.item is spacer:
                model.add_paragraph(style="a03_ementa")
            p = add_formatted_paragraph(model, block.item)
            if (block_index + 1 < len(extraction.blocks)
                    and isinstance(extraction.blocks[block_index + 1], TableBlock)):
                p.paragraph_format.keep_with_next = True
            if block.kind == "date":
                p.paragraph_format.space_before = Pt(7)
        else:
            p = model.add_paragraph()
            p.style = "Normal"
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if block.kind == "unit" else (
                WD_ALIGN_PARAGRAPH.CENTER if block.kind == "annex_head" else WD_ALIGN_PARAGRAPH.LEFT
            )
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            run = p.add_run(block.text)
            run.font.name = "Arial"
            run.font.size = Pt(7 if block.kind == "annex_head" else 6.4)
            run.bold = block.kind == "annex_head"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix="sefa_pdf_", suffix=".docx", dir=output_path.parent, delete=False) as temp:
        temporary = Path(temp.name)
    try:
        model.save(temporary)
        check = Document(temporary)
        if len(check.tables) != len(extraction.tables):
            raise RuntimeError("A saída não preservou todas as tabelas.")
        for expected, actual in zip(extraction.tables, check.tables):
            if len(expected.matrix) != len(actual.rows):
                raise RuntimeError("A saída perdeu linhas da tabela.")
            covered = {(ri, ci) for r0, c0, r1, c1 in expected.merges
                       for ri in range(r0, r1 + 1) for ci in range(c0, c1 + 1)
                       if (ri, ci) != (r0, c0)}
            presented, _ = formatted_table_matrix(expected, normalize_municipalities)
            for ri, row in enumerate(presented):
                for ci, value in enumerate(row):
                    if (ri, ci) not in covered and actual.cell(ri, ci).text != value:
                        raise RuntimeError("A saída divergiu de uma célula da tabela.")
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)


def write_pdf_report(extraction: Extraction, report_path: Path, normalize_municipalities: bool = True) -> None:
    sections = []
    for block in extraction.blocks:
        if isinstance(block, ParagraphBlock):
            attention = "<br>".join(html.escape(s) for s in block.attention)
            sections.append(f'<p><small>Página {block.page} · {html.escape(block.kind)} · {html.escape(block.role)}</small><br>{html.escape(block.text)}'
                            + (f'<br><strong>Conferir: {attention}</strong>' if attention else '') + '</p>')
        elif isinstance(block, ImageBlock):
            sections.append(f'<p><small>Página {block.page} · imagem do anexo incorporada</small></p>')
        else:
            presented, adjustments = formatted_table_matrix(block, normalize_municipalities)
            values = sum(monetary_values(block.matrix).values())
            differences = list(block.issues)
            if matrix_signature(block.matrix) != block.source_signature:
                differences.append("O texto do Word difere dos caracteres presentes nesta região do PDF.")
            if monetary_values(block.matrix) != block.source_numbers:
                differences.append("Há divergência entre os valores monetários do PDF e do Word.")
            if block.matrix != block.original_matrix:
                differences.append("Algumas células foram corrigidas na interface.")
            if differences:
                status = ("Conferência manual declarada. " if block.reviewed else "Revisão necessária. ") + "; ".join(differences)
            else:
                status = "Texto e valores da extração preservados; confira a associação entre as linhas no PDF"
            table_rows = []
            for row in presented:
                table_rows.append('<tr>' + ''.join(f'<td>{html.escape(value).replace(chr(10), "<br>")}</td>' for value in row) + '</tr>')
            parts = table_page_parts(block)
            source_pages = list(dict.fromkeys(part.page for part in parts))
            page_label = f"página {source_pages[0]}" if len(source_pages) == 1 else (
                "páginas " + " e ".join(map(str, source_pages)))
            locations = "; ".join(
                f"página {part.page}" + (f", coluna {part.column}" if part.column else "") +
                f", linhas {part.first_row + 1} a {part.last_row}"
                for part in parts)
            adjustment_rows = ''.join(
                '<tr>' + ''.join(f'<td>{html.escape(str(v))}</td>' for v in
                                (ri + 1, ci + 1, before, after, rule)) + '</tr>'
                for ri, ci, before, after, rule in adjustments)
            adjustment_report = (f'<details><summary>{len(adjustments)} ajustes automáticos de grafia '
                                 '(maiúsculas/minúsculas e quebras de cabeçalho)</summary>'
                                 '<p>Os números, acentos e a ordem das linhas foram preservados. '
                                 'Cada alteração abaixo compara a matriz extraída com o texto gravado no Word.</p>'
                                 '<table><tr><th>Linha</th><th>Coluna</th><th>Extraído do PDF</th>'
                                 '<th>No Word</th><th>Regra</th></tr>' + adjustment_rows + '</table></details>') if adjustments else ''
            sections.append(f'<h2>Tabela em {html.escape(page_label)} · detecção por {html.escape(block.strategy)}</h2><p>{len(block.matrix)} linhas, {len(block.matrix[0])} colunas, '
                            f'{values} valores monetários. Origem: {html.escape(locations)}. '
                            f'{html.escape(status)}.</p>{adjustment_report}<table>{"".join(table_rows)}</table>')
    issues = "".join(f'<li>{html.escape(issue)}</li>' for issue in extraction.issues + extraction.blocking_issues)
    page = f'''<!doctype html><html lang="pt-BR"><meta charset="utf-8"><title>Conferência do PDF SEFA</title>
<style>body{{font:14px/1.45 Arial,sans-serif;max-width:1400px;margin:32px auto;padding:0 22px;color:#222}}
table{{border-collapse:collapse;width:100%;font-size:12px;margin-bottom:25px}}
td{{border:1px solid #aaa;padding:4px;vertical-align:top;overflow-wrap:anywhere}}
td:not(:first-child){{text-align:right}}small{{color:#555}}h1{{font-size:22px}}h2{{font-size:17px}}</style>
<h1>Conferência da importação do PDF</h1>
<p>Identificador: {html.escape(extraction.query)} · páginas {', '.join(map(str, extraction.pages))}.
Trechos de texto entre tabelas verificados: {extraction.intertable_segments_checked}.
O conteúdo das tabelas é editável no Word. Confira a correspondência entre rótulos, números e o PDF antes do uso.</p>
<p>Padronização da grafia de municípios: {'ativada' if normalize_municipalities else 'desativada'}.
Somente campos identificados pelo cabeçalho Município recebem iniciais maiúsculas.
Os ajustes de apresentação aparecem em uma lista própria, sem substituir a matriz usada na auditoria.</p>
<ul>{issues}</ul>{''.join(sections)}
<p>Marcador de término identificado: {html.escape(extraction.protocol or 'não encontrado')}.
Conferência manual dos limites: {'sim' if extraction.limits_reviewed else 'não'}.</p></html>'''
    report_path.write_text(page, encoding="utf-8")
