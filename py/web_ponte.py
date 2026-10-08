"""Ponte entre a página web e o Formatador SEFA, rodando no navegador (Pyodide).

Os módulos do programa desktop (extrator_pdf, importar_pdf, formatador_sefa,
links_no_word, links_legislacao) entram SEM alteração. Tudo o que é específico do
navegador fica aqui:

- o pdfplumber é instalado sem o pypdfium2 (que não tem versão para esta edição do
  Pyodide); as poucas chamadas a page.to_image() passam a ser desenhadas pelo PyMuPDF;
- a busca de links consulta apenas a tabela_links.tsv embutida no site, porque a página
  não consegue falar com o site de legislação da SEFA nem com o Planalto;
- arquivos entram e saem pela pasta de trabalho /work (o JavaScript grava o PDF lá e lê
  o Word e o relatório gerados).

Cada função pública devolve um texto JSON, que o JavaScript converte com JSON.parse.
"""

from __future__ import annotations

import io
import json
import re
import shutil
import traceback
import warnings
from pathlib import Path

warnings.filterwarnings("ignore", message=".*fitz.*deprecated.*")

import fitz  # noqa: E402
import pdfplumber  # noqa: E402
import pdfplumber.page  # noqa: E402
from PIL import Image  # noqa: E402

TRABALHO = Path("/work")
TRABALHO.mkdir(exist_ok=True)
ESTADO: dict = {"extracao": None, "pdf": None, "consulta": "", "municipios": True,
                "versao": 0, "previa": None}
_HISTORICO: list = []        # cópias da análise antes de cada correção, para desfazer


# ---------------------------------------------------------------------------
# 1) page.to_image() sem pypdfium2: desenha o recorte com o PyMuPDF
# ---------------------------------------------------------------------------
class _ImagemPagina:
    """Imita o pdfplumber.display.PageImage no que o extrator usa (.original)."""

    def __init__(self, original: Image.Image):
        self.original = original


def _to_image_pymupdf(self, resolution=None, width=None, height=None, antialias=False, **_):
    caminho = ESTADO["pdf"]
    if caminho is None:
        raise RuntimeError("PDF de origem não definido para desenhar a página.")
    x0, top, x1, bottom = self.bbox
    with fitz.open(str(caminho)) as documento:
        pagina = documento[self.page_number - 1]
        if resolution is None:
            if width:
                resolution = 72 * width / max(x1 - x0, 1)
            elif height:
                resolution = 72 * height / max(bottom - top, 1)
            else:
                resolution = 72
        pix = pagina.get_pixmap(dpi=int(round(resolution)), clip=fitz.Rect(x0, top, x1, bottom) & pagina.rect,
                                alpha=False)
        imagem = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    return _ImagemPagina(imagem)


pdfplumber.page.Page.to_image = _to_image_pymupdf

# Só agora os módulos do programa (eles importam pdfplumber e fitz)
import extrator_pdf as ex  # noqa: E402
import importar_pdf as imp  # noqa: E402
import links_legislacao as ll  # noqa: E402
import links_no_word as lnw  # noqa: E402
from formatador_sefa import DEFAULT_MODEL, STYLES  # noqa: E402


# ---------------------------------------------------------------------------
# 2) links: só a tabela embutida
# ---------------------------------------------------------------------------
SITUACAO_ONLINE = "Não buscado (a versão online consulta só a tabela de links)"


class _SemSite:
    def __init__(self, *_a, **_k):
        pass

    def __enter__(self):
        raise RuntimeError("a versão online consulta só a tabela de links")

    def __exit__(self, *exc):
        return False


ll.SiteSEFA = _SemSite
ll.S_SITE_OFF = SITUACAO_ONLINE
ll.buscar_planalto = lambda o, site=None: ("", SITUACAO_ONLINE)
ll.salvar_tabela = lambda d: None  # nada novo para guardar; o arquivo do site é só leitura

# Busca assistida: os links que a pessoa colou na página (guardados no navegador dela)
# entram junto com a tabela do site. Em caso de conflito, vale a tabela do site.
LINKS_USUARIO: dict[str, str] = {}
_tabela_do_site = ll.carregar_tabela
ll.carregar_tabela = lambda: {**LINKS_USUARIO, **_tabela_do_site()}

# Guarda o resultado completo da última etapa de links (com a chave de cada ato),
# que o links_no_word resume sem as chaves.
_ULTIMO_LINKS: dict = {}
_criar_links_original = ll.criar_links


def _criar_links_registrando(doc, log=print, parar=None):
    resultado = _criar_links_original(doc, log, parar)
    _ULTIMO_LINKS["res"] = resultado
    return resultado


ll.criar_links = _criar_links_registrando


def endereco_busca(o) -> str:
    """Página de resultados do site da SEFA já filtrada pelo número e ano do ato."""
    if o.tipo == "Constituição" or not o.numero:
        return ll.CONFIG["pagina_busca"]
    params = {"idStatus": 3, "numeroAto": o.numero, "ano": o.ano or "", "ordenacao": "MAIS_RECENTES"}
    return ll.CONFIG["pagina_busca"] + "/resultados?" + "&".join(f"{k}={v}" for k, v in params.items() if v != "")


def _pendentes_para_busca() -> tuple[list[dict], list[str]]:
    res = _ULTIMO_LINKS.get("res")
    if res is None:
        return [], []
    busca, fora = [], []
    for chave, a in res.atos.items():
        if a["url"]:
            continue
        if a["sit"] == ll.S_FORA_SITE:
            fora.append(a["exemplo"])
            continue
        busca.append({"chave": chave, "exemplo": a["exemplo"], "qtd": a["qtd"], "situacao": a["sit"],
                      "busca": endereco_busca(a["o"])})
    return busca, fora


def definir_links_usuario(texto_json: str) -> str:
    try:
        dados = json.loads(texto_json or "{}")
        LINKS_USUARIO.clear()
        LINKS_USUARIO.update({str(k): str(v) for k, v in dados.items()
                              if str(v).startswith(("https://", "http://")) and "|" in str(k)})
        return _json(quantidade=len(LINKS_USUARIO))
    except Exception as exc:
        return _erro(exc)


# ---------------------------------------------------------------------------
# utilidades
# ---------------------------------------------------------------------------
def _json(ok: bool = True, **dados) -> str:
    return json.dumps({"ok": ok, **dados}, ensure_ascii=False)


def _erro(exc: BaseException) -> str:
    return _json(False, erro=str(exc) or exc.__class__.__name__, detalhe=traceback.format_exc())


def _extracao():
    if ESTADO["extracao"] is None:
        raise RuntimeError("Analise um PDF primeiro.")
    return ESTADO["extracao"]


def nome_sugerido(consulta: str) -> str:
    return re.sub(r"[^\w.-]+", "_", consulta).strip("_") + "_SEFA.docx"


def estado() -> str:
    """Tudo o que a tela precisa para desenhar a prévia (equivale ao refresh() da janela)."""
    try:
        e = _extracao()
        municipios = ESTADO["municipios"]
        paragrafos = [{"i": i, "pagina": b.page, "parte": b.kind, "estilo": b.role, "texto": b.text,
                       "atencao": list(b.attention), "formatado": bool(getattr(b, "formatacao_manual", None))}
                      for i, b in enumerate(e.paragraphs)]
        tabelas = []
        for ti, t in enumerate(e.tables):
            apresentada, _ = ex.formatted_table_matrix(t, municipios)
            partes = ex.table_page_parts(t)
            tabelas.append({"i": ti, "linhas": apresentada, "problemas": list(t.issues),
                            "conferida": t.reviewed,
                            "paginas": sorted({p.page for p in partes})})
        ordem, ip, it = [], 0, 0
        for bloco in e.blocks:  # sequência real do documento: parágrafos, tabelas e imagens
            if isinstance(bloco, ex.ParagraphBlock):
                ordem.append(["p", ip, bloco.page]); ip += 1
            elif isinstance(bloco, ex.TableBlock):
                ordem.append(["t", it, bloco.page]); it += 1
            else:
                ordem.append(["img", 0, bloco.page])
        valores = sum(sum(t.source_numbers.values()) for t in e.tables)
        erros_tabela = imp.check_table_numbers(e)
        return _json(paginas=list(e.pages), paragrafos=paragrafos, tabelas=tabelas, ordem=ordem, valores=valores,
                     erros_tabela=erros_tabela, bloqueios=list(e.blocking_issues),
                     limites=list(e.issues), limites_conferidos=e.limits_reviewed,
                     estilos=list(STYLES), municipios=municipios, pode_desfazer=len(_HISTORICO),
                     removidos=len(getattr(e, "removidos_web", [])),
                     nome_sugerido=nome_sugerido(ESTADO["consulta"]))
    except Exception as exc:
        return _erro(exc)


# ---------------------------------------------------------------------------
# desfazer e formatação manual
# ---------------------------------------------------------------------------
def _guardar_para_desfazer() -> None:
    import copy as _copy

    _HISTORICO.append(_copy.deepcopy(_extracao()))
    del _HISTORICO[:-30]
    ESTADO["versao"] += 1


def _anotar(bloco, nota: str) -> None:
    if nota not in bloco.attention:
        bloco.attention.append(nota)


def _limpar_runs(runs) -> list[dict]:
    """[{t, b, i, u}] sem trechos vazios, juntando vizinhos com a mesma formatação e sem
    espaços nas pontas do parágrafo."""
    limpos: list[dict] = []
    for r in runs or []:
        t = str(r.get("t", ""))
        if not t:
            continue
        f = {"b": bool(r.get("b")), "i": bool(r.get("i")), "u": bool(r.get("u"))}
        if limpos and all(limpos[-1][k] == f[k] for k in f):
            limpos[-1]["t"] += t
        else:
            limpos.append({"t": t, **f})
    if limpos:
        limpos[0]["t"] = limpos[0]["t"].lstrip()
        limpos[-1]["t"] = limpos[-1]["t"].rstrip()
    return [r for r in limpos if r["t"]]


def _texto_runs(runs) -> str:
    return "".join(r["t"] for r in runs)


def _ligar_paragrafos(doc, extracao):
    """Liga cada parágrafo de texto do corpo do Word ao trecho da análise, pela ordem e
    pelo texto sem espaços. Devolve ({índice do trecho: elemento w:p}, mapa para a tela)."""
    trechos = [_chave_texto(b.text) for b in extracao.paragraphs]
    ponteiro, ligados, mapa = 0, {}, []
    for filho in doc.element.body.iterchildren(ex.qn("w:p")):
        chave = _chave_texto(_texto_paragrafo_docx(filho))
        if not chave:
            continue
        indice = -1
        for passo in range(8):
            j = ponteiro + passo
            if j < len(trechos) and trechos[j] == chave:
                indice, ponteiro = j, j + 1
                break
        if indice >= 0:
            ligados[indice] = filho
        mapa.append({"k": chave, "i": indice})
    return ligados, mapa


def _valor_estilo(estilo, atributo):
    while estilo is not None:
        valor = getattr(estilo.font, atributo, None)
        if valor is not None:
            return valor
        estilo = estilo.base_style
    return None


def _formatacao_efetiva(run, paragrafo, atributo) -> bool:
    valor = getattr(run, atributo)
    if valor is None:
        try:
            valor = _valor_estilo(run.style, atributo)
        except Exception:
            valor = None
    if valor is None:
        valor = _valor_estilo(paragrafo.style, atributo)
    return bool(valor)


def _runs_do_paragrafo(doc, p_el) -> list[dict]:
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run

    paragrafo = Paragraph(p_el, doc._body)
    runs = []
    for r_el in p_el.iter(ex.qn("w:r")):
        run = Run(r_el, paragrafo)
        runs.append({"t": run.text, "b": _formatacao_efetiva(run, paragrafo, "bold"),
                     "i": _formatacao_efetiva(run, paragrafo, "italic"),
                     "u": _formatacao_efetiva(run, paragrafo, "underline")})
    return _limpar_runs(runs)


def _runs_do_trecho(indice: int) -> list[dict]:
    """Formatação atual do trecho: a manual, se houver; senão, a do Word da última
    visualização (se ainda estiver em dia); senão, o texto sem formatação."""
    from docx import Document as _Document

    bloco = _extracao().paragraphs[indice]
    manual = getattr(bloco, "formatacao_manual", None)
    if manual:
        return [dict(r) for r in manual]
    previa = ESTADO.get("previa")
    if previa and previa["versao"] == ESTADO["versao"] and Path(previa["caminho"]).exists():
        try:
            doc = _Document(previa["caminho"])
            ligados, _ = _ligar_paragrafos(doc, _extracao())
            if indice in ligados:
                runs = _runs_do_paragrafo(doc, ligados[indice])
                if _chave_texto(_texto_runs(runs)) == _chave_texto(bloco.text):
                    return runs
        except Exception:
            pass
    return [{"t": bloco.text, "b": False, "i": False, "u": False}]


def _substituir_runs(doc, p_el, runs) -> None:
    import copy as _copy
    from docx.text.paragraph import Paragraph

    modelo = None
    for r_el in p_el.iter(ex.qn("w:r")):
        rpr = r_el.find(ex.qn("w:rPr"))
        if rpr is not None:
            modelo = _copy.deepcopy(rpr)
            break
    for filho in list(p_el):
        if filho.tag != ex.qn("w:pPr"):
            p_el.remove(filho)
    paragrafo = Paragraph(p_el, doc._body)
    for r in runs:
        novo = paragrafo.add_run()
        if modelo is not None:
            novo._r.insert(0, _copy.deepcopy(modelo))
        novo.text = r["t"]
        novo.bold, novo.italic, novo.underline = r["b"], r["i"], r["u"]


def _aplicar_formatacao_manual(docx_path: Path, extracao) -> int:
    """Depois de montar o Word, troca os trechos com formatação manual pelo que a pessoa
    definiu. Só vale se o texto do trecho ainda for o mesmo da formatação."""
    from docx import Document as _Document

    blocos = extracao.paragraphs
    if not any(getattr(b, "formatacao_manual", None) for b in blocos):
        return 0
    doc = _Document(str(docx_path))
    ligados, _ = _ligar_paragrafos(doc, extracao)
    feitos = 0
    for i, bloco in enumerate(blocos):
        runs = getattr(bloco, "formatacao_manual", None)
        if not runs or i not in ligados:
            continue
        if _chave_texto(_texto_runs(runs)) != _chave_texto(bloco.text):
            continue
        _substituir_runs(doc, ligados[i], runs)
        feitos += 1
    doc.save(str(docx_path))
    return feitos


def _secao_removidos(extracao) -> str:
    import html as _html

    removidos = getattr(extracao, "removidos_web", [])
    if not removidos:
        return ""
    linhas = "".join(f"<li><small>Página {r['pagina']} · {_html.escape(r['estilo'])}</small><br>"
                     f"{_html.escape(r['texto'])}</li>" for r in removidos)
    return ('<section id="removidos-web"><h2>Trechos apagados na visualização do Word</h2>'
            f"<p>Estes trechos foram lidos do PDF e apagados manualmente antes de gerar o Word.</p><ol>{linhas}</ol></section>")


def _anexar_ao_relatorio(caminho: Path, secao: str) -> None:
    if not secao:
        return
    texto = caminho.read_text(encoding="utf-8")
    fim = texto.rfind("</html>")
    caminho.write_text(texto[:fim] + secao + texto[fim:] if fim >= 0 else texto + secao, encoding="utf-8")


def trecho_formatado(indice: int) -> str:
    try:
        bloco = _extracao().paragraphs[int(indice)]
        return _json(runs=_runs_do_trecho(int(indice)), estilo=bloco.role, atencao=list(bloco.attention))
    except Exception as exc:
        return _erro(exc)


def aplicar_trecho(indice: int, runs_json: str, estilo: str, formatou: bool) -> str:
    try:
        e = _extracao()
        bloco = e.paragraphs[int(indice)]
        runs = _limpar_runs(json.loads(runs_json))
        texto = _texto_runs(runs)
        if not texto.strip():
            raise ValueError("O trecho ficou vazio. Para tirá-lo do Word, use Apagar parágrafo.")
        _guardar_para_desfazer()
        bloco = _extracao().paragraphs[int(indice)]
        if texto != bloco.text:
            _anotar(bloco, "Trecho corrigido manualmente na prévia.")
        if formatou or getattr(bloco, "formatacao_manual", None):
            bloco.formatacao_manual = runs
            if formatou:
                _anotar(bloco, "Formatação ajustada manualmente na visualização do Word.")
        bloco.text = texto
        bloco.role = estilo or bloco.role
        return estado()
    except Exception as exc:
        return _erro(exc)


def _posicao_no_documento(e, bloco) -> int:
    return next(k for k, x in enumerate(e.blocks) if x is bloco)


def apagar_paragrafo(indice: int) -> str:
    try:
        _guardar_para_desfazer()
        e = _extracao()
        bloco = e.paragraphs[int(indice)]
        e.removidos_web = list(getattr(e, "removidos_web", [])) + [
            {"pagina": bloco.page, "estilo": bloco.role, "texto": bloco.text}]
        del e.blocks[_posicao_no_documento(e, bloco)]
        return estado()
    except Exception as exc:
        return _erro(exc)


def juntar_paragrafos(indice: int) -> str:
    """Junta o trecho indice com o seguinte (indice + 1)."""
    try:
        e = _extracao()
        indice = int(indice)
        if not 0 <= indice < len(e.paragraphs) - 1:
            raise ValueError("Não há outro parágrafo para juntar nessa direção.")
        a, b = e.paragraphs[indice], e.paragraphs[indice + 1]
        if _posicao_no_documento(e, b) != _posicao_no_documento(e, a) + 1:
            raise ValueError("Há uma tabela ou imagem entre esses dois trechos; não dá para juntá-los.")
        manual = bool(getattr(a, "formatacao_manual", None) or getattr(b, "formatacao_manual", None))
        runs = _runs_do_trecho(indice) + [{"t": " ", "b": False, "i": False, "u": False}] + _runs_do_trecho(indice + 1)
        _guardar_para_desfazer()
        e = _extracao()
        a, b = e.paragraphs[indice], e.paragraphs[indice + 1]
        a.text = a.text.rstrip() + " " + b.text.lstrip()
        if manual:
            a.formatacao_manual = _limpar_runs(runs)
        for nota in b.attention:
            _anotar(a, nota)
        _anotar(a, "Parágrafos unidos manualmente na visualização do Word.")
        del e.blocks[_posicao_no_documento(e, b)]
        return estado()
    except Exception as exc:
        return _erro(exc)


def dividir_paragrafo(indice: int, antes_json: str, depois_json: str, estilo: str, formatou: bool) -> str:
    try:
        antes, depois = _limpar_runs(json.loads(antes_json)), _limpar_runs(json.loads(depois_json))
        if not _texto_runs(antes).strip() or not _texto_runs(depois).strip():
            raise ValueError("Clique no ponto do texto onde o parágrafo deve ser dividido, sem ser no começo nem no fim.")
        _guardar_para_desfazer()
        e = _extracao()
        bloco = e.paragraphs[int(indice)]
        manual = formatou or getattr(bloco, "formatacao_manual", None)
        novo = ex.ParagraphBlock(_texto_runs(depois), estilo or bloco.role, bloco.page, bloco.kind, None,
                                 ["Parágrafo criado ao dividir um trecho na visualização do Word."])
        bloco.text, bloco.role = _texto_runs(antes), estilo or bloco.role
        _anotar(bloco, "Trecho dividido manualmente na visualização do Word.")
        if manual:
            bloco.formatacao_manual, novo.formatacao_manual = antes, depois
        elif hasattr(bloco, "formatacao_manual"):
            del bloco.formatacao_manual
        e.blocks.insert(_posicao_no_documento(e, bloco) + 1, novo)
        return estado()
    except Exception as exc:
        return _erro(exc)


def desfazer() -> str:
    try:
        if not _HISTORICO:
            raise ValueError("Não há correção para desfazer.")
        ESTADO["extracao"] = _HISTORICO.pop()
        ESTADO["versao"] += 1
        return estado()
    except Exception as exc:
        return _erro(exc)


# ---------------------------------------------------------------------------
# ações da tela
# ---------------------------------------------------------------------------
def analisar(caminho_pdf: str, consulta: str, max_paginas: int = 12) -> str:
    try:
        consulta = (consulta or "").strip()
        if not consulta:
            raise ValueError("Digite o número do ato, por exemplo 494/2026-SEFA.GS.")
        max_paginas = int(max_paginas)
        if not 1 <= max_paginas <= 100:
            raise ValueError("Informe um número entre 1 e 100 páginas.")
        caminho = Path(caminho_pdf)
        ESTADO["pdf"] = caminho
        ESTADO["extracao"] = ex.extract_act(caminho, consulta, max_pages=max_paginas)
        ESTADO["consulta"] = consulta
        ESTADO["versao"] += 1
        _HISTORICO.clear()
        return estado()
    except Exception as exc:
        ESTADO["extracao"] = None
        return _erro(exc)


def definir_municipios(valor: bool) -> str:
    ESTADO["municipios"] = bool(valor)
    return estado()


def editar_paragrafo(indice: int, texto: str, estilo: str) -> str:
    try:
        _guardar_para_desfazer()
        bloco = _extracao().paragraphs[int(indice)]
        novo = (texto or "").strip()
        manual = getattr(bloco, "formatacao_manual", None)
        if manual and _chave_texto(novo) != _chave_texto(bloco.text):
            del bloco.formatacao_manual
            _anotar(bloco, "A formatação manual foi descartada porque o texto foi corrigido na prévia.")
        bloco.text = novo
        bloco.role = estilo or bloco.role
        bloco.attention.append("Trecho corrigido manualmente na prévia.")
        return estado()
    except Exception as exc:
        return _erro(exc)


def editar_celula(tabela: int, linha: int, coluna: int, texto: str) -> str:
    try:
        _guardar_para_desfazer()
        t = _extracao().tables[int(tabela)]
        t.matrix[int(linha)][int(coluna)] = (texto or "").strip()
        t.reviewed = False
        return estado()
    except Exception as exc:
        return _erro(exc)


def conferir_tabela(tabela: int) -> str:
    try:
        _guardar_para_desfazer()
        _extracao().tables[int(tabela)].reviewed = True
        return estado()
    except Exception as exc:
        return _erro(exc)


def conferir_limites() -> str:
    try:
        _guardar_para_desfazer()
        _extracao().limits_reviewed = True
        return estado()
    except Exception as exc:
        return _erro(exc)


def recorte_tabela(tabela: int, linha: int, destino: str) -> str:
    """Grava em destino o PNG do trecho do PDF onde está a linha da tabela (Ver tabela no PDF)."""
    try:
        e = _extracao()
        t = e.tables[int(tabela)]
        partes = ex.table_page_parts(t)
        parte = next((p for p in partes if p.first_row <= int(linha) < p.last_row), partes[0])
        with fitz.open(str(e.pdf_path)) as pdf:
            pagina = pdf[parte.page - 1]
            caixa = fitz.Rect(*parte.bbox) + (-15, -15, 15, 15)
            pix = pagina.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), clip=caixa & pagina.rect, alpha=False)
            Path(destino).write_bytes(pix.tobytes("png"))
        local = f", coluna {parte.column}" if parte.column else ""
        return _json(legenda=f"Tabela {int(tabela) + 1}, página {parte.page}{local} do PDF")
    except Exception as exc:
        return _erro(exc)


def gerar(saida: str, modelo: str | None, relatorio: bool, links: bool, links_usuario: str = "{}") -> str:
    """Gera o Word (e o relatório, se pedido) e aplica os links da tabela embutida
    e dos links que a pessoa já colou na busca assistida."""
    try:
        e = _extracao()
        definir_links_usuario(links_usuario)
        _ULTIMO_LINKS.clear()
        saida_p = Path(saida)
        relatorio_p = saida_p.with_name(saida_p.stem + "_CONFERIR.html") if relatorio else None
        for antigo in (saida_p, relatorio_p):
            if antigo is not None and antigo.exists():
                antigo.unlink()
        modelo_p = Path(modelo) if modelo else DEFAULT_MODEL
        imp.export(e, modelo_p, saida_p, relatorio_p, force=True,
                   normalize_municipalities=ESTADO["municipios"])
        _aplicar_formatacao_manual(saida_p, e)
        if relatorio_p is not None:
            _anexar_ao_relatorio(relatorio_p, _secao_removidos(e))
        registro: list[str] = []
        resumo_txt, aviso, busca, fora = "", "", [], []
        if links:
            resumo = lnw.aplicar_e_registrar(saida_p, relatorio_p, registro.append)
            resumo_txt = resumo.frase()
            if resumo.concluido:
                busca, fora = _pendentes_para_busca()
            else:
                aviso = resumo.aviso
        return _json(word=str(saida_p), relatorio=str(relatorio_p) if relatorio_p else "",
                     resumo_links=resumo_txt, aviso_links=aviso, pendentes_busca=busca,
                     fora_do_site=fora, registro=registro)
    except Exception as exc:
        return _erro(exc)


def _chave_texto(texto: str) -> str:
    """Texto sem nenhum espaço, para casar o parágrafo do Word com o desenho na tela
    (quebras de linha e tabulações viram nada nos dois lados)."""
    return re.sub(r"\s+", "", texto or "")


def _texto_paragrafo_docx(p) -> str:
    partes = []
    for el in p.iter():
        if el.tag == ex.qn("w:t"):
            partes.append(el.text or "")
        elif el.tag in (ex.qn("w:tab"), ex.qn("w:br"), ex.qn("w:cr")):
            partes.append(" ")
    return "".join(partes)


def _mapa_paragrafos(docx_path: Path, extracao) -> list[dict]:
    """Para cada parágrafo de texto do corpo do Word, na ordem: a chave do texto e o índice
    do trecho correspondente na prévia (-1 quando é espaço, imagem etc.)."""
    from docx import Document as _Document

    _, mapa = _ligar_paragrafos(_Document(str(docx_path)), extracao)
    return mapa


def previa_word(destino: str, modelo: str, links: bool, links_usuario: str = "{}") -> str:
    """Monta o Word como ficará no download (com links e destaques), só para visualizar.
    Usa uma cópia da extração: nada aqui marca conferências nem altera o que será exportado,
    e tabelas ainda não conferidas aparecem como estão."""
    try:
        import copy as _copy

        e = _copy.deepcopy(_extracao())
        for t in e.tables:
            t.reviewed = True  # só para permitir desenhar; a exportação de verdade continua exigindo
        definir_links_usuario(links_usuario)
        destino_p = Path(destino)
        if destino_p.exists():
            destino_p.unlink()
        modelo_p = Path(modelo) if modelo else DEFAULT_MODEL
        ex.build_word(e, modelo_p, destino_p, overwrite=True, normalize_municipalities=ESTADO["municipios"])
        _aplicar_formatacao_manual(destino_p, e)
        sem_link = 0
        if links:
            resumo = lnw.aplicar_links(destino_p, lambda *_: None)
            sem_link = resumo.destacadas if resumo.concluido else 0
        ESTADO["previa"] = {"versao": ESTADO["versao"], "caminho": str(destino_p)}
        return _json(word=str(destino_p), mapa=_mapa_paragrafos(destino_p, e), sem_link=sem_link)
    except Exception as exc:
        return _erro(exc)


def versao() -> str:
    return _json(pymupdf=fitz.VersionBind, pdfplumber=pdfplumber.__version__,
                 tabela_links=len(_tabela_do_site()), pagina_busca=ll.CONFIG["pagina_busca"])
