#!/usr/bin/env python3
"""Importa um ato do PDF do diário oficial para Word editável da SEFA."""

from __future__ import annotations

import argparse
import base64
import re
import sys
from pathlib import Path

import fitz
from docx import Document

from extrator_pdf import (Extraction, ParagraphBlock, build_word, extract_act,
                         formatted_table_matrix, matrix_signature, monetary_values, table_page_parts,
                         write_pdf_report)
from formatador_sefa import DEFAULT_MODEL, STYLES, show_and_open
from links_no_word import aplicar_e_registrar, executar_com_progresso, texto_final


def check_table_numbers(extraction: Extraction) -> list[str]:
    errors = []
    for index, table in enumerate(extraction.tables, 1):
        concerns = []
        actual = monetary_values(table.matrix)
        if actual != table.source_numbers:
            concerns.append(f"valores monetários divergem (PDF: {sum(table.source_numbers.values())}; "
                            f"Word: {sum(actual.values())})")
        if matrix_signature(table.matrix) != table.source_signature:
            concerns.append("o texto das células diverge do PDF (inclui rótulos e percentuais)")
        if table.matrix != table.original_matrix:
            concerns.append("houve edição manual de células")
        concerns.extend(table.issues)
        if concerns and not table.reviewed:
            errors.append(f"Tabela {index}: " + "; ".join(concerns))
    return errors


def export(extraction: Extraction, model: Path, output: Path, report: Path | None, force: bool = False,
           normalize_municipalities: bool = True) -> None:
    """Gera o Word; o relatório HTML só é gravado quando report não é None."""
    if not model.is_file():
        raise FileNotFoundError(f"Modelo Word não encontrado: {model}")
    if output.resolve() == model.resolve():
        raise ValueError("O nome da saída deve ser diferente do documento modelo.")
    if report is not None and output.resolve() == report.resolve():
        raise ValueError("O Word e o relatório devem ter nomes diferentes.")
    if (output.exists() or (report is not None and report.exists())) and not force:
        raise FileExistsError("O Word ou o relatório já existe. Escolha outro nome ou autorize a substituição.")
    errors = check_table_numbers(extraction)
    errors.extend(extraction.blocking_issues)
    if extraction.issues and not extraction.limits_reviewed:
        errors.append("Confirme o início e o fim do ato antes de exportar: " + "; ".join(extraction.issues))
    if errors:
        raise ValueError("\n".join(errors))
    Document(model)  # verifica o modelo antes de começar a gravar a saída
    build_word(extraction, model, output, overwrite=force, normalize_municipalities=normalize_municipalities)
    if report is not None:
        write_pdf_report(extraction, report, normalize_municipalities)


def run_gui() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    class App:
        def __init__(self) -> None:
            self.root = tk.Tk()
            self.root.title("SEFA 3.18 — importar ato completo do PDF")
            self.root.geometry("1100x780")
            self.root.minsize(800, 580)
            self.pdf_var = tk.StringVar()
            self.act_var = tk.StringVar()
            self.max_pages_var = tk.StringVar(value="12")
            self.model_var = tk.StringVar(value=str(DEFAULT_MODEL))
            self.municipalities_var = tk.BooleanVar(value=True)
            self.links_var = tk.BooleanVar(value=True)
            self.report_var = tk.BooleanVar(value=False)
            self.status = tk.StringVar(value="Escolha o PDF e digite o número do ato, por exemplo 494/2026-SEFA.GS.")
            self.role = tk.StringVar()
            self.extraction: Extraction | None = None
            self.selection: tuple[str, int, int, int] | None = None
            self.read_pdf = ""
            self.read_act = ""
            self.read_max_pages = ""
            self.layout()

        def layout(self) -> None:
            root = self.root
            root.columnconfigure(0, weight=1)
            root.rowconfigure(1, weight=1)
            inputs = ttk.Frame(root, padding=10)
            inputs.grid(row=0, column=0, sticky="ew")
            inputs.columnconfigure(1, weight=1)
            ttk.Label(inputs, text="PDF do diário oficial").grid(row=0, column=0, sticky="w")
            ttk.Entry(inputs, textvariable=self.pdf_var).grid(row=0, column=1, sticky="ew", padx=7, pady=3)
            ttk.Button(inputs, text="Escolher PDF…", command=self.pick_pdf).grid(row=0, column=2)
            ttk.Label(inputs, text="Número do ato").grid(row=1, column=0, sticky="w")
            ttk.Entry(inputs, textvariable=self.act_var).grid(row=1, column=1, sticky="ew", padx=7, pady=3)
            ttk.Button(inputs, text="Analisar PDF", command=self.analyze).grid(row=1, column=2, sticky="ew")
            ttk.Label(inputs, text="Modelo Word da SEFA").grid(row=2, column=0, sticky="w")
            ttk.Entry(inputs, textvariable=self.model_var).grid(row=2, column=1, sticky="ew", padx=7, pady=3)
            ttk.Button(inputs, text="Outro modelo…", command=self.pick_model).grid(row=2, column=2)
            ttk.Label(inputs, text="Máximo de páginas do ato").grid(row=3, column=0, sticky="w")
            ttk.Entry(inputs, textvariable=self.max_pages_var, width=8).grid(row=3, column=1, sticky="w", padx=7, pady=3)
            ttk.Checkbutton(inputs, text="Padronizar grafia de municípios (os ajustes vão para o relatório, quando gerado)",
                            variable=self.municipalities_var, command=self.change_case).grid(
                                row=4, column=1, columnspan=2, sticky="w", padx=7, pady=3)
            ttk.Checkbutton(inputs, text="Criar hyperlinks das leis, decretos etc. citados "
                            "(tabela local, site de legislação da SEFA e Planalto)",
                            variable=self.links_var).grid(row=5, column=1, columnspan=2, sticky="w", padx=7, pady=3)
            ttk.Checkbutton(inputs, text="Gerar relatório de conferência (HTML ao lado do Word)",
                            variable=self.report_var).grid(row=6, column=1, columnspan=2, sticky="w", padx=7, pady=3)
            notebook = ttk.Notebook(root)
            notebook.grid(row=1, column=0, sticky="nsew", padx=10)
            for title, tag, columns in (
                ("Texto do ato", "paragraph", (("page", "Página", 75), ("kind", "Parte", 125),
                                                   ("role", "Estilo", 145), ("text", "Texto", 700))),
                ("Células das tabelas", "cell", (("table", "Tabela", 70), ("row", "Linha", 70),
                                                     ("col", "Coluna", 70), ("text", "Conteúdo", 760))),
            ):
                frame = ttk.Frame(notebook)
                notebook.add(frame, text=title)
                frame.columnconfigure(0, weight=1)
                frame.rowconfigure(0, weight=1)
                tree = ttk.Treeview(frame, columns=[name for name, _, _ in columns], show="headings", selectmode="browse")
                for name, label, width in columns:
                    tree.heading(name, text=label)
                    tree.column(name, width=width, stretch=name == "text")
                tree.grid(row=0, column=0, sticky="nsew")
                ybar = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
                ybar.grid(row=0, column=1, sticky="ns")
                tree.configure(yscrollcommand=ybar.set)
                tree.tag_configure("attention", background="#fff1d2")
                tree.bind("<<TreeviewSelect>>", lambda event, section=tag: self.selected(section))
                setattr(self, tag + "_tree", tree)
            editor = ttk.LabelFrame(root, text="Editar o item selecionado", padding=8)
            editor.grid(row=2, column=0, sticky="ew", padx=10, pady=6)
            editor.columnconfigure(0, weight=1)
            self.edit_text = tk.Text(editor, height=3, wrap="word", font=("Arial", 10))
            self.edit_text.grid(row=0, column=0, columnspan=3, sticky="ew")
            ttk.Label(editor, text="Tipo do parágrafo:").grid(row=1, column=0, sticky="w")
            self.role_combo = ttk.Combobox(editor, textvariable=self.role, state="readonly", width=24,
                                            values=tuple(STYLES))
            self.role_combo.grid(row=1, column=1, sticky="w")
            controls = ttk.Frame(editor)
            controls.grid(row=1, column=2, sticky="e")
            ttk.Button(controls, text="Ver tabela no PDF", command=self.view_table).pack(side="left", padx=(0, 8))
            ttk.Button(controls, text="Conferi a tabela no PDF", command=self.review_table).pack(side="left", padx=(0, 8))
            ttk.Button(controls, text="Aplicar correção", command=self.edit).pack(side="left")
            bottom = ttk.Frame(root, padding=(10, 0, 10, 10))
            bottom.grid(row=3, column=0, sticky="ew")
            ttk.Label(bottom, textvariable=self.status, wraplength=830).pack(side="left", fill="x", expand=True)
            ttk.Button(bottom, text="Conferi início e fim", command=self.review_limits).pack(side="right", padx=(8, 0))
            ttk.Button(bottom, text="Salvar Word…", command=self.save).pack(side="right", padx=(8, 0))

        def pick_pdf(self) -> None:
            path = filedialog.askopenfilename(title="PDF do diário oficial", filetypes=[("PDF", "*.pdf")])
            if path:
                self.pdf_var.set(path)

        def pick_model(self) -> None:
            path = filedialog.askopenfilename(title="Modelo Word da SEFA", filetypes=[("Word", "*.docx")])
            if path:
                self.model_var.set(path)

        def change_case(self) -> None:
            self.selection = None
            self.role.set("")
            self.edit_text.delete("1.0", "end")
            self.refresh()

        def analyze(self) -> None:
            path = Path(self.pdf_var.get().strip())
            query = self.act_var.get().strip()
            if not path.is_file():
                messagebox.showerror("PDF", "Escolha um arquivo PDF existente.")
                return
            try:
                max_pages = int(self.max_pages_var.get())
                if not 1 <= max_pages <= 100:
                    raise ValueError("Informe um número entre 1 e 100 páginas.")
                self.extraction = extract_act(path, query, max_pages=max_pages)
            except Exception as error:
                messagebox.showerror("Não foi possível analisar o PDF", str(error))
                return
            self.read_pdf, self.read_act = self.pdf_var.get(), query
            self.read_max_pages = self.max_pages_var.get()
            self.selection = None
            self.refresh()

        def refresh(self) -> None:
            if not self.extraction:
                return
            self.paragraph_tree.delete(*self.paragraph_tree.get_children())
            self.cell_tree.delete(*self.cell_tree.get_children())
            for idx, block in enumerate(self.extraction.paragraphs):
                self.paragraph_tree.insert("", "end", iid=str(idx),
                                           values=(block.page, block.kind, block.role, block.text[:600]),
                                           tags=("attention",) if block.attention else ())
            for ti, table in enumerate(self.extraction.tables):
                presented, _ = formatted_table_matrix(table, self.municipalities_var.get())
                for ri, row in enumerate(presented):
                    for ci, value in enumerate(row):
                        self.cell_tree.insert("", "end", iid=f"{ti}:{ri}:{ci}",
                                              values=(ti + 1, ri + 1, ci + 1, value[:700]),
                                              tags=("attention",) if table.issues else ())
            n = sum(sum(t.source_numbers.values()) for t in self.extraction.tables)
            errors = check_table_numbers(self.extraction)
            self.status.set(f"Páginas {', '.join(map(str, self.extraction.pages))}. "
                            f"{len(self.extraction.paragraphs)} trechos, {len(self.extraction.tables)} tabelas, "
                            f"{n} valores monetários lidos do PDF. " +
                            ("Revise: " + errors[0] if errors else "Confira o texto e a diagramação no Word.") +
                            (" BLOQUEIO: " + "; ".join(self.extraction.blocking_issues)
                             if self.extraction.blocking_issues else "") +
                            (" CONFIRA OS LIMITES: " + "; ".join(self.extraction.issues)
                             if self.extraction.issues and not self.extraction.limits_reviewed else ""))

        def selected(self, section: str) -> None:
            tree = self.paragraph_tree if section == "paragraph" else self.cell_tree
            if not tree.selection() or not self.extraction:
                return
            if section == "paragraph":
                index = int(tree.selection()[0])
                block = self.extraction.paragraphs[index]
                self.selection = (section, index, 0, 0)
                self.role.set(block.role)
                if block.attention:
                    self.status.set("Confira este trecho: " + "; ".join(block.attention))
            else:
                ti, ri, ci = map(int, tree.selection()[0].split(":"))
                block = self.extraction.tables[ti]
                self.selection = (section, ti, ri, ci)
                self.role.set("")
                if block.issues:
                    self.status.set("Confira esta tabela: " + "; ".join(block.issues))
            value = (block.text if section == "paragraph" else
                     formatted_table_matrix(block, self.municipalities_var.get())[0][ri][ci])
            self.edit_text.delete("1.0", "end")
            self.edit_text.insert("1.0", value)

        def edit(self) -> None:
            if not self.extraction or not self.selection:
                messagebox.showinfo("Selecione um item", "Selecione um parágrafo ou uma célula na prévia.")
                return
            section, first, row, col = self.selection
            value = self.edit_text.get("1.0", "end-1c").strip()
            if section == "paragraph":
                block = self.extraction.paragraphs[first]
                block.text, block.role = value, self.role.get() or block.role
                block.attention.append("Trecho corrigido manualmente na prévia.")
                current_tree, current_id = self.paragraph_tree, str(first)
            else:
                table = self.extraction.tables[first]
                table.matrix[row][col] = value
                table.reviewed = False
                current_tree, current_id = self.cell_tree, f"{first}:{row}:{col}"
            self.refresh()
            current_tree.selection_set(current_id)
            current_tree.see(current_id)

        def view_table(self) -> None:
            if not self.extraction or not self.selection or self.selection[0] != "cell":
                messagebox.showinfo("Selecione uma célula", "Selecione uma célula da tabela para ver seu recorte no PDF.")
                return
            table = self.extraction.tables[self.selection[1]]
            row_index = self.selection[2]
            source_part = next((part for part in table_page_parts(table)
                                if part.first_row <= row_index < part.last_row),
                               table_page_parts(table)[0])
            with fitz.open(str(self.extraction.pdf_path)) as pdf:
                page = pdf[source_part.page - 1]
                box = fitz.Rect(*source_part.bbox) + (-15, -15, 15, 15)
                pix = page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), clip=box & page.rect, alpha=False)
                picture = tk.PhotoImage(data=base64.b64encode(pix.tobytes("png")).decode("ascii"))
            window = tk.Toplevel(self.root)
            location = f" · coluna {source_part.column}" if source_part.column else ""
            window.title(f"Tabela {self.selection[1] + 1} · página {source_part.page}{location} do PDF")
            window.geometry("950x700")
            frame = ttk.Frame(window)
            frame.pack(fill="both", expand=True)
            canvas = tk.Canvas(frame, scrollregion=(0, 0, picture.width(), picture.height()))
            canvas.pack(side="left", fill="both", expand=True)
            bar = ttk.Scrollbar(frame, orient="vertical", command=canvas.yview)
            bar.pack(side="right", fill="y")
            canvas.configure(yscrollcommand=bar.set)
            canvas.create_image(0, 0, image=picture, anchor="nw")
            window.picture = picture  # mantém a imagem visível enquanto a janela estiver aberta

        def review_table(self) -> None:
            if not self.extraction or not self.selection or self.selection[0] != "cell":
                messagebox.showinfo("Selecione uma célula", "Selecione uma célula da tabela que foi conferida com o PDF.")
                return
            ti = self.selection[1]
            table = self.extraction.tables[ti]
            if messagebox.askyesno("Revisão manual", f"Você comparou cada rótulo e cada valor da tabela {ti + 1} com o PDF original? "
                                   "Diferenças automáticas serão aceitas após sua conferência; a revisão ficará registrada no relatório, se for gerado."):
                table.reviewed = True
                self.refresh()

        def review_limits(self) -> None:
            if not self.extraction:
                messagebox.showinfo("Analise o PDF", "Clique em Analisar PDF antes de conferir o ato.")
                return
            if messagebox.askyesno("Início e fim do ato", "Você conferiu no PDF que o título e todas as páginas "
                                   "do ato foram incluídos, até sua assinatura, protocolo ou início do ato seguinte?"):
                self.extraction.limits_reviewed = True
                self.refresh()

        def save(self) -> None:
            if (not self.extraction or self.pdf_var.get() != self.read_pdf or
                    self.act_var.get().strip() != self.read_act or self.max_pages_var.get() != self.read_max_pages):
                messagebox.showinfo("Analise o PDF", "Escolha o PDF e clique em Analisar PDF antes de salvar.")
                return
            suggested = re.sub(r"[^\w.-]+", "_", self.read_act).strip("_") + "_SEFA.docx"
            chosen = filedialog.asksaveasfilename(title="Salvar ato completo", defaultextension=".docx",
                                                   initialfile=suggested, filetypes=[("Word", "*.docx")])
            if not chosen:
                return
            output = Path(chosen)
            report = output.with_name(output.stem + "_CONFERIR.html") if self.report_var.get() else None
            force = False
            if output.exists() or (report and report.exists()):
                force = messagebox.askyesno("Arquivos existentes", "Substituir o Word ou o relatório existente?")
                if not force:
                    return
            try:
                export(self.extraction, Path(self.model_var.get().strip()), output, report, force,
                       self.municipalities_var.get())
            except Exception as error:
                messagebox.showerror("Não foi possível gerar o Word", str(error))
                return
            arquivos = f"Word: {output}" + (f"\nRelatório: {report}" if report else "")
            if not self.links_var.get():
                show_and_open(messagebox, "Arquivos criados", f"{arquivos}\n\nCompare com o PDF antes de usar.", output)
                return
            executar_com_progresso(
                self.root, "Criando hyperlinks de legislação",
                lambda log, parar: aplicar_e_registrar(output, report, log, parar),
                lambda resumo, erro: show_and_open(
                    messagebox, "Arquivos criados",
                    f"{arquivos}\n\n{texto_final(resumo, erro, report is not None)}\n\n"
                    "Compare com o PDF antes de usar.", output))

    App().root.mainloop()


def main() -> int:
    parser = argparse.ArgumentParser(description="Importa um ato e tabelas do PDF para Word editável da SEFA.")
    parser.add_argument("--pdf", type=Path, help="PDF do diário oficial; sem este argumento, abre uma janela")
    parser.add_argument("--act", help="Trecho único da epígrafe, por exemplo 494/2026-SEFA.GS")
    parser.add_argument("--model", type=Path, default=DEFAULT_MODEL)
    parser.add_argument("--max-pages", type=int, default=12, help="Limite de páginas pesquisadas para o ato (padrão: 12)")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--relatorio", action="store_true",
                        help="Gera o relatório HTML de conferência ao lado do Word")
    parser.add_argument("--report", type=Path, help="Nome do relatório HTML (implica --relatorio)")
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--preservar-grafia", action="store_true",
                        help="Mantém maiúsculas/minúsculas do PDF nas colunas de municípios")
    parser.add_argument("--confirmar-limites", action="store_true",
                        help="Confirma que o começo e o fim do ato foram conferidos no PDF")
    parser.add_argument("--sem-links", action="store_true",
                        help="Não cria os hyperlinks das leis, decretos etc. citados")
    args = parser.parse_args()
    try:
        if args.pdf:
            if not args.act or not args.output:
                parser.error("--pdf requer --act e --output")
            extraction = extract_act(args.pdf, args.act, max_pages=args.max_pages)
            extraction.limits_reviewed = args.confirmar_limites
            report = None
            if args.relatorio or args.report:
                report = args.report or args.output.with_name(args.output.stem + "_CONFERIR.html")
            export(extraction, args.model, args.output, report, args.force,
                   not args.preservar_grafia)
            print(f"Word: {args.output}" + (f"\nRelatório: {report}" if report else "") +
                  f"\nPáginas: {extraction.pages}; tabelas: {len(extraction.tables)}; "
                  f"conferir: {len(extraction.issues)}")
            if not args.sem_links:
                resumo = aplicar_e_registrar(args.output, report)
                print(texto_final(resumo, None, com_relatorio=report is not None))
        else:
            run_gui()
    except Exception as error:
        print(f"Erro: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
