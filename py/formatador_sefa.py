#!/usr/bin/env python3
"""Formata texto normativo colado usando estilos de um documento Word da SEFA.

Funciona sem IA e sem acesso à internet. Não interpreta a vigência nem altera
o texto jurídico. A classificação pode ser corrigida antes da exportação.
"""

from __future__ import annotations

import argparse
import html
import os
import re
import subprocess
from collections import Counter
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from links_no_word import aplicar_e_registrar, executar_com_progresso, texto_final


HERE = Path(__file__).resolve().parent
DEFAULT_MODEL = HERE / "modelo_sefa.docx"

STYLES = {
    "epigrafe": ("a01_epigrafe",),
    "publicacao": ("a02_publicacao",),
    "ementa": ("a03_ementa",),
    "ordem": ("a04_ordem_preamb",),
    "autoridade": ("a05_autoridade",),
    "cargo": ("a06_cargo",),
    "texto": ("a06_texto", "a20_texto"),
    "titulo": ("a10_titulo", "a08_titulo"),
    "anexo": ("a10_titulo", "a01_epigrafe_an"),
    "subtitulo": ("a11_subtitulo",),
    "redacao_nova": ("a22_nova_redacao",),
    "remissao": ("a30_remiss_nova_red", "a10_nr_remiss"),
    "remissao_antiga": ("a31_remiss_red_ant", "a11_rdant_remiss"),
    "redacao_antiga": ("a23_redacao_ant", "a12_rdant"),
}

EPIGRAPH = re.compile(
    r"^(?:LEI(?:\s+COMPLEMENTAR)?|DECRETO(?:\s+LEGISLATIVO)?|"
    r"PORTARIA(?:\s+CONJUNTA)?|INSTRU[CÇ][AÃ]O\s+NORMATIVA|"
    r"RESOLU[CÇ][AÃ]O(?:\s+[A-Z]{2,}(?:/[A-Z]+)?)?|"
    r"ATO\s+COTEPE(?:/ICMS)?|CONV[EÊ]NIO\s+ICMS|AJUSTE\s+SINIEF|"
    r"PROTOCOLO\s+ICMS|"
    # DOU (3.12): "ATO TÉCNICO CONJUNTO RFB/SUFIS/CGIBS/DIRETORIA EXECUTIVA Nº 6, DE ..."
    r"ATO\s+T[EÉ]CNICO(?:\s+CONJUNTO)?|ATO\s+DECLARAT[OÓ]RIO(?:\s+(?:EXECUTIVO|INTERPRETATIVO))?|"
    r"MEDIDA\s+PROVIS[OÓ]RIA|(?-i:DESPACHO))\b.{0,75}\bN[º°O.]?\s*\d[\d./-]*\b",
    re.I,
)
PUBLICATION = re.compile(
    r"^(?:PUBLICAD[OA]|REPUBLICAD[OA]|RETIFICAD[OA]|ALTERAD[OA]\s+PELOS?|"
    r"EDI[CÇ][AÃ]O\s+(?:EXTRA|N[º°O.]?\s*\d)|"
    r"D(?:O\s*[-–]?\s*E|OE|OU)\b|VIDE\s+(?:DECRETO|LEI|IN|PORTARIA))\b",
    re.I,
)
PREAMBLE = re.compile(r"^(?:CONSIDERANDO\b|(?:O|A|OS|AS)\s+.{3,600}?\b(?:NO USO|NO DESEMPENHO|NO EXERC[IÍ]CIO|USANDO|FAZ SABER|COM FUNDAMENTO)\b)", re.I)
# Também com letras espaçadas, comum no DOE: "R E S O L V E:" (Portaria nº 2005/2026-SEASTER e outras).
COMMAND = re.compile(r"^(?:D\s?E\s?C\s?R\s?E\s?T\s?A(?:\s?M)?|R\s?E\s?S\s?O\s?L\s?V\s?E(?:\s?M)?|D\s?E\s?L\s?I\s?B\s?E\s?R\s?A(?:\s?M)?"
                     r"|PROMULGAM?|FAZ SABER|FAZEM SABER)\s*[:.;]?$", re.I)
HEADING = re.compile(r"^(?:PARTE|LIVRO|T[IÍ]TULO|CAP[IÍ]TULO|SE[CÇ][AÃ]O|SUBSE[CÇ][AÃ]O)\s+(?:[IVXLCDM]+|[\dº°-]+)\s*[:.-]?$", re.I)
ANNEX = re.compile(r"^(?:ANEXO|AP[EÊ]NDICE)\s+(?:(?:[IVXLCDM]+|[\dº°-]+|[ÚU]NIC[OA]|[A-Z])(?:\s*[-–].*)?$"
                   r"|(?:[ÀA]|AO|DA|DO)\s+(?:PORTARIA|RESOLU[CÇ][AÃ]O|DECRETO|LEI|INSTRU[CÇ][AÃ]O|ATO|DELIBERA[CÇ][AÃ]O)\b)", re.I)
# Dispositivos: artigo, parágrafo, inciso, alínea e item ("1. Anexo VI ...", LC 95/1998, art. 10, IV)
# Dispositivo do PRÓPRIO ato: rótulo seguido de texto que começa com maiúscula ("Art. 1º Fica", "Art. 14. A",
# "Art. 3º - Esta", "Art. 2º-A. Fica", "§ 1º O", "Parágrafo único. O"). Referência a dispositivo não abre
# parágrafo nem leva negrito: "art. 7º, inciso VI, do Decreto", "§ 1º do art. 3º da Lei", "parágrafo único,
# inciso I", "Art. 27, parágrafo único, que versa", "Art. 24 da Lei" (DOE de 02 e 05.10.2026).
_LABEL_TEXT = r"\s*(?:[º°]|o(?=\s))?(?:\s*-\s*[A-Z](?![A-Za-zÀ-ÿ]))?\s*\.?\s*(?:[-–]\s*)?(?=[A-ZÀ-Ý\"“(]|$)"
ARTICLE = re.compile(r"^(?:(?:Art|ART)\.?\s*\d+" + _LABEL_TEXT + r"|§\s*\d+" + _LABEL_TEXT
                     + r"|(?:Par[aá]grafo|PAR[AÁ]GRAFO)\s+(?:[UuÚú]nico|[UÚ]NICO)\s*[.:–-]?\s*(?=[A-ZÀ-Ý\"“(]|$)"
                     + r"|[IVXLCDM]+\s*[-–]|[A-Za-z]\s*\)|\d{1,2}\.\s+\S)")
LABEL_WITHOUT_SPACE = re.compile(r"^((?:Art|ART)\.?\s*\d+\s*[º°]|§\s*\d+\s*[º°])(?=[A-ZÀ-Ý])")
ARTICLE_LABEL = re.compile(r"^((?:Art|ART)\.?\s*\d+(?:\s*[º°])?(?:\s*[-–]\s*[A-Za-z0-9]+)?\.?)")
# Preâmbulo que termina na própria fórmula de ordem, sem linha "RESOLVE:" separada (DOU)
INLINE_COMMAND_END = re.compile(r"\b(?:RESOLVEM?|DECRETAM?|DELIBERAM?|DETERMINAM?)\s*:\s*$", re.I)
PREAMBLE_AUTHORITY = re.compile(r"^((?:O|A|OS|AS)\s+.+?)(?=,?\s+(?:no uso|no desempenho|no exercício|usando)\b|;\s*$)", re.I)
EDITORIAL_NEW = re.compile(r"^(?:REDA[CÇ][AÃ]O\s+DADA|ACRESCID[OA]|VIDE\s+(?:CONVALIDA[CÇ][AÃ]O|SOLU[CÇ][AÃ]O))\b", re.I)
EDITORIAL_OLD = re.compile(r"^(?:REDA[CÇ][AÃ]O\s+ORIGINAL|REDA[CÇ][AÃ]O\s+ANTERIOR|REVOGAD[OA])\b", re.I)
WEB_NOISE = re.compile(
    r"^(?:P[AÁ]GINA\s+\d+|P[AÁ]G\.?\s*\d+|DOCUMENTO\s+ASSINADO\s+ELETRONICAMENTE|"
    r"ESTE\s+CONTE[UÚ]DO\s+N[AÃ]O\s+SUBSTITUI|DI[AÁ]RIO\s+OFICIAL\s+(?:DA|DO)\b)", re.I
)
LEGAL_QUOTE = re.compile(r'^[“"]\s*(?:ART\.?\s*\d|§\s*\d|PAR[AÁ]GRAFO\b|[IVXLCDM]+\s*[-–]|[a-z]\s*\)|\.{3,})', re.I)
QUOTE_END = re.compile(r'[”"]\s*(?:\(NR\))?\s*[.;]?$' , re.I)
# Cargos de quem assina. CARGO (acima) também abre parágrafo em qualquer ponto do PDF; esta
# lista mais ampla só vale perto de um nome (assinatura), porque "Procurador do Estado",
# "Delegado de Polícia" etc. podem começar uma linha no meio de um artigo.
SIGNER_CARGO = re.compile(
    r"^(?:VICE[-\s])?(?:GOVERNADOR(?:A)?|(?:SUB)?SECRET[AÁ]RI[OA]|PRESIDENTE|MINISTR[OA]|DIRETOR(?:A)?|"
    r"PROCURADOR(?:A)?|DELEGAD[OA]|DEFENSOR(?:A)?|CONTROLADOR(?:A)?|CORREGEDOR(?:A)?|COORDENADOR(?:A)?|"
    r"CHEFE|COMANDANTE|SUPERINTENDENTE|PREFEIT[OA]|OUVIDOR(?:A)?|PROMOTOR(?:A)?|CONSELHEIR[OA]|"
    r"AUDITOR(?:A)?|REITOR(?:A)?)\b", re.I)
# Palavras de órgão: "Estado do Pará" (continuação de cargo) não é nome de gente.
INSTITUTION_WORDS = {"estado", "pará", "para", "fazenda", "justiça", "justica", "secretaria", "governo",
                     "polícia", "policia", "civil", "militar", "pública", "publica", "procuradoria", "ministério",
                     "ministerio", "receita", "tribunal", "assembleia", "legislativa", "república", "republica",
                     "federal", "estadual", "municipal", "geral", "adjunto", "adjunta", "executivo", "executiva"}
# Barra seguida de espaço antes do ano em número de ato: "nº 041/ 2002" (DOE de 02.10.2026).
# Só com marca de número e ano de 4 dígitos que não seja parte de data ("nº 5.810/ 24.01.1994"
# fica como está) nem fora de citação ("competência 12/ 2025" fica como está).
SLASH_YEAR_SPACE = re.compile(r"(\b[nN](?:\.?[º°oO]|\.)\.?[ \u00a0]*\d[\d.]*)/[ \u00a0]+((?:19|20)\d{2})(?!\d)(?![./-]\d)")
NAME = re.compile(r"^[A-ZÀ-Ý][A-ZÀ-Ý'-]+(?:\s+[A-ZÀ-Ý][A-ZÀ-Ý'-]+){1,5}$")
CARGO = re.compile(r"^(?:GOVERNADOR[AA]?|SECRET[AÁ]RI[OA]|PRESIDENTE|MINISTR[OA]|DIRETOR[AA]?)\b", re.I)
# Os quatro últimos verbos só com inicial maiúscula: "disciplina", "define" etc. podem
# começar uma linha no meio de um artigo do PDF sem abrir parágrafo.
EMENTA_START = re.compile(r"^(?:ABRE|ALTERA|DISP[OÕ]E|INSTITUI|REGULAMENTA|ESTABELECE|APROVA|REVOGA|HOMOLOGA|PRORROGA|MODIFICA|ACRESCENTA"
                          r"|(?-i:Disciplina|DISCIPLINA|Divulga|DIVULGA|Define|DEFINE|Fixa|FIXA))\b", re.I)


@dataclass
class Block:
    text: str
    first_line: int
    last_line: int
    blank_before: bool = False


@dataclass
class Item:
    text: str
    first_line: int
    last_line: int
    role: str
    reason: str
    attention: list[str] = field(default_factory=list)
    style: str = ""
    overridden: bool = False


def open_file(path: Path) -> str | None:
    """Abre o arquivo no programa padrão do computador (o Word, para .docx).
    Devolve a mensagem de erro, ou None se deu certo."""
    try:
        if sys.platform.startswith("win"):
            os.startfile(str(path))  # type: ignore[attr-defined]
        else:
            subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", str(path)],
                             stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except Exception as error:
        return str(error)
    return None


def show_and_open(messagebox, title: str, text: str, path: Path) -> None:
    """Mensagem final das janelas; ao clicar em OK, o Word gerado é aberto. A mensagem vem
    antes para que a janela do Word não cubra o resumo (links pendentes, avisos)."""
    messagebox.showinfo(title, f"{text}\n\nAo clicar em OK, o documento será aberto.")
    error = open_file(path)
    if error:
        messagebox.showwarning("Não foi possível abrir o documento",
                               f"O arquivo foi salvo, mas não abriu automaticamente:\n{path}\n\n{error}")


def looks_uppercase(text: str) -> bool:
    letters = [c for c in text if c.isalpha()]
    return bool(letters) and len(text) <= 140 and all(c == c.upper() for c in letters)


AUTHORITY_START = re.compile(
    r"^(?:O|A|OS|AS)\s+(?:VICE[-\s])?(?:GOVERNADOR|(?:SUB)?SECRET[AÁ]RI|PROCURADOR|PRESIDENTE|DIRETOR|DELEGAD|"
    r"DEFENSOR|CONTROLADOR|CORREGEDOR|COMANDANTE|SUPERINTENDENTE|PREFEIT|MINISTR|COORDENADOR|CHEFE|AUDITOR|"
    r"OUVIDOR|PROMOTOR|CONSELHEIR)", re.I)


def is_preamble(text: str) -> bool:
    """Preâmbulo: "O SECRETÁRIO ..., no uso ..." ou "Considerando ...", e também a lista de autoridades
    sem "no uso", terminada em ";" ("A Procuradora-Geral do Estado, o Secretário de Estado da Fazenda,
    ... e o Presidente da PRODEPA;" - Portaria Conjunta nº 5/2026, DOE de 05.10.2026)."""
    return bool(PREAMBLE.match(text) or (AUTHORITY_START.match(text) and text.rstrip().endswith((";", ":"))))


# Hífen no fim da linha: separação silábica ("Se-" / "cretário") ou hífen da palavra ("cumpra-" / "se").
LINE_END_HYPHEN = re.compile(r"([A-Za-zÀ-ÖØ-öø-ÿ]+)-$")
WORD = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ]+(?:-[A-Za-zÀ-ÖØ-öø-ÿ]+)*")
HYPHEN_PREFIXES = {"vice", "pós", "pré", "pró", "recém", "além", "aquém", "grão", "grã"}
VOWEL_PREFIXES = {"auto", "anti", "contra", "extra", "infra", "intra", "supra", "ultra", "semi", "micro", "macro",
                  "mini", "multi", "neo", "proto", "pseudo", "sobre", "arqui", "entre"}
WEEKDAYS = {"segunda", "terça", "quarta", "quinta", "sexta"}


def vocabulary(lines: list[str]) -> Counter:
    """Palavras inteiras do texto (minúsculas), sem os pedaços de palavra partidos no fim da linha."""
    counts: Counter = Counter()
    broken = False
    for line in lines:
        words = WORD.findall(line)
        if broken and words:
            words = words[1:]
        broken = bool(LINE_END_HYPHEN.search(line.rstrip()))
        if broken and words:
            words = words[:-1]
        counts.update(w.lower() for w in words)
    return counts


def keep_line_hyphen(left: str, right: str, vocab: Counter | None = None) -> bool:
    """left: pedaço antes do hífen do fim da linha; right: começo da linha seguinte.
    True se o hífen é da palavra ("cumpra-se", "sexta-feira", "Delegado-Geral"); False se é separação
    silábica, a imensa maioria nos diários ("Se-cretário", "ADMINIS-TRATIVAS", "étni-co-racial")."""
    word = (WORD.match(right) or [""])[0]
    if not word:
        return False
    first = word.split("-")[0]
    if vocab:
        joined, hyphened = (left + word).lower(), (left + "-" + word).lower()
        if vocab[hyphened] > vocab[joined]:
            return True
        if vocab[joined]:
            return False
    low_left, low_first = left.lower(), first.lower()
    if low_first in ("se", "lhe", "lhes") and not low_left.endswith("s"):
        return True                                            # cumpra-se, aplica-se, atribuir-lhe
    if low_first in ("lo", "la", "los", "las") and low_left[-1:] in "áâêéíóô":
        return True                                            # concedê-lo, torná-la
    if low_first in ("no", "na", "nos", "nas") and low_left.endswith(("ão", "õe", "m")):
        return True                                            # dão-no, fazem-na
    if low_left in WEEKDAYS and low_first.startswith("feira"):
        return True
    if low_left in HYPHEN_PREFIXES:
        return True
    if low_left in VOWEL_PREFIXES and (low_first[:1] == "h" or low_first[:1] == low_left[-1:]):
        return True                                            # anti-inflamatório, micro-ondas
    if first[:1].isupper() and not (left.isupper() and first.isupper()):
        return True                                            # Delegado-Geral, Econômico-Fiscais
    return False


def join_line(previous: str, line: str, vocab: Counter | None = None) -> str:
    """Junta a linha seguinte ao parágrafo: sem espaço depois de barra ("SEIRDH/" + "PRODEPA") e de
    identificador com hífen ("5/2026-" + "PGE"); hífen silábico retirado; nos demais casos, um espaço."""
    head = previous.rstrip()
    found = LINE_END_HYPHEN.search(head)
    if found and line[:1].isalpha():
        return head + line if keep_line_hyphen(found.group(1), line, vocab) else head[:-1] + line
    if head.endswith("/") or (re.search(r"\d-$", head) and line[:1].isalpha()):
        return head + line
    return f"{head} {line}"


# "DÊ-SE CIÊNCIA, REGISTRE-SE, PUBLIQUE-SE E CUMPRA-SE." e "CASA MILITAR DA GOVERNADORIA DO ESTADO, 1º DE
# OUTUBRO DE 2026." são fórmula de encerramento e local/data, não linhas a conferir.
CLOSING = re.compile(r"^(?:D[ÊE]-SE CI[ÊE]NCIA|CIENTIFIQUE-SE|REGISTRE-SE|PUBLIQUE-SE|CUMPRA-SE)\b", re.I)
PLACE_DATE_LINE = re.compile(r"^[^,;:]{3,140},\s*(?:EM\s+)?\d{1,2}[º°]?\s+(?:DE\s+)?[A-ZÇ]+\s+(?:DE\s+)?\d{4}\.?$", re.I)


def looks_like_signer_name(text: str) -> bool:
    """Nome de signatário em linha própria, em maiúsculas ou não ("Ana Carolina Lobo Gluck Paúl").
    Mais estrito que looks_like_person_name: começa por palavra com inicial maiúscula e não
    contém palavras de órgão, para não confundir "Estado do Pará" (resto de um cargo) com nome."""
    words = text.split()
    return (looks_like_person_name(text) and words[0][0].isupper()
            and words[0].casefold() not in {"de", "da", "do", "das", "dos", "e"}
            and not any(w.casefold() in INSTITUTION_WORDS for w in words))


EPIGRAPH_YEAR_END = re.compile(r"\b(?:19|20)\d{2}\s*\.?\s*$")


def continues_epigraph(epigraph: str, line: str) -> bool:
    """Segunda linha de uma epígrafe quebrada no diário. Ex.: "INSTRUÇÃO NORMATIVA Nº 024 -"
    seguida de "SEFA.GS, DE 01 DE OUTUBRO DE 2026" (DO-e/SEFA de 02.10.2026). Só une quando
    a epígrafe ainda não terminou no ano e a linha seguinte, curta e em maiúsculas, traz o
    ano no fim ou completa um "-", "," ou "/" deixado em aberto."""
    if not EPIGRAPH.match(epigraph) or EPIGRAPH_YEAR_END.search(epigraph):
        return False
    if len(line) > 90 or not looks_uppercase(line) or looks_like_person_name(line):
        return False
    if any(rx.match(line) for rx in (EPIGRAPH, COMMAND, PREAMBLE, ARTICLE, HEADING, ANNEX, CARGO)):
        return False
    return bool(EPIGRAPH_YEAR_END.search(line)) or epigraph.rstrip().endswith(("-", "–", ",", "/"))


SIGNER_SUFFIX = re.compile(
    r"\s*[-–]\s*(?:(?:CEL|TC|TEN(?:ENTE)?[\s-]*CEL|MAJ|CAP|[12][º°]?\s*TEN|SUBTEN|[123][º°]?\s*SGT|CB|SD)\b"
    r"|(?:RG|MF|MAT(?:R[IÍ]CULA)?|CPF|ID\.?\s*FUNC(?:IONAL)?)\b).*$", re.I)


def looks_like_person_name(text: str) -> bool:
    """Aceita nomes em maiúsculas ou com iniciais maiúsculas; exige contexto na classificação. Ignora posto
    e registro depois do nome ("OSMAR VIEIRA DA COSTA JÚNIOR - CEL QOPM RG 9916", DOE de 02.10.2026)."""
    text = SIGNER_SUFFIX.sub("", text)
    words = text.split()
    return (2 <= len(words) <= 8 and not CARGO.match(text) and not SIGNER_CARGO.match(text)
            and not HEADING.match(text) and not ANNEX.match(text)
            and all(re.fullmatch(r"[^\W\d_]+(?:['’\-][^\W\d_]+)*", word)
                    and (word.casefold() in {"de", "da", "do", "das", "dos", "e"}
                         or word[0].isupper()) for word in words))


def ementa_spacer_before(items: list[Item]) -> Item | None:
    """Reproduz o separador do exemplo revisado, somente na ausência de ementa."""
    if not any(item.role == "epigrafe" for item in items):
        return None
    for index, item in enumerate(items):
        if item.role == "texto" and is_preamble(item.text):
            prefix = items[:index]
            return item if (any(p.role == "publicacao" for p in prefix)
                            and not any(p.role == "ementa" for p in prefix)) else None
        if ARTICLE.match(item.text) or item.role == "ordem":
            break
    return None


# Linha que termina em abreviatura ("identidade funcional n.") continua na seguinte
# quando esta começa por número ("57199022/3 - titular; e").
ENDS_WITH_ABBREVIATION = re.compile(r"(?:^|[\s(])(?:n|n[º°]|art|arts|inc|fl|fls|pág|p)\.$", re.I)


def logical_blocks(source: str, join_soft_wraps: bool = False) -> list[Block]:
    """Apenas separa linhas; juntar quebra de página/coluna requer revisão."""
    source = source.replace("\r\n", "\n").replace("\r", "\n")
    vocab = vocabulary(source.split("\n")) if join_soft_wraps else None
    blocks: list[Block] = []
    blank = False
    for line_number, raw in enumerate(source.split("\n"), start=1):
        line = raw.strip().replace("\u00a0", " ")
        if not line:
            blank = True
            continue
        if join_soft_wraps and blocks and not blank and (
            (re.match(r"^[a-zà-ÿ]", line)
             and not re.search(r'[.!?:;”"]\s*$', blocks[-1].text)
             and not re.match(r"^[a-z]\s*\)", line))
            or (ENDS_WITH_ABBREVIATION.search(blocks[-1].text) and line[:1].isdigit())
        ):
            blocks[-1].text = join_line(blocks[-1].text, line, vocab)
            blocks[-1].last_line = line_number
        else:
            blocks.append(Block(line, line_number, line_number, blank))
        blank = False
    return blocks


def classify(source: str, join_soft_wraps: bool = False) -> list[Item]:
    blocks = logical_blocks(source, join_soft_wraps)
    items: list[Item] = []
    front = True
    found_epigraph = False
    seen_preamble = False
    in_quote = False
    last_heading = False
    after_palace = False
    last_was_name = False
    skipped: set[int] = set()

    for block_index, block in enumerate(blocks):
        if block_index in skipped:
            continue
        t = block.text
        u = t.upper()
        role, reason = "texto", "Texto comum"
        attention: list[str] = []
        epigraph_joined = False
        inline_command = None

        if EPIGRAPH.match(t) and (front or not found_epigraph):
            following = block_index + 1
            while following < len(blocks) and continues_epigraph(t, blocks[following].text):
                t = join_line(t, blocks[following].text)
                block.text, block.last_line = t, blocks[following].last_line
                skipped.add(following)
                following += 1
                epigraph_joined = True
            u = t.upper()
            role, reason = "epigrafe", "Tipo, número e data do ato"
            found_epigraph = True
            front = True
            seen_preamble = False
            in_quote = False
        elif front and PUBLICATION.match(t):
            role, reason = "publicacao", "Linha de publicação ou referência antes do texto"
        elif COMMAND.match(t):
            role, reason = "ordem", "Fórmula que abre a parte normativa"
            front = False
        elif in_quote or LEGAL_QUOTE.match(t):
            role, reason = "redacao_nova", "Trecho normativo entre aspas"
            in_quote = not bool(QUOTE_END.search(t))
        elif ANNEX.match(t):
            role, reason = "anexo", "Identificador de anexo ou apêndice"
            front = False
        elif HEADING.match(t):
            role, reason = "titulo", "Divisão do ato normativo"
            front = False
        elif last_heading and looks_uppercase(t) and not ARTICLE.match(t) and not t.endswith(('.', ';', ':')):
            role, reason = "titulo", "Nome da divisão logo após o título"
        elif (looks_like_person_name(t) and not front
              and (after_palace or (block_index + 1 < len(blocks)
                                   and SIGNER_CARGO.match(blocks[block_index + 1].text)))):
            role, reason = "autoridade", "Assinatura após local/data ou seguida de cargo"
        elif last_was_name and SIGNER_CARGO.match(t):
            role, reason = "cargo", "Cargo após a assinatura"
        elif front and is_preamble(t):
            role, reason = "texto", "Preâmbulo ou considerando"
            seen_preamble = True
            if INLINE_COMMAND_END.search(t):      # DOU: "... no uso das atribuições ..., resolvem:"
                front = False
                inline_command = INLINE_COMMAND_END.search(t)
        elif front and seen_preamble and not ARTICLE.match(t):
            role, reason = "texto", "Continuação do preâmbulo"
            if INLINE_COMMAND_END.search(t):      # última linha do preâmbulo copiado: "..., resolvem:"
                front = False
                inline_command = INLINE_COMMAND_END.search(t)
            if not PREAMBLE.match(t):
                attention.append("Confira se esta linha continua o preâmbulo.")
        elif front and found_epigraph and not ARTICLE.match(t):
            role, reason = "ementa", "Texto entre epígrafe e preâmbulo"
            if not EMENTA_START.match(t) and not (items and items[-1].role == "ementa"):
                attention.append("Confirme se este trecho pertence à ementa.")
        elif front and not found_epigraph and EMENTA_START.match(t):
            role, reason = "ementa", "Possível ementa sem epígrafe anterior"
            attention.append("Não foi encontrada uma epígrafe anterior.")
        elif ARTICLE.match(t):
            role, reason = "texto", "Artigo ou subdivisão do dispositivo"
            front = False
        elif not front and EDITORIAL_NEW.match(t):
            role, reason = "remissao", "Possível nota editorial sobre redação vigente"
            attention.append("Confira se é uma nota editorial, e não parte do ato publicado.")
        elif not front and EDITORIAL_OLD.match(t):
            role, reason = "remissao_antiga", "Possível nota editorial sobre redação anterior"
            attention.append("Confira se a nota é histórica ou se trata de uma revogação vigente.")
        elif WEB_NOISE.match(t):
            attention.append("Possível cabeçalho, rodapé ou aviso da fonte. O texto foi preservado.")
        elif not front and (CLOSING.match(t) or PLACE_DATE_LINE.match(t)):
            reason = "Fórmula de encerramento ou local e data"
        elif not front and looks_uppercase(t) and len(t) < 120:
            attention.append("Linha em maiúsculas não reconhecida como título ou assinatura.")
        elif front and not found_epigraph:
            attention.append("Início do ato não identificado; confirme o estilo.")
        else:
            reason = "Texto comum ou continuação; confira a separação de parágrafos"
            if not ARTICLE.match(t) and not PREAMBLE.match(t):
                attention.append("Confira se a quebra de linha corresponde a um parágrafo.")

        if LABEL_WITHOUT_SPACE.match(t):
            t = block.text = LABEL_WITHOUT_SPACE.sub(r"\1 ", t, count=1)
            attention.append("Espaço inserido depois do rótulo do dispositivo (estava sem espaço no diário).")
        if role != "epigrafe" and SLASH_YEAR_SPACE.search(t):
            fixed = SLASH_YEAR_SPACE.sub(r"\1/\2", t)
            examples = ", ".join(sorted({m.group(0) for m in SLASH_YEAR_SPACE.finditer(t)}))
            attention.append(f"Espaço retirado depois da barra no número do ato ({examples}).")
            t = block.text = fixed
        if epigraph_joined:
            attention.append("A epígrafe estava em duas linhas na origem e foi reunida; confira.")
        elif block.last_line > block.first_line:
            attention.append("Linhas da origem foram reunidas; confira o texto.")
        command = INLINE_COMMAND_END.search(t) if inline_command else None
        if command and t[:command.start()].strip():
            items.append(Item(t[:command.start()].rstrip(), block.first_line, block.last_line, role, reason, attention))
            item = Item(command.group(0).strip(), block.last_line, block.last_line, "ordem",
                        "Fórmula que abre a parte normativa (estava no fim do preâmbulo)")
        else:
            item = Item(t, block.first_line, block.last_line, role, reason, attention)
        items.append(item)

        last_heading = role == "titulo" and bool(HEADING.match(t))
        last_was_name = role == "autoridade"
        after_palace = bool(re.match(r"^(?:PAL[AÁ]CIO|GABINETE)\b", u)
                            or re.match(r"^[^,;:]+,\s*\d{1,2}\s+de\s+\w+\s+de\s+\d{4}[.]?$", t, re.I)) or last_was_name
        if role == "cargo":
            after_palace = False

    if in_quote and items:
        items[-1].attention.append("Aspas de redação nova não foram fechadas no texto fornecido.")
    return items


def paragraph_style_names(model: Document) -> set[str]:
    return {s.name for s in model.styles if s.type == WD_STYLE_TYPE.PARAGRAPH}


def apply_styles(items: list[Item], model: Document) -> None:
    available = paragraph_style_names(model)
    for item in items:
        if item.overridden:
            if item.style not in available:
                raise ValueError(f"O modelo não contém o estilo escolhido: {item.style}")
            continue
        candidates = STYLES[item.role]
        item.style = next((name for name in candidates if name in available), "")
        if not item.style:
            item.style = next((n for n in STYLES["texto"] if n in available), "Normal")
            notice = f"O modelo não possui o estilo {candidates[0]}; foi aplicado {item.style}."
            if notice not in item.attention:
                item.attention.append(notice)


def clear_model_contents(model: Document) -> None:
    """Usa estilos e página do modelo, sem levar o conteúdo do ato de origem."""
    for element in list(model._element.body):
        if not element.tag.endswith("}sectPr"):
            model._element.body.remove(element)
    for section in model.sections:
        for part in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            for paragraph in part.paragraphs:
                paragraph.text = ""
            for table in part.tables:
                table._element.getparent().remove(table._element)
    model.core_properties.title = "Documento formatado no padrão SEFA"
    model.core_properties.subject = ""
    model.core_properties.author = ""
    model.core_properties.last_modified_by = ""
    model.core_properties.comments = ""


def add_formatted_paragraph(model: Document, item: Item):
    """Aplica o estilo e os negritos estruturais encontrados nos exemplos SEFA."""
    p = model.add_paragraph(style=item.style)
    label = ARTICLE_LABEL.match(item.text) if item.role == "texto" and ARTICLE.match(item.text) else None
    if not label and item.role == "texto":
        label = PREAMBLE_AUTHORITY.match(item.text)
    if label:
        p.add_run(label.group(1)).bold = True
        p.add_run(item.text[label.end():])
    else:
        p.add_run(item.text)
    return p


def build_document(items: list[Item], model_path: Path, output_path: Path, overwrite: bool = False) -> None:
    if not items:
        raise ValueError("Cole ou forneça algum texto antes de gerar o Word.")
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"O arquivo já existe: {output_path}")
    if model_path.resolve() == output_path.resolve():
        raise ValueError("O arquivo de saída deve ser diferente do modelo. Escolha outro nome.")
    model = Document(str(model_path))
    apply_styles(items, model)
    clear_model_contents(model)
    spacer = ementa_spacer_before(items) if "a03_ementa" in paragraph_style_names(model) else None
    expected = []
    for item in items:
        if item is spacer:
            model.add_paragraph(style="a03_ementa")
            expected.append(("", "a03_ementa"))
        add_formatted_paragraph(model, item)
        expected.append((item.text, item.style))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix="sefa_", suffix=".docx", dir=output_path.parent, delete=False) as tmp:
        temp_path = Path(tmp.name)
    try:
        model.save(str(temp_path))
        # Verifica a quantidade e o texto antes de disponibilizar a saída.
        produced = Document(str(temp_path))
        actual = [(p.text, p.style.name) for p in produced.paragraphs]
        if actual != expected:
            raise RuntimeError("A conferência do Word produzido falhou; o arquivo de destino não foi alterado.")
        os.replace(temp_path, output_path)
    finally:
        temp_path.unlink(missing_ok=True)


def write_report(items: list[Item], report_path: Path) -> None:
    review = sum(bool(item.attention) for item in items)
    rows = []
    for index, item in enumerate(items, 1):
        info = "<br>".join(html.escape(note) for note in item.attention) or "—"
        css_class = ' class="review"' if item.attention else ""
        rows.append(
            f"<tr{css_class}><td>{index}</td><td>{item.first_line}–{item.last_line}</td>"
            f"<td><code>{html.escape(item.style)}</code></td>"
            f"<td>{html.escape(item.text)}</td><td>{info}</td></tr>"
        )
    page = """<!doctype html><html lang="pt-BR"><meta charset="utf-8">
<title>Conferência da formatação SEFA</title><style>
body{font:15px/1.5 Arial,sans-serif;max-width:1400px;margin:35px auto;padding:0 20px;color:#222}
table{border-collapse:collapse;width:100%;table-layout:fixed} th,td{border:1px solid #bbb;padding:8px;vertical-align:top;overflow-wrap:anywhere}
th{background:#eee;text-align:left}.review{background:#fff4dc}td:nth-child(1){width:35px}td:nth-child(2){width:60px}
td:nth-child(3){width:160px}td:nth-child(5){width:260px}code{font-size:13px}h1{font-size:24px}
@media print{body{margin:12mm}.review{print-color-adjust:exact}}
</style><h1>Conferência da formatação SEFA</h1>
<p>__COUNT__ trechos formatados; __REVIEW__ com pontos para conferir. O programa preservou todas as linhas não vazias.
As cores indicam apenas dúvida de formatação ou segmentação, não avaliação jurídica.</p>
<p>Confira especialmente a epígrafe, a ementa, as aspas de nova redação, os anexos, a assinatura e as quebras de linha da origem.</p>
<table><thead><tr><th>Nº</th><th>Linhas</th><th>Estilo aplicado</th><th>Texto</th><th>Conferir</th></tr></thead><tbody>
__ROWS__</tbody></table></html>"""
    page = page.replace("__COUNT__", str(len(items))).replace("__REVIEW__", str(review)).replace("__ROWS__", "\n".join(rows))
    report_path.write_text(page, encoding="utf-8")


def read_text(path: Path) -> str:
    data = path.read_bytes()
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252")


def run_cli(args: argparse.Namespace) -> None:
    source = read_text(args.input)
    items = classify(source, args.join_soft_wraps)
    output = args.output or args.input.with_name(args.input.stem + "_SEFA.docx")
    report = None
    if args.relatorio or args.report:
        report = args.report or output.with_name(output.stem + "_CONFERIR.html")
        if report.exists() and not args.force:
            raise FileExistsError(f"O relatório já existe: {report}")
    build_document(items, args.model, output, args.force)
    if report:
        write_report(items, report)
    print(f"Word: {output}" + (f"\nConferência: {report}" if report else "") +
          f"\nTrechos: {len(items)}; conferir: {sum(bool(x.attention) for x in items)}")
    if not args.sem_links:
        resumo = aplicar_e_registrar(output, report)
        print(texto_final(resumo, None, com_relatorio=report is not None))


def run_gui() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    class App:
        def __init__(self) -> None:
            self.root = tk.Tk()
            self.root.title("Formatador SEFA 3.18 — textos copiados")
            self.root.geometry("1080x780")
            self.root.minsize(760, 550)
            self.model_var = tk.StringVar(value=str(DEFAULT_MODEL))
            self.join_var = tk.BooleanVar(value=False)
            self.links_var = tk.BooleanVar(value=True)
            self.report_var = tk.BooleanVar(value=False)
            self.style_var = tk.StringVar()
            self.status_var = tk.StringVar(value="Cole o texto, analise, confira os estilos e salve.")
            self.items: list[Item] = []
            self.last_source = ""
            self.last_model = ""
            self.last_join = False
            self.layout()

        def layout(self) -> None:
            root = self.root
            root.columnconfigure(0, weight=1)
            root.rowconfigure(2, weight=2)
            root.rowconfigure(5, weight=3)
            top = ttk.Frame(root, padding=10)
            top.grid(row=0, column=0, sticky="ew")
            top.columnconfigure(1, weight=1)
            ttk.Label(top, text="Modelo Word SEFA").grid(row=0, column=0, sticky="w")
            ttk.Entry(top, textvariable=self.model_var).grid(row=0, column=1, sticky="ew", padx=8)
            ttk.Button(top, text="Escolher…", command=self.choose_model).grid(row=0, column=2)
            ttk.Label(root, text="1  Cole aqui o texto da publicação (Ctrl+V)", padding=(10, 3)).grid(row=1, column=0, sticky="w")
            frame = ttk.Frame(root, padding=(10, 0, 10, 0))
            frame.grid(row=2, column=0, sticky="nsew")
            frame.columnconfigure(0, weight=1)
            frame.rowconfigure(0, weight=1)
            self.source = tk.Text(frame, wrap="word", undo=True, font=("Arial", 10))
            self.source.grid(row=0, column=0, sticky="nsew")
            text_scroll = ttk.Scrollbar(frame, orient="vertical", command=self.source.yview)
            text_scroll.grid(row=0, column=1, sticky="ns")
            self.source.configure(yscrollcommand=text_scroll.set)
            actions = ttk.Frame(root, padding=10)
            actions.grid(row=3, column=0, sticky="ew")
            ttk.Button(actions, text="Colar da área de transferência", command=self.paste).pack(side="left", padx=(0, 8))
            ttk.Checkbutton(actions, text="Juntar linhas quebradas antes de palavra minúscula (revisar)", variable=self.join_var).pack(side="left", padx=(0, 8))
            ttk.Button(actions, text="2  Analisar texto", command=self.analyze).pack(side="right")
            ttk.Label(root, text="3  Confira a prévia; selecione uma linha para trocar o estilo", padding=(10, 3)).grid(row=4, column=0, sticky="w")
            tree_frame = ttk.Frame(root, padding=(10, 0, 10, 0))
            tree_frame.grid(row=5, column=0, sticky="nsew")
            tree_frame.columnconfigure(0, weight=1)
            tree_frame.rowconfigure(0, weight=1)
            self.table = ttk.Treeview(tree_frame, columns=("line", "style", "flag", "text"), show="headings", selectmode="browse")
            for key, title, width in [("line", "Linha", 75), ("style", "Estilo", 190), ("flag", "Atenção", 95), ("text", "Texto", 620)]:
                self.table.heading(key, text=title)
                self.table.column(key, width=width, stretch=(key == "text"))
            self.table.grid(row=0, column=0, sticky="nsew")
            self.table.bind("<<TreeviewSelect>>", self.select_item)
            scroll = ttk.Scrollbar(tree_frame, orient="vertical", command=self.table.yview)
            scroll.grid(row=0, column=1, sticky="ns")
            self.table.configure(yscrollcommand=scroll.set)
            self.table.tag_configure("review", background="#fff2d8")
            opcoes = ttk.Frame(root, padding=(10, 8, 10, 0))
            opcoes.grid(row=6, column=0, sticky="ew")
            ttk.Label(opcoes, text="Ao salvar:").pack(side="left")
            ttk.Checkbutton(opcoes, text="Criar hyperlinks de legislação",
                            variable=self.links_var).pack(side="left", padx=(8, 16))
            ttk.Checkbutton(opcoes, text="Gerar relatório de conferência",
                            variable=self.report_var).pack(side="left")
            bottom = ttk.Frame(root, padding=10)
            bottom.grid(row=7, column=0, sticky="ew")
            ttk.Label(bottom, text="Estilo da linha:").pack(side="left")
            self.styles = ttk.Combobox(bottom, textvariable=self.style_var, state="readonly", width=27)
            self.styles.pack(side="left", padx=8)
            ttk.Button(bottom, text="Aplicar estilo", command=self.override).pack(side="left")
            ttk.Button(bottom, text="4  Salvar Word", command=self.export).pack(side="right")
            ttk.Label(root, textvariable=self.status_var, padding=(10, 2)).grid(row=8, column=0, sticky="w")

        def choose_model(self) -> None:
            chosen = filedialog.askopenfilename(title="Escolha um Word da SEFA como modelo", filetypes=[("Word", "*.docx")])
            if chosen:
                self.model_var.set(chosen)

        def paste(self) -> None:
            try:
                value = self.root.clipboard_get()
                self.source.delete("1.0", "end")
                self.source.insert("1.0", value)
            except tk.TclError:
                messagebox.showerror("Área de transferência", "Não encontrei texto copiado.")

        def analyze(self) -> None:
            raw = self.source.get("1.0", "end-1c")
            model_path = Path(self.model_var.get())
            try:
                model = Document(str(model_path))
                self.items = classify(raw, self.join_var.get())
                apply_styles(self.items, model)
                available = paragraph_style_names(model)
                self.styles["values"] = sorted(available)
            except Exception as error:
                messagebox.showerror("Não foi possível analisar", str(error))
                return
            self.last_source, self.last_model, self.last_join = raw, self.model_var.get(), self.join_var.get()
            self.refresh()

        def refresh(self) -> None:
            self.table.delete(*self.table.get_children())
            for idx, item in enumerate(self.items):
                preview = item.text.replace("\t", " ")
                self.table.insert("", "end", iid=str(idx), values=(
                    f"{item.first_line}–{item.last_line}", item.style,
                    "Conferir" if item.attention else "", preview[:450]),
                    tags=("review",) if item.attention else ())
            n = sum(bool(x.attention) for x in self.items)
            self.status_var.set(f"{len(self.items)} trechos; {n} pedem conferência. Selecione uma linha amarela para ver o motivo.")

        def select_item(self, _event=None) -> None:
            selection = self.table.selection()
            if selection:
                item = self.items[int(selection[0])]
                self.style_var.set(item.style)
                detail = item.attention[0] if item.attention else item.reason
                self.status_var.set(f"Linha {item.first_line}: {detail}")

        def override(self) -> None:
            selection = self.table.selection()
            if not selection or not self.style_var.get():
                return
            item = self.items[int(selection[0])]
            item.style = self.style_var.get()
            item.overridden = True
            item.attention.append("Estilo escolhido manualmente na prévia.")
            self.refresh()
            self.table.selection_set(selection[0])

        def export(self) -> None:
            current = self.source.get("1.0", "end-1c")
            if not self.items or current != self.last_source or self.model_var.get() != self.last_model or self.join_var.get() != self.last_join:
                messagebox.showinfo("Atualize a prévia", "O texto, o modelo ou a opção de junção mudou. Clique em Analisar texto antes de salvar.")
                return
            chosen = filedialog.asksaveasfilename(title="Salvar documento formatado", defaultextension=".docx", filetypes=[("Word", "*.docx")], initialfile="ato_SEFA.docx")
            if not chosen:
                return
            output = Path(chosen)
            report = output.with_name(output.stem + "_CONFERIR.html") if self.report_var.get() else None
            if ((output.exists() or (report and report.exists()))
                    and not messagebox.askyesno("Substituir arquivos", "O Word ou seu relatório já existe. Deseja substituí-lo?")):
                return
            try:
                build_document(self.items, Path(self.model_var.get()), output, overwrite=True)
                if report:
                    write_report(self.items, report)
            except Exception as error:
                messagebox.showerror("Não foi possível salvar", str(error))
                return
            arquivos = f"Documento: {output}" + (f"\nConferência: {report}" if report else "")
            final = "Confira os itens destacados antes de usar o documento."
            if not self.links_var.get():
                show_and_open(messagebox, "Arquivos salvos", f"{arquivos}\n\n{final}", output)
                return
            executar_com_progresso(
                self.root, "Criando hyperlinks de legislação",
                lambda log, parar: aplicar_e_registrar(output, report, log, parar),
                lambda resumo, erro: show_and_open(
                    messagebox, "Arquivos salvos",
                    f"{arquivos}\n\n{texto_final(resumo, erro, report is not None)}\n\n{final}", output))

    App().root.mainloop()


def main() -> int:
    parser = argparse.ArgumentParser(description="Aplica os estilos da SEFA a texto normativo copiado.")
    parser.add_argument("--input", type=Path, help="Arquivo .txt da publicação; sem --input abre a interface gráfica")
    parser.add_argument("--model", type=Path, default=DEFAULT_MODEL, help="Documento .docx com estilos SEFA; seu conteúdo não será copiado")
    parser.add_argument("--output", type=Path, help="Arquivo .docx de saída")
    parser.add_argument("--relatorio", action="store_true",
                        help="Gera o relatório HTML de conferência ao lado do Word")
    parser.add_argument("--report", type=Path, help="Nome do relatório HTML (implica --relatorio)")
    parser.add_argument("--join-soft-wraps", action="store_true", help="Junta linhas quebradas antes de palavra minúscula")
    parser.add_argument("--force", action="store_true", help="Substitui arquivos de saída existentes")
    parser.add_argument("--sem-links", action="store_true",
                        help="Não cria os hyperlinks das leis, decretos etc. citados")
    args = parser.parse_args()
    try:
        if args.input:
            run_cli(args)
        else:
            run_gui()
    except Exception as error:
        print(f"Erro: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
